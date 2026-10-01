from __future__ import annotations

import io
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import urlparse

import pytest
from conftest import make_job
from docx import Document
from fastapi.testclient import TestClient

from clue_ai.database import get_profile, get_settings, save_profile, save_search_run
from clue_ai.domain import CandidateProfile, SearchCriteria
from clue_ai.repository import (
    add_source,
    get_run,
    get_run_results,
    get_source,
    list_sources,
    save_jobs,
    update_run,
)
from clue_ai.sources import FetchOutcome
from clue_ai.web import create_app

ORIGIN = {"Origin": "http://127.0.0.1"}


def test_local_views_render_without_a_jev_key(settings):
    client = TestClient(create_app(settings), base_url="http://127.0.0.1")

    for route in ("/", "/profile", "/search", "/x-leads", "/saved", "/hidden", "/sources", "/settings", "/privacy", "/health"):
        response = client.get(route)
        assert response.status_code == 200, route
    assert client.get("/health").json() == {"status": "ok", "storage": "local"}


@pytest.mark.parametrize(
    ("route", "current_href"),
    (
        ("/", "/"),
        ("/profile", "/profile"),
        ("/search", "/search"),
        ("/saved", "/saved"),
        ("/hidden", "/hidden"),
        ("/sources", "/sources"),
        ("/settings", "/settings"),
    ),
)
def test_primary_navigation_names_groups_and_marks_current_page(settings, route, current_href):
    client = TestClient(create_app(settings), base_url="http://127.0.0.1")

    response = client.get(route)

    assert response.status_code == 200
    assert 'aria-label="Your search"' in response.text
    assert 'aria-label="Workspace"' in response.text
    assert f'<a href="{current_href}" class="is-active" aria-current="page">' in response.text


def test_result_status_poller_only_runs_for_active_searches(settings):
    client = TestClient(create_app(settings), base_url="http://127.0.0.1")

    save_search_run(settings.database_path, "run-polling", SearchCriteria())
    active = client.get("/searches/run-polling")
    assert active.status_code == 200
    assert '/static/poll.js' in active.text

    update_run(
        settings.database_path,
        "run-polling",
        status="complete",
        stage="complete",
        message="Search complete.",
        completed=True,
    )
    complete = client.get("/searches/run-polling")
    assert complete.status_code == 200
    assert '/static/poll.js' not in complete.text


def test_jev_settings_state_external_data_and_local_budget_limits(settings):
    client = TestClient(create_app(settings), base_url="http://127.0.0.1")

    response = client.get("/settings")

    assert response.status_code == 200
    assert "Work history and skills can still identify you." in response.text
    assert "Clue does not inspect your TypeSafe account terms" in response.text
    assert "it cannot guarantee account-wide charges" in response.text
    assert "Clue does not read or control account refill settings" in response.text
    assert "I understand which profile facts and listing details leave this device" in response.text
    assert "I have reviewed the data disclosure and TypeSafe terms above" not in response.text
    assert "The built-in connector definitions and manual X lead marker remain." in response.text


def test_localhost_and_same_origin_boundaries_reject_cross_origin_requests(settings):
    app = create_app(settings)
    local = TestClient(app, base_url="http://127.0.0.1")
    remote = TestClient(app, base_url="http://attacker.example")

    assert remote.get("/").status_code == 421
    response = local.post(
        "/settings/jev-consent", data={"agree": "on"}, headers={"Origin": "https://attacker.example"}
    )
    assert response.status_code == 403


@pytest.mark.parametrize(
    "origin",
    ("http://127.0.0.1:8000", "http://localhost:8000", "http://[::1]:8000"),
)
def test_same_port_loopback_aliases_are_accepted_for_local_forms(settings, origin):
    client = TestClient(create_app(settings), base_url="http://127.0.0.1:8000")

    response = client.post(
        "/settings/jev-consent",
        data={"agree": "on"},
        headers={"Origin": origin},
        follow_redirects=False,
    )

    assert response.status_code == 303


@pytest.mark.parametrize("origin", ("https://127.0.0.1:8000", "http://127.0.0.1:8001"))
def test_local_form_boundary_rejects_scheme_or_port_mismatch(settings, origin):
    client = TestClient(create_app(settings), base_url="http://127.0.0.1:8000")

    response = client.post(
        "/settings/jev-consent",
        data={"agree": "on"},
        headers={"Origin": origin},
        follow_redirects=False,
    )

    assert response.status_code == 403


def test_source_management_requires_review_before_enabling(settings):
    client = TestClient(create_app(settings), base_url="http://127.0.0.1")
    response = client.post(
        "/sources/add",
        data={"kind": "greenhouse", "company": "Example", "identifier": "example"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    source_id = next(
        item["id"] for item in list_sources(settings.database_path) if item["id"].startswith("user-")
    )
    source = get_source(settings.database_path, source_id)
    assert response.status_code == 303
    assert source["state"] == "review"
    assert source["enabled"] == 0

    incomplete = client.post(
        f"/sources/{source_id}/approve",
        data={"terms_reviewed": "on"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert incomplete.status_code == 303
    assert get_source(settings.database_path, source_id)["enabled"] == 0

    reviewed = client.post(
        f"/sources/{source_id}/approve",
        data={
            key: "on"
            for key in (
                "terms_reviewed",
                "zero_cost",
                "attribution_confirmed",
                "robots_confirmed",
                "public_access_confirmed",
            )
        },
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert reviewed.status_code == 303
    assert get_source(settings.database_path, source_id)["state"] == "approved"
    assert get_source(settings.database_path, source_id)["enabled"] == 1


def test_x_manual_lead_is_user_entered_and_never_becomes_a_fetch_source(settings, monkeypatch):
    from clue_ai import services
    from clue_ai.filters import filter_jobs
    from clue_ai.repository import all_active_jobs, manual_x_leads, set_source_enabled, sources_due
    from clue_ai.sources import fetch_source

    client = TestClient(create_app(settings), base_url="http://127.0.0.1")
    built = client.get(
        "/x-leads",
        params={"roles": "Product Designer, UX Designer", "work_from": "Milan, Italy"},
    )
    assert built.status_code == 200
    assert "https://x.com/search?q=" in built.text
    assert "choose the Latest results tab" in built.text

    base_form = {
        "post_url": "https://twitter.com/hiring/status/123456789?ref=post",
        "job_url": "https://careers.example.com/jobs/123?utm_source=x",
        "title": "Product Designer",
        "company": "Example Studio",
        "location": "Remote in Italy and Europe",
        "workplace_type": "remote",
        "description": "Design product workflows for customers in Italy. Required: prototyping and research.",
        "reviewed": "on",
    }
    unsafe = dict(base_form, job_url="https://t.co/short")
    rejected = client.post(
        "/x-leads/add", data=unsafe, headers=ORIGIN, follow_redirects=False
    )
    assert rejected.status_code == 303
    assert "shortened" in rejected.headers["location"]
    assert manual_x_leads(settings.database_path) == []

    saved = client.post(
        "/x-leads/add", data=base_form, headers=ORIGIN, follow_redirects=False
    )
    assert saved.status_code == 303
    lead = manual_x_leads(settings.database_path)[0]
    assert lead["post_url"] == "https://x.com/hiring/status/123456789"
    assert lead["job_url"] == "https://careers.example.com/jobs/123"
    page = client.get("/x-leads")
    assert "Open listing · careers.example.com" in page.text
    assert "Open X post" in page.text
    indexed = filter_jobs(
        all_active_jobs(settings.database_path),
        SearchCriteria(roles="Product Designer", work_from="Italy", workplace="remote"),
    )
    assert len(indexed) == 1
    assert indexed[0]["freshness_status"] == "manual"

    set_source_enabled(settings.database_path, "x_manual", True)
    assert get_source(settings.database_path, "x_manual")["enabled"] == 0
    assert "x_manual" not in {source["id"] for source in sources_due(settings.database_path)}
    outcome = fetch_source(get_source(settings.database_path, "x_manual"), SearchCriteria(), settings)
    assert outcome.skipped
    assert "never fetched" in outcome.message

    cannot_enable = client.post("/sources/x_manual/enable", headers=ORIGIN, follow_redirects=False)
    assert cannot_enable.status_code == 404

    monkeypatch.setattr(
        services,
        "fetch_source",
        lambda *_args, **_kwargs: FetchOutcome(message="Fixture; no network request.", skipped=True),
    )
    searched = client.post(
        "/search",
        data={"roles": "Product Designer", "work_from": "Italy", "workplace": "remote"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert searched.status_code == 303
    results = client.get(searched.headers["location"])
    assert results.status_code == 200
    assert "Manual lead · live status unverified" in results.text
    assert "Employer listing · careers.example.com" in results.text
    assert "X post" in results.text
    assert "availability not rechecked" in results.text


def test_lever_source_form_records_the_selected_region(settings):
    client = TestClient(create_app(settings), base_url="http://127.0.0.1")
    response = client.post(
        "/sources/add",
        data={
            "kind": "lever",
            "company": "Prima",
            "identifier": "prima",
            "lever_region": "eu",
        },
        headers=ORIGIN,
        follow_redirects=False,
    )
    source = next(item for item in list_sources(settings.database_path) if item["name"] == "Prima")

    assert response.status_code == 303
    assert source["state"] == "review"
    assert source["enabled"] == 0
    assert source["endpoint"] == "https://api.eu.lever.co/v0/postings/prima?mode=json"
    assert source["config"]["region"] == "eu"


def test_search_to_results_save_hide_and_delete_uses_mocked_sources_only(
    settings, monkeypatch
):
    from clue_ai import services

    client = TestClient(create_app(settings), base_url="http://127.0.0.1")
    add_source(
        settings.database_path,
        name="Under review",
        kind="greenhouse",
        endpoint="https://boards-api.greenhouse.io/v1/boards/review/jobs",
        config={"board_token": "review"},
        attribution="Under review",
    )
    calls = []

    base_job = make_job()

    def fake_fetch(source, *_args):
        calls.append(source["id"])
        job = make_job(
            source_id=source["id"],
            url=base_job.canonical_url,
            title=base_job.title,
            company=base_job.company,
            location=base_job.location_raw,
            description=base_job.description,
        )
        job.posted_at = base_job.posted_at
        return FetchOutcome(
            jobs=[job],
            checked=1,
            response_bytes=1_024,
            raw_records=1,
            status_counts={"status_200": 1},
        )

    monkeypatch.setattr(services, "fetch_source", fake_fetch)
    response = client.post(
        "/search",
        data={"roles": "Software Engineer", "work_from": "Italy", "workplace": "remote"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert response.status_code == 303
    results_url = response.headers["location"]
    run_id = results_url.rsplit("/", 1)[-1]
    results = client.get(results_url)

    assert results.status_code == 200
    assert "Software Engineer" in results.text
    assert 'href="https://jobs.example.org/openings/software-engineer"' in results.text
    assert "explicitly enable Jev scoring" in results.text
    assert "HTTP 200" in results.text
    assert set(calls) == {"jobicy", "remotejobs", "remoteok", "remotefirstjobs", "startupjobs"}
    assert "Powered by RemoteJobs.org" in results.text
    assert not any(call.startswith("user-") for call in calls)
    row = get_run_results(settings.database_path, run_id)[0]
    assert row["eligibility_status"] == "eligible"
    assert row["eligibility_evidence"]
    assert row["score_state"] == "unscored"

    saved = client.post(
        f"/jobs/{row['id']}/save",
        data={"return_to": results_url},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert saved.status_code == 303
    saved_page = client.get("/saved")
    assert "Software Engineer" in saved_page.text
    assert "Powered by RemoteJobs.org" in saved_page.text
    hidden = client.post(
        f"/jobs/{row['id']}/hide",
        data={"return_to": results_url},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert hidden.status_code == 303
    hidden_page = client.get("/hidden")
    assert "Software Engineer" in hidden_page.text
    assert "Powered by RemoteJobs.org" in hidden_page.text

    deleted = client.post(
        "/data/delete",
        data={"confirmation": "DELETE"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert deleted.status_code == 303
    assert "Local+profile" in deleted.headers["location"]
    assert get_profile(settings.database_path) == CandidateProfile()


def test_cv_upload_is_held_for_review_then_saved_locally_and_removable(settings):
    document = Document()
    document.add_heading("Summary", level=1)
    document.add_paragraph(
        "Synthetic candidate profile with enough selectable text to test local document extraction."
    )
    output = io.BytesIO()
    document.save(output)
    client = TestClient(create_app(settings), base_url="http://127.0.0.1")
    extracted = client.post(
        "/profile/extract",
        files={"cv_file": ("synthetic.docx", output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert extracted.status_code == 303
    pending_id = urlparse(extracted.headers["location"]).query.split("=")[-1]
    review_page = client.get(extracted.headers["location"])
    assert "Synthetic candidate profile" in review_page.text

    saved = client.post(
        "/profile/save",
        data={
            "pending_id": pending_id,
            "target_roles": "Software Engineer",
            "skills": "Python, SQL",
            "profile_language": "en",
        },
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert saved.status_code == 303
    profile = get_profile(settings.database_path)
    assert profile.target_roles == "Software Engineer"
    assert profile.cv_path
    assert profile.extracted_text.startswith("Summary")

    removed = client.post("/profile/remove-cv", headers=ORIGIN, follow_redirects=False)
    assert removed.status_code == 303
    profile = get_profile(settings.database_path)
    assert profile.target_roles == "Software Engineer"
    assert profile.cv_path == ""
    assert profile.extracted_text == ""


def test_cv_first_workflow_saves_profile_searches_and_scores_automatically(settings, monkeypatch):
    from clue_ai import jev, services
    from clue_ai.services import run_jev_scoring as score_jev

    configured = replace(settings, api_key="test-key")
    client = TestClient(create_app(configured), base_url="http://127.0.0.1")
    monkeypatch.setattr(services, "sources_due", lambda *_args: [])
    save_jobs(
        settings.database_path,
        [
            make_job(
                title="Senior Product Designer",
                location="Italy and Europe",
                description=(
                    "We need a senior product designer for our fully remote European team. "
                    "You will lead user research, build prototypes, improve design systems, "
                    "work with product managers, and make accessible software for customers."
                ),
            )
        ],
    )

    document = Document()
    document.add_paragraph("Alex Rivera")
    document.add_paragraph("alex@example.test")
    document.add_heading("Summary", level=1)
    document.add_paragraph(
        "Product designer focused on accessible software and clear customer experiences."
    )
    document.add_heading("Skills", level=1)
    document.add_paragraph("Figma, user research, prototyping, design systems, accessibility")
    document.add_heading("Experience", level=1)
    document.add_paragraph("Senior Product Designer | Northstar Studio | 2022 – Present")
    document.add_paragraph(
        "Led product design for a remote team, improving onboarding and usability."
    )
    document.add_heading("Languages", level=1)
    document.add_paragraph("English C1, Italian B2")
    file_bytes = io.BytesIO()
    document.save(file_bytes)
    sent_states = []
    scoring_stages = []

    def capture_scoring_stage(database_path, runtime_settings, run_id):
        run = get_run(database_path, run_id)
        scoring_stages.append((run["status"], run["stage"]))
        score_jev(database_path, runtime_settings, run_id)

    monkeypatch.setattr(services, "run_jev_scoring", capture_scoring_stage)

    class FakeClient:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def system_one(self, *, state, questions):
            sent_states.append(state)
            choices = {
                name: SimpleNamespace(choice="4", confidence=0.9, probabilities={"4": 1.0})
                for name in questions
            }
            return SimpleNamespace(
                usage=SimpleNamespace(input_tokens=500),
                model="jev-test",
                choices=choices,
            )

    monkeypatch.setattr(jev, "_default_client_factory", lambda **_kwargs: FakeClient())
    response = client.post(
        "/workflow/start",
        data={"jev_auto_score": "on"},
        files={
            "cv_file": (
                "candidate.docx",
                file_bytes.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=ORIGIN,
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"].startswith("/searches/")
    run_id = response.headers["location"].rsplit("/", 1)[-1]
    profile = get_profile(settings.database_path)
    assert profile.target_roles == "Senior Product Designer"
    assert profile.profile_language == "en"
    assert profile.cv_path and Path(profile.cv_path).is_file()
    assert "alex@example.test" in profile.extracted_text
    run = get_run(settings.database_path, run_id)
    assert run["status"] == "complete"
    assert run["criteria"]["roles"] == "Senior Product Designer"
    assert run["criteria"]["work_from"] == "Italy"
    assert run["criteria"]["workplace"] == "remote"
    result = get_run_results(settings.database_path, run_id)[0]
    assert result["score_state"] == "scored"
    assert result["combined_score"] == 1.0
    assert get_settings(settings.database_path)["jev_consent_at"]
    assert len(sent_states) == 1
    assert scoring_stages == [("scoring", "jev")]
    assert "alex@example.test" not in str(sent_states[0])
    assert "candidate.docx" not in str(sent_states[0])

    page = client.get(response.headers["location"])
    assert "1 listings scored with Jev" in page.text
    assert "Score with Jev" not in page.text


def test_cv_first_workflow_keeps_listings_unscored_without_opt_in(settings, monkeypatch):
    from clue_ai import jev, services

    client = TestClient(create_app(settings), base_url="http://127.0.0.1")
    monkeypatch.setattr(services, "sources_due", lambda *_args: [])
    monkeypatch.setattr(
        jev,
        "_default_client_factory",
        lambda **_kwargs: pytest.fail("Jev client must not be created without opt-in and a key"),
    )
    save_jobs(settings.database_path, [make_job(title="Product Designer")])

    document = Document()
    document.add_heading("Summary", level=1)
    document.add_paragraph(
        "Product designer with experience building useful software and supporting remote teams."
    )
    document.add_heading("Experience", level=1)
    document.add_paragraph("Product Designer | Example Studio | 2021 – Present")
    document.add_paragraph(
        "Designed product workflows, developed research plans, and worked with distributed teams."
    )
    file_bytes = io.BytesIO()
    document.save(file_bytes)

    response = client.post(
        "/workflow/start",
        files={
            "cv_file": (
                "candidate.docx",
                file_bytes.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=ORIGIN,
        follow_redirects=False,
    )

    assert response.status_code == 303
    run_id = response.headers["location"].rsplit("/", 1)[-1]
    run = get_run(settings.database_path, run_id)
    assert run["status"] == "complete"
    row = get_run_results(settings.database_path, run_id)[0]
    assert row["score_state"] == "unscored"
    assert "explicitly enable Jev scoring" in row["score_reason"]
    assert not get_settings(settings.database_path)["jev_consent_at"]


def test_repeat_cv_workflow_uses_saved_profile_and_search_preferences(settings, monkeypatch):
    from clue_ai import services

    client = TestClient(create_app(settings), base_url="http://127.0.0.1")
    monkeypatch.setattr(services, "sources_due", lambda *_args: [])
    save_profile(
        settings.database_path,
        CandidateProfile(
            target_roles="UX Researcher",
            skills="User research",
            profile_language="en",
        ),
    )
    save_search_run(
        settings.database_path,
        "previous-preferences",
        SearchCriteria(
            roles="UX Researcher",
            work_from="Milan, Italy",
            workplace="remote",
            minimum_salary="55000",
        ),
    )
    update_run(
        settings.database_path,
        "previous-preferences",
        status="complete",
        stage="done",
        message="Previous search completed.",
        completed=True,
    )

    response = client.post(
        "/workflow/start",
        headers=ORIGIN,
        follow_redirects=False,
    )

    assert response.status_code == 303
    run_id = response.headers["location"].rsplit("/", 1)[-1]
    run = get_run(settings.database_path, run_id)
    assert run["criteria"]["roles"] == "UX Researcher"
    assert run["criteria"]["work_from"] == "Milan, Italy"
    assert run["criteria"]["minimum_salary"] == "55000"


def test_search_worker_marks_unexpected_workflow_failure(settings, database, monkeypatch):
    from clue_ai import services

    save_search_run(settings.database_path, "run-worker-failure", SearchCriteria())

    def fail_search(*_args, **_kwargs):
        raise RuntimeError("synthetic worker failure")

    monkeypatch.setattr(services, "run_search", fail_search)
    services.run_search_worker(settings.database_path, settings, "run-worker-failure")

    run = get_run(settings.database_path, "run-worker-failure")
    assert run["status"] == "failed"
    assert run["stage"] == "error"
    assert run["error"] == "Local search error: RuntimeError."
