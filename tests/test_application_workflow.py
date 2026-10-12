from __future__ import annotations

import io
import json
import re
import urllib.error
from copy import deepcopy
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from email import policy
from email.parser import BytesParser
from hashlib import sha256
from pathlib import Path
from typing import ClassVar
from urllib.parse import urlsplit

import pytest
from conftest import make_job
from docx import Document
from fastapi.testclient import TestClient

from clue_ai.application_followups import list_followups, schedule_followup
from clue_ai.application_prep import (
    add_source,
    get_preparation,
    list_claims,
    request_preparation,
    retry_preparation,
    review_claim,
    set_source_options,
    suggest_claims,
)
from clue_ai.application_prompts import (
    CRAWL_TOOL,
    LETTER_REVIEW_DIMENSION_NAMES,
    OUTPUT_SCHEMA_VERSION,
    OUTPUT_SCHEMAS,
    PROMPT_VERSION,
    ROLE_PROMPTS,
)
from clue_ai.application_workflow import (
    MAX_OUTPUT_TOKENS,
    OpenAIProviderError,
    StalePreparationError,
    _api_call,
    _apply_grounding_results,
    _assert_packet_quality_gates,
    _bind_recruiter_requirements,
    _call_stage,
    _contains_role_duty_inventory,
    _cover_letter_quality,
    _cover_letter_title,
    _filter_diagnoser_references,
    _get_bound_inputs,
    _grounding_assertions,
    _letter_repair_context,
    _looks_like_heading,
    _looks_like_subheading,
    _merge_letter_revisions,
    _packet_review_note,
    _parse_stage_result,
    _payload,
    _preferred_cover_letter_profile_claim_ids,
    _preferred_profile_letter_anchors,
    _render_artifacts,
    _repair_letter_requirement_ids,
    _research_context,
    _resume_contact_lines,
    _resume_render_blocks,
    _resume_signature_name,
    _rewrite_change_summary,
    _run_researcher,
    _select_technical_profile_evidence,
    _technical_profile_prompt_context,
    _validate_letter_review,
    _validate_research,
    _validate_rewrite,
    create_practice_session,
    get_packet,
    queue_practice_answers,
    run_interview_practice_assessment,
    run_interview_practice_questions,
    run_preparation,
    verify_packet_approval,
)
from clue_ai.config import Settings
from clue_ai.database import (
    connect,
    delete_personal_data,
    get_settings,
    recover_interrupted_gmail_drafts,
    save_openai_controls,
    save_search_run,
    set_gmail_connection,
)
from clue_ai.domain import SearchCriteria
from clue_ai.gmail_drafts import (
    GMAIL_DRAFTS_URL,
    GmailDraftError,
    create_approved_packet_draft,
    create_unsent_draft,
)
from clue_ai.openai_provider import (
    DEFAULT_GENERATION_TIMEOUT_SECONDS,
    INPUT_TOKEN_COUNT_TIMEOUT_SECONDS,
    LONG_CONTEXT_INPUT_RATE_MULTIPLIER,
    LONG_CONTEXT_OUTPUT_RATE_MULTIPLIER,
    LONG_CONTEXT_SURCHARGE_THRESHOLD_TOKENS,
    MODEL_CONTEXT_WINDOW_TOKENS,
    MODEL_ID,
    REASONING_EFFORT,
    OpenAIBudgetError,
    OpenAIConfigurationError,
    ResponsesClient,
    ensure_context_fits,
    reserve_usage,
    settle_usage,
)
from clue_ai.repository import save_jobs, save_run_results, update_run
from clue_ai.reset import reset_search_data
from clue_ai.scrapling_research import BoundedResearchCrawler, ResearchCrawlError
from clue_ai.web import create_app

ORIGIN = {"Origin": "http://127.0.0.1"}
JOB_URL = "https://jobs.example.org/openings/software-engineer"
APPLICATION_URL = "https://jobs.example.org/openings/software-engineer/apply"
OFFICIAL_SOURCE_URL = "https://careers.example.org/jobs/software-engineer-intern"

def _opening_text(
    role_title="AI Engineer",
    company="Northstar",
    responsibilities=("model deployment", "agent-assisted workflows"),
):
    employer = f" at {company}" if company else ""
    first, second = list(responsibilities[:2]) if len(responsibilities) >= 2 else (
        "the advertised technical work", "the team's stated responsibilities"
    )
    return (
        f"I am applying for the {role_title} role{employer}. "
        f"The posting emphasizes {first} and {second}, core responsibilities in its scientific AI work and collaboration across research teams, with reliable operation across research and internal environments."
    )


def _complete_quality_letter(
    *,
    preferred_profile_claim_ids=None,
    registration_first=False,
    role_title="AI Engineer",
    company="Northstar",
    responsibilities=("model deployment", "agent-assisted workflows"),
):
    preferred = list(preferred_profile_claim_ids or [])
    if registration_first:
        registration_claim = preferred[0] if preferred else "claim-registration"
        ayni_claim = preferred[1] if len(preferred) > 1 else "claim-ayni"
    else:
        ayni_claim = preferred[0] if preferred else "claim-ayni"
        registration_claim = preferred[1] if len(preferred) > 1 else "claim-registration"
    opening = {
        "section": "opening",
        "text": _opening_text(role_title, company, responsibilities),
        "claim_ids": [],
        "requirement_ids": ["REQ-01"],
        "source_urls": [JOB_URL],
    }
    ayni = {
        "section": "evidence",
        "text": (
            "I worked on Ayni, a platform for creating and operating specialized AI-agent workflows. "
            "I contributed to workflow design in Python services, including model routing, retrieval, tool execution, structured run validation, "
            "and interfaces for starting runs and reviewing results. The project also includes evaluation and "
            "approval workflows with recovery safeguards. Northstar calls for developing agent-assisted workflows in shared services."
        ),
        "claim_ids": [ayni_claim],
        "requirement_ids": ["REQ-02"],
        "source_urls": [JOB_URL],
    }
    registration = {
        "section": "evidence",
        "text": (
            "For Medical Image Registration, I deployed a VoxelMorph pipeline on the AMD Kria KV260 DPU. "
            "The quantized model ran on FPGA hardware, cutting inference and image warping by 66% against "
            "a CPU baseline while maintaining alignment quality. This keeps runtime "
            "and registration quality in view during deployment. Northstar names model deployment as an AI Engineer responsibility."
        ),
        "claim_ids": [registration_claim],
        "requirement_ids": ["REQ-01"],
        "source_urls": [JOB_URL],
    }
    evidence = [registration, ayni] if registration_first else [ayni, registration]
    contribution = {
        "section": "contribution",
        "text": (
            "Together, these projects span agent workflow operation and scientific model deployment, including controls for how runs are executed and quality is measured. "
            "That evaluation-and-operations focus matters for Northstar's shared AI services, where systems must support research teams and remain maintainable in production."
        ),
        "claim_ids": [ayni_claim, registration_claim],
        "requirement_ids": ["REQ-02"],
        "source_urls": [JOB_URL],
    }
    closing = {
        "section": "closing",
        "text": (
            "Thank you for considering my application. I would welcome a conversation about Northstar's scientific AI work in this position."
        ),
        "claim_ids": [],
        "requirement_ids": ["REQ-01"],
        "source_urls": [JOB_URL],
    }
    return {"cover_letter_paragraphs": [opening, *evidence, contribution, closing]}


def _supported_letter_grounding(count=5):
    return {
        "results": [
            {
                "id": f"cover_letter_paragraph_{index}",
                "output_type": "cover_letter_paragraph",
                "status": "supported",
            }
            for index in range(count)
        ]
    }



def _approve_synthetic_tailored_resume(*_args, **_kwargs):
    return {
        "rubric_version": "tailored-resume-fit-v2",
        "status": "approved",
        "model_status": "approved",
        "confidence": 0.92,
        "confidence_policy": "diagnostic only",
        "reason": "Synthetic Jev reviewer approved the role fit.",
        "diagnostic_reason_code": "no_material_issue",
        "diagnostic_reason": "No concrete material issue prevents owner review.",
        "model": "jev-test",
        "actual_input_tokens": 0,
        "actual_cost_usd": 0.0,
    }


@pytest.fixture(autouse=True)
def synthetic_claim_support(monkeypatch):
    """Keep workflow tests provider-free; claim_support has its own SDK contract tests."""
    def accept_synthetic_assertions(_database_path, _settings, _request_id, assertions):
        return {
            "status": "complete",
            "model": "jev-test",
            "actual_cost_usd": 0.0,
            "results": [
                {"id": item["id"], "status": "supported", "confidence": 0.9}
                for item in assertions
            ],
        }

    monkeypatch.setattr(
        "clue_ai.application_workflow.check_generated_claim_support",
        accept_synthetic_assertions,
    )


def _save_match(
    database_path: Path,
    *,
    filter_status="match",
    eligibility_status="eligible",
    description=None,
) -> str:
    job = make_job(url=JOB_URL)
    if description is not None:
        job = replace(job, description=description)
    save_jobs(database_path, [job])
    with connect(database_path) as db:
        job_id = db.execute("SELECT id FROM jobs LIMIT 1").fetchone()["id"]
    save_search_run(database_path, "run-workflow", SearchCriteria())
    update_run(database_path, "run-workflow", status="complete", completed=True)
    save_run_results(
        database_path,
        "run-workflow",
        [{"id": job_id, "filter_status": filter_status, "eligibility_status": eligibility_status}],
        score_state="scored",
    )
    return job_id


def _approved_claim(database_path: Path, source_id: str) -> str:
    suggest_claims(database_path, source_id)
    # Multiple suggestions can share a creation timestamp; UUID ordering is random.
    # Pick by source text so multi-line profile fixtures stay deterministic.
    claim = min(
        list_claims(database_path, source_id),
        key=lambda item: (item["claim_text"].casefold(), item["claim_text"]),
    )
    review_claim(
        database_path,
        claim["id"],
        status="approved",
        evidence_level="implemented",
        category="systems",
        role_family="applied_ai_llm",
        owner_note="Synthetic owner-approved test evidence.",
    )
    return claim["id"]


def _create_preparation(
    settings: Settings,
    database_path: Path,
    *,
    structure_policy="preserve",
    technical_text=(
        "Implemented reliable Python services and a typed local workflow.\n"
        "Tested a simulated maintenance workflow with synthetic reports to measure duplicate handling.\n"
    ),
    technical_authorship="unknown",
    descriptive_text="I prefer direct, modest writing and collaborative technical teams.\n",
    descriptive_authorship="owner_written",
    filter_status="match",
    eligibility_status="eligible",
    confirm_jev_conflict=False,
    confirm_location_eligibility=False,
    description=None,
    additional_source_url="",
):
    technical = add_source(
        database_path,
        settings,
        "synthetic-technical.md",
        "technical_profile",
        technical_text.encode(),
        authorship_label=technical_authorship,
    )
    set_source_options(
        database_path,
        technical["id"],
        permitted=True,
        default_cv=False,
        structure_policy="preserve",
    )
    claim_id = _approved_claim(database_path, technical["id"])
    resume_text = (
        "Education\nMSc Computer Science\nExperience\n"
        "- Built Python services for Example Labs.\n"
    )
    cv = add_source(
        database_path,
        settings,
        "synthetic-resume.md",
        "resume",
        resume_text.encode(),
    )
    set_source_options(
        database_path,
        cv["id"],
        permitted=True,
        default_cv=True,
        structure_policy=structure_policy,
    )
    descriptive = add_source(
        database_path,
        settings,
        "synthetic-descriptive.md",
        "descriptive_profile",
        descriptive_text.encode(),
        authorship_label=descriptive_authorship,
    )
    set_source_options(
        database_path,
        descriptive["id"],
        permitted=True,
        default_cv=False,
        structure_policy="preserve",
    )
    sample = add_source(
        database_path,
        settings,
        "synthetic-writing-sample.md",
        "writing_sample",
        b"This synthetic sample is not selected and must stay out of prompts.\n",
        authorship_label="owner_written",
    )
    job_id = _save_match(
        database_path,
        filter_status=filter_status,
        eligibility_status=eligibility_status,
        description=description,
    )
    record, created = request_preparation(
        database_path,
        job_id,
        "run-workflow",
        cv["id"],
        additional_source_url,
        confirm_jev_conflict,
        confirm_location_eligibility,
    )
    assert created
    return get_preparation(database_path, record["id"]), cv, technical, descriptive, sample, claim_id, resume_text


def _enable_synthetic_openai(settings: Settings, database_path: Path) -> Settings:
    save_openai_controls(
        database_path,
        consent=True,
        monthly_cap_usd=5.0,
        opportunity_cap_usd=2.0,
        input_usd_per_million=0.1,
        output_usd_per_million=0.5,
        rate_card_revision=datetime.now(timezone.utc).date().isoformat(),
    )
    return replace(settings, openai_api_key="synthetic-openai-key")


@pytest.mark.parametrize(
    ("stage", "expected_message"),
    [
        ("diagnoser", "Diagnoser is checking selected CV lines"),
        ("recruiter", "Recruiter is mapping the job requirements"),
        ("rewriter", "Rewriter is drafting the tailored CV"),
        ("letter_reviewer", "Jev is reviewing the cover letter"),
        ("letter_reviser", "Rewriter is revising the cover letter"),
    ],
)
def test_generation_agent_stage_is_saved_before_api_call(
    settings, database, monkeypatch, stage, expected_message
):
    record, *_ = _create_preparation(settings, database)
    monkeypatch.setattr("clue_ai.application_workflow._payload", lambda *_args, **_kwargs: {"input": []})
    monkeypatch.setattr("clue_ai.application_workflow._api_call", lambda *_args, **_kwargs: {})
    monkeypatch.setattr("clue_ai.application_workflow._parse_stage_result", lambda *_args: {"ok": True})

    result = _call_stage(database, settings, {}, record, stage, {}, object(), [])

    assert result == {"ok": True}
    saved = get_preparation(database, record["id"])
    assert saved["state"] == "generating"
    assert saved["status_message"].startswith(expected_message)
    assert saved["worker_heartbeat_at"]


def test_blocked_public_crawl_is_reported_while_research_continues(settings, database, monkeypatch):
    record, *_ = _create_preparation(settings, database)
    calls = iter(
        [
            {
                "output": [
                    {
                        "type": "function_call",
                        "name": "crawl_public_page",
                        "call_id": "crawl-1",
                        "arguments": json.dumps(
                            {"url": "https://blocked.example.org/jobs/role", "research_question": "Find application details."}
                        ),
                    }
                ]
            },
            {"output": [{"type": "message", "content": []}]},
        ]
    )

    class BlockedCrawler:
        def __init__(self, _settings, _allowed_urls):
            self.allowed_urls = {JOB_URL}
            self.pages = {}

        def crawl(self, _url, _question):
            raise RuntimeError(
                "Max retries exceeded for blocked request: https://blocked.example.org/jobs/role"
            )

    monkeypatch.setattr(
        "clue_ai.application_workflow._api_call",
        lambda *_args, **_kwargs: next(calls),
    )
    monkeypatch.setattr("clue_ai.application_workflow._parse_stage_result", lambda *_args: {"findings": []})
    monkeypatch.setattr(
        "clue_ai.application_workflow._validate_research",
        lambda research, _pages: research,
    )

    result = _run_researcher(
        database,
        settings,
        record,
        record["snapshot"],
        object(),
        lambda _key: object(),
        BlockedCrawler,
        [],
    )

    saved = get_preparation(database, record["id"])
    assert result["findings"] == []
    assert result["crawl_warnings"][0]["host"] == "blocked.example.org"
    assert result["crawl_warnings"][0]["reason"] == "the site blocked or challenged the request"
    assert saved["state"] == "researching"
    assert "blocked.example.org" in saved["status_message"]
    assert "blocked or unavailable" in saved["status_message"]
    assert "saved listing and any reachable allowed sources" in saved["status_message"]


class FakeCrawler:
    instances: ClassVar[list[FakeCrawler]] = []
    inspect_application_page: ClassVar[bool] = False
    application_instruction_quote: ClassVar[str] = (
        "To apply, email your CV and cover letter to recruitment@example.org."
    )

    def __init__(self, _settings, allowed_urls):
        self.allowed_urls = set(allowed_urls)
        self.pages = {}
        self.__class__.instances.append(self)

    def crawl(self, url, research_question):
        if self.inspect_application_page and url == APPLICATION_URL:
            page = {
                "url": url,
                "title": "Apply for Software Engineer",
                "text": (
                    self.application_instruction_quote + " "
                    "Tell us about a Python service you have built."
                ),
                "links": [],
                "form_fields": [
                    {
                        "field_id": "why-python-service",
                        "label": "Tell us about a Python service you have built.",
                        "control_type": "textarea",
                        "required": True,
                        "options": [],
                        "is_open_question": True,
                    }
                ],
                "observed_at": "2026-10-06T10:01:00+00:00",
                "research_question": research_question,
            }
            self.pages[url] = page
            return page
        page = {
            "url": url,
            "title": "Example Labs Team",
            "text": (
                "Example Labs builds Python services for software teams. "
                "Alex Rivera, Engineering Manager at Example Labs, can be contacted at alex@example.org."
            ),
            "links": (
                [{"label": "Apply now", "url": APPLICATION_URL}]
                if self.inspect_application_page
                else []
            ),
            "observed_at": "2026-10-06T10:00:00+00:00",
            "research_question": research_question,
        }
        self.pages[url] = page
        return page


def _message_result(value):
    return {
        "model": MODEL_ID,
        "status": "completed",
        "output": [
            {"type": "message", "content": [{"type": "output_text", "text": json.dumps(value)}]}
        ],
        "usage": {"input_tokens": 150, "output_tokens": 80},
    }


class FakePreparationClient:
    instances: ClassVar[list[FakePreparationClient]] = []
    omit_preferred_profile_claims: ClassVar[bool] = False
    short_letter_mode: ClassVar[bool] = False
    force_letter_review: ClassVar[bool] = False
    force_profile_coverage_gap: ClassVar[bool] = False
    inspect_application_page: ClassVar[bool] = False

    def __init__(self, _key):
        self.calls = []
        self.letter_review_calls = 0
        self.instances.append(self)

    def count_input_tokens(self, payload):
        encoded = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        return max(200, len(encoded) // 4)

    def create(self, payload):
        self.calls.append(payload)
        assert payload["model"] == MODEL_ID
        assert payload["reasoning"]["effort"] == REASONING_EFFORT
        assert payload["store"] is False
        stage = payload["text"]["format"]["name"].removeprefix("clue_").rsplit("_v", 1)[0]
        prompt_input = payload["input"]
        if stage == "researcher" and payload.get("tools"):
            assert payload["tools"] == [CRAWL_TOOL]
            tool_outputs = [
                item for item in prompt_input
                if item.get("type") == "function_call_output"
            ]
            if not tool_outputs:
                next_url = JOB_URL
            elif (
                self.inspect_application_page
                and json.loads(tool_outputs[-1].get("output", "{}")).get("url") == JOB_URL
            ):
                next_url = APPLICATION_URL
            else:
                next_url = ""
            if next_url:
                return {
                    "model": MODEL_ID,
                    "status": "completed",
                    "output": [
                        {
                            "type": "function_call",
                            "name": "crawl_public_page",
                            "arguments": json.dumps(
                                {"url": next_url, "research_question": "Inspect role and application details"}
                            ),
                            "call_id": "synthetic-call-1",
                        }
                    ],
                    "usage": {"input_tokens": 100, "output_tokens": 30},
                }
            return _message_result(
                {
                    "company_summary": "Example Labs builds Python services.",
                    "findings": [
                        {
                            "topic": "engineering team",
                            "fact": "The team builds Python services.",
                            "source_url": JOB_URL,
                            "quote": "Example Labs builds Python services",
                            "confidence": "high",
                        }
                    ],
                    "contacts": [
                        {
                            "name": "Alex Rivera",
                            "role": "Engineering Manager",
                            "organization": "Example Labs",
                            "contact_type": "likely_team_lead",
                            "source_url": JOB_URL,
                            "quote": "Alex Rivera, Engineering Manager at Example Labs",
                            "public_email": "alex@example.org",
                            "confidence": "high",
                            "function_match": "high",
                        }
                    ],
                    "no_contact_found_reason": "",
                    "unresolved_questions": [],
                }
            )
        assert "tools" not in payload
        context = json.loads(prompt_input[0]["content"])
        if stage == "diagnoser":
            cv = context["cv"]
            return _message_result(
                {
                    "diagnostics": [
                        {
                            "source_id": cv["source_id"],
                            "line_id": cv["lines"][0]["id"],
                            "issue": "A section heading could be clearer.",
                            "severity": "low",
                            "suggested_fix": "Use a standard section name.",
                        }
                    ],
                    "overall_note": "These are visible text-parseability observations.",
                }
            )
        if stage == "recruiter":
            claim_id = context["approved_claims"][0]["id"]
            profile_claim_ids = [
                evidence["id"]
                for profile in context.get("technical_profile_sources", [])
                for evidence in profile.get("evidence", [])
            ]
            mapped_claim_ids = [claim_id, *profile_claim_ids]
            requirements = [
                ("Python service experience", "required", "L0003"),
                ("Reliable local workflow testing", "preferred", "L0004"),
            ]
            if self.short_letter_mode:
                requirements = [
                    ("Build agent-assisted workflows", "required", "L0003"),
                    ("Deploy AI models on edge hardware", "preferred", "L0004"),
                ]
            return _message_result(
                {
                    "selected_cv_id": context["permitted_cv_id"],
                    "requirements": [
                        {
                            "requirement": requirement,
                            "priority": priority,
                            "document_coverage": "covered",
                            "claim_ids": mapped_claim_ids,
                            "cv_line_ids": [line_id],
                            "notes": "Supported by owner-approved evidence.",
                        }
                        for requirement, priority, line_id in requirements
                    ],
                    "coverage_note": "Document evidence only; Jev remains the match authority.",
                    "owner_questions": [],
                }
            )
        if stage == "rewriter":
            claim_id = context["approved_claims"][0]["id"]
            preferred_profile_claim_ids = context.get(
                "preferred_technical_profile_claim_ids", []
            )
            alternate_profile_claim_id = next(
                (
                    evidence["id"]
                    for profile in context.get("technical_profile_sources", [])
                    for evidence in profile.get("evidence", [])
                    if evidence["id"] != claim_id
                ),
                claim_id,
            )
            requirement_ids = [
                item["requirement_id"] for item in context["recruiter_evidence_map"]
            ]
            lines = context["cv"]["lines"]
            contact_context = context["research"]["contacts"]
            assert isinstance(contact_context, list)
            assert '"public_email":' not in json.dumps(contact_context)
            assert "alex@example.org" not in json.dumps(contact_context)
            assert "recruitment@example.org" not in json.dumps(context["job"])
            assert "applications@example.org" not in json.dumps(context["job"])
            if self.inspect_application_page:
                assert "recruitment@example.org" not in json.dumps(context["research"])
            assert context["owner_writing_context"][0]["source_type"] == "descriptive_profile"
            assert "synthetic-writing-sample" not in json.dumps(context)
            order = [line["id"] for line in lines]
            if context["cv"]["structure_policy"] == "allow_improvements":
                order = ["L0003", "L0004", "L0001", "L0002"]
            available_profile_ids = [
                evidence["id"]
                for profile in context.get("technical_profile_sources", [])
                for evidence in profile.get("evidence", [])
            ]
            mapped_profile_ids = list(dict.fromkeys([
                *preferred_profile_claim_ids,
                *available_profile_ids,
                claim_id,
                alternate_profile_claim_id,
            ]))[:2]
            letter = _complete_quality_letter(
                preferred_profile_claim_ids=mapped_profile_ids,
                role_title=context["job"]["title"],
                company=context["job"]["company"],
                responsibilities=("Python service work", "reliable local workflow testing"),
            )["cover_letter_paragraphs"]
            if self.short_letter_mode:
                letter = [
                    {
                        "section": "evidence",
                        "text": (
                            "I worked on Ayni, a platform for creating and operating specialized AI-agent workflows. "
                            "Agent-assisted workflows are an explicit part of this role."
                        ),
                        "claim_ids": [mapped_profile_ids[0]],
                        "cv_line_ids": [],
                        "requirement_ids": [requirement_ids[0]],
                        "research_finding_ids": [],
                        "source_urls": [JOB_URL],
                        "preference_source_ids": [],
                    },
                    {
                        "section": "evidence",
                        "text": (
                            "For Medical Image Registration, I deployed a VoxelMorph pipeline on the AMD Kria KV260 DPU. "
                            "Model deployment is also an explicit responsibility."
                        ),
                        "claim_ids": [mapped_profile_ids[1]],
                        "cv_line_ids": [],
                        "requirement_ids": [requirement_ids[1]],
                        "research_finding_ids": [],
                        "source_urls": [JOB_URL],
                        "preference_source_ids": [],
                    },
                ]
            letter[0]["requirement_ids"] = [requirement_ids[0]]
            for index, paragraph in enumerate(letter):
                paragraph.update(
                    {
                        "research_finding_ids": ["RF01"] if index == 2 else [],
                        "preference_source_ids": [],
                    }
                )
                if index == 1:
                    paragraph["cv_line_ids"] = ["L0004"]
                    paragraph["requirement_ids"] = [requirement_ids[0]]
                elif index == 2:
                    paragraph["cv_line_ids"] = []
                    paragraph["requirement_ids"] = [requirement_ids[1]]
            if self.short_letter_mode:
                letter[1]["requirement_ids"] = [requirement_ids[1]]
            if self.omit_preferred_profile_claims:
                letter[1]["claim_ids"] = [claim_id]
                letter[2]["claim_ids"] = [claim_id]
            application_questions = context.get("application_questions", [])
            application_answers = [
                {
                    "question": question,
                    "answer": (
                        "I have implemented reliable Python services and tested a simulated "
                        "maintenance workflow with synthetic reports to measure duplicate handling."
                    ),
                    "claim_ids": [claim_id],
                    "cv_line_ids": ["L0004"],
                    "source_urls": [APPLICATION_URL],
                    "needs_owner_input": False,
                }
                for question in application_questions
            ]
            application_email = context["research"].get("application_email_instruction", {})
            if application_email.get("required"):
                outreach_drafts = [
                    {
                        "contact_source_url": application_email["source_url"],
                        "recipient_name": "Recruitment team",
                        "subject": "Application for Software Engineer",
                        "body": (
                            "Hello Recruitment team, I am applying for the Software Engineer role. "
                            "My experience includes implementing reliable Python services and testing "
                            "a simulated maintenance workflow with synthetic reports. I will attach "
                            "my tailored CV and cover letter when sending this message. Thank you "
                            "for considering my application; I would welcome a conversation about "
                            "the role."
                        ),
                        "claim_ids": [claim_id],
                        "cv_line_ids": ["L0004"],
                        "source_urls": [JOB_URL, application_email["source_url"]],
                    }
                ]
            else:
                outreach_drafts = ([
                    {
                        "contact_source_url": JOB_URL,
                        "recipient_name": "Alex Rivera",
                        "subject": "Question about the Software Engineer role",
                        "body": (
                            "Hello Alex, I am interested in the Software Engineer role and the team's work on Python services. "
                            "My experience includes building a Python service for Example Labs, where I focused on reliable behavior and clear interfaces. "
                            "I would be grateful for any guidance on the priorities for this opening and how the engineering team approaches service evaluation. "
                            "I understand you may be busy, so even a brief pointer to the most relevant team information would be helpful. "
                            "Thank you for your time and consideration."
                        ),
                        "claim_ids": [claim_id],
                        "cv_line_ids": ["L0004"],
                        "source_urls": [JOB_URL],
                    }
                ] if contact_context else [])
            return _message_result(
                {
                    "selected_cv_id": context["cv"]["source_id"],
                    "resume_line_order": order,
                    "resume_bullet_edits": [
                        {
                            "line_id": "L0004",
                            "original_text": lines[3]["text"],
                            "revised_text": "Built reliable Python services for Example Labs.",
                            "claim_ids": [claim_id],
                            "cv_line_ids": ["L0004"],
                            "edit_reason": "Clarifies the supported contribution without inventing a metric.",
                        }
                    ],
                    "cover_letter_title": "Application for Software Engineer",
                    "cover_letter_paragraphs": letter,
                    "application_answers": application_answers,
                    "outreach_drafts": outreach_drafts,
                    "unresolved_questions": [],
                    "change_summary": "Reordered existing sections and clarified one supported bullet.",
                }
            )
        if stage == "letter_reviewer":
            context = json.loads(prompt_input[0]["content"])
            self.letter_review_calls += 1
            needs_revision = (
                (self.short_letter_mode or self.force_letter_review)
                and self.letter_review_calls == 1
            )
            profile_assessment = []
            missing_profile_claim = None
            for item in context.get("selected_profile_evidence_for_review", []):
                paragraph_indices = item.get("body_paragraph_indices", [])
                if (
                    self.force_profile_coverage_gap
                    and self.letter_review_calls == 1
                    and not paragraph_indices
                    and missing_profile_claim is None
                ):
                    missing_profile_claim = item["claim_id"]
                    profile_assessment.append({
                        "claim_id": item["claim_id"],
                        "status": "material_gap",
                        "paragraph_index": None,
                        "rationale": "This separate selected profile claim supports another important role responsibility.",
                    })
                else:
                    profile_assessment.append({
                        "claim_id": item["claim_id"],
                        "status": "included" if paragraph_indices else "not_needed",
                        "paragraph_index": paragraph_indices[0] if paragraph_indices else None,
                        "rationale": "The claim is included in a developed paragraph." if paragraph_indices else "This claim would not add a distinct role example.",
                    })
            return _message_result(
                {
                    "decision": "needs_revision" if needs_revision else "ready",
                    "dimensions": [
                        {
                            "name": name,
                            "score": 2 if needs_revision and name in {
                                "letter_completeness", "candidate_evidence_development"
                            } else 4,
                            "reason": "The synthetic reviewer requests one quality revision." if needs_revision else "The synthetic sample meets this criterion.",
                        }
                        for name in LETTER_REVIEW_DIMENSION_NAMES
                    ],
                    "issues": ([
                        {
                            "issue": "The synthetic sample needs a stronger application argument.",
                            "paragraph_indices": list(range(len(context["letter_paragraphs"]))),
                        }
                    ] if needs_revision else []),
                    "profile_evidence_assessment": profile_assessment,
                }
            )
        if stage == "letter_reviser":
            context = json.loads(prompt_input[0]["content"])
            preferred_profile_claim_ids = context.get(
                "preferred_technical_profile_claim_ids", []
            )
            available_profile_ids = [
                evidence["id"]
                for profile in context.get("technical_profile_sources", [])
                for evidence in profile.get("evidence", [])
            ]
            mapped_profile_ids = list(dict.fromkeys([
                *preferred_profile_claim_ids,
                *available_profile_ids,
            ]))[:2]
            letter = _complete_quality_letter(
                preferred_profile_claim_ids=mapped_profile_ids,
                role_title=context["job"]["title"],
                company=context["job"]["company"],
                responsibilities=tuple(
                    item.get("requirement", "")
                    for item in context["recruiter_evidence_map"][:2]
                ),
            )["cover_letter_paragraphs"]
            revisions = []
            evidence_templates = [item for item in letter if item["section"] == "evidence"]
            all_paragraphs = [
                *context["locked_supported_paragraphs"],
                *context["paragraphs_for_repair"],
            ]
            for item in context["paragraphs_for_repair"]:
                assert "claim_ids" in item and "cv_line_ids" in item
                paragraph_index = item["paragraph_index"]
                section = item.get("section", "evidence")
                if section == "evidence" and evidence_templates:
                    evidence_position = sum(
                        1
                        for prior in all_paragraphs
                        if prior.get("section", "evidence") == "evidence"
                        and prior.get("paragraph_index", 0) < paragraph_index
                    )
                    paragraph = dict(evidence_templates[min(evidence_position, len(evidence_templates) - 1)])
                else:
                    paragraph = dict(next(
                        (candidate for candidate in letter if candidate["section"] == section),
                        letter[0],
                    ))
                allowed_claims = item.get("allowed_claim_ids", [])
                claim_ids = [value for value in paragraph.get("claim_ids", []) if value in allowed_claims]
                if not claim_ids:
                    claim_ids = [value for value in item.get("claim_ids", []) if value in allowed_claims]
                allowed_lines = item.get("allowed_cv_line_ids", [])
                cv_line_ids = [value for value in item.get("cv_line_ids", []) if value in allowed_lines]
                paragraph.update(
                    {
                        "paragraph_index": paragraph_index,
                        "section": section,
                        "claim_ids": claim_ids,
                        "cv_line_ids": cv_line_ids,
                        "requirement_ids": item.get("requirement_ids") or paragraph.get("requirement_ids", []),
                        "research_finding_ids": [],
                        "source_urls": [JOB_URL],
                        "preference_source_ids": [],
                    }
                )
                revisions.append(paragraph)
            return _message_result({"revisions": revisions, "unresolved_questions": []})
        if stage == "hiring_manager":
            questions = context.get("questions") or [
                "How would you measure reliability in a Python service?",
                "Describe a difficult technical tradeoff you made.",
                "How would you investigate a production regression?",
            ]
            answers = context.get("answers") or []
            return _message_result(
                {
                    "questions": questions,
                    "answer_assessments": [
                        {
                            "question": answer["question"],
                            "answer": answer["answer"],
                            "technical_evidence": "The answer names a concrete implementation detail.",
                            "reasoning": "The reasoning follows a clear sequence.",
                            "clarity": "The explanation is concise.",
                            "suggested_practice": "Add the outcome and what you learned.",
                        }
                        for answer in answers
                    ],
                    "practice_note": "Practice feedback only; this is not a hiring prediction.",
                }
            )
        raise AssertionError(f"Unexpected synthetic workflow stage: {stage}")


def test_research_and_writing_only_run_after_click_and_bind_all_selected_inputs(settings, database):
    record, cv, _technical, _descriptive, sample, claim_id, _resume_text = _create_preparation(
        settings, database, technical_authorship="owner_written"
    )
    source, claims, preferences, writing, technical_profiles = _get_bound_inputs(database, settings, record)
    assert source["id"] == cv["id"]
    assert [claim["id"] for claim in claims] == [claim_id]
    assert claims[0]["source_type"] == "technical_profile"
    assert preferences == ""
    assert [item["source_type"] for item in writing] == ["descriptive_profile"]
    assert writing[0]["filename"] == "synthetic-descriptive.md"
    assert len(technical_profiles) == 1
    assert technical_profiles[0]["source_id"] == _technical["id"]
    assert technical_profiles[0]["authorship_label"] == "owner_written"
    assert technical_profiles[0]["text"] == _technical["extracted_text"]
    assert sample["permitted"] == 0

    settings_without_key = settings

    def must_not_create_client(_key):
        raise AssertionError("OpenAI client must not be created while the key is blank.")

    run_preparation(database, settings_without_key, record["id"], client_factory=must_not_create_client)
    blocked = get_preparation(database, record["id"])
    assert blocked["state"] == "blocked"
    with connect(database) as db:
        assert db.execute("SELECT COUNT(*) FROM openai_usage").fetchone()[0] == 0

    set_source_options(
        database,
        sample["id"],
        permitted=True,
        default_cv=False,
        structure_policy="preserve",
    )
    with pytest.raises(StalePreparationError, match="writing references changed"):
        _get_bound_inputs(database, settings, record)
    assert get_settings(database)["openai_consent_at"] == ""


def test_technical_profile_claim_changes_stale_the_listing_trigger(settings, database):
    record, _cv, technical, _descriptive, _sample, _claim_id, _resume = _create_preparation(
        settings,
        database,
        technical_text=(
            "Implemented a typed local workflow.\n"
            "Measured regression evaluation across three synthetic runs.\n"
        ),
    )
    pending = next(
        item for item in list_claims(database, technical["id"])
        if item["status"] == "unreviewed"
    )
    review_claim(
        database,
        pending["id"],
        status="rejected",
        evidence_level="unknown",
        category="systems",
        role_family="applied_ai_llm",
        owner_note="Synthetic test rejection.",
    )

    with pytest.raises(StalePreparationError, match="Technical profile evidence changed"):
        _get_bound_inputs(database, settings, record)


def test_owner_trigger_routes_full_descriptive_context_and_approved_claims_to_rewriter_only(settings, database):
    profile_tail = "PROFILE_TAIL_MARKER: use concise, direct language."
    descriptive_text = (
        "Descriptive profile opening.\n"
        + ("A long locally stored profile paragraph with general personal detail.\n" * 250)
        + profile_tail
    )
    claim_marker = "APPROVED_TECHNICAL_CLAIM_MARKER"
    technical_tail = "UNREVIEWED_PROFILE_MARKER: implemented reliable Python service testing workflows."
    record, _cv, _technical, _descriptive, _sample, _claim_id, _resume = _create_preparation(
        settings,
        database,
        technical_text=(
            f"Implemented a typed local workflow. {claim_marker}\n"
            "Demonstrated reliable local workflow testing using synthetic maintenance reports.\n"
            f"{technical_tail}\n"
        ),
        descriptive_text=descriptive_text,
    )
    enabled = _enable_synthetic_openai(settings, database)
    FakePreparationClient.instances.clear()

    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )

    completed = get_preparation(database, record["id"])
    assert completed["state"] == "review", completed
    calls = FakePreparationClient.instances[-1].calls
    stage_inputs = {}
    profile_prompt_contexts = {}
    rewriter_context = None
    for payload in calls:
        stage = payload["text"]["format"]["name"].removeprefix("clue_").rsplit("_v", 1)[0]
        stage_inputs.setdefault(stage, []).append(json.dumps(payload["input"], ensure_ascii=False))
        if stage in {"recruiter", "rewriter"}:
            parsed_context = json.loads(payload["input"][0]["content"])
            profile_prompt_contexts[stage] = parsed_context["technical_profile_sources"]
            if stage == "rewriter":
                rewriter_context = parsed_context

    assert {"researcher", "diagnoser", "recruiter", "rewriter"} <= set(stage_inputs)
    for isolated_stage in ("researcher", "diagnoser"):
        assert all(profile_tail not in value for value in stage_inputs[isolated_stage])
        assert all(claim_marker not in value for value in stage_inputs[isolated_stage])
        assert all(technical_tail not in value for value in stage_inputs[isolated_stage])
    assert all(profile_tail not in value for value in stage_inputs["recruiter"])
    assert all(claim_marker in value for value in stage_inputs["recruiter"])
    assert all(technical_tail in value for value in stage_inputs["recruiter"])
    assert all(profile_tail in value for value in stage_inputs["rewriter"])
    assert all(claim_marker in value for value in stage_inputs["rewriter"])
    assert all(technical_tail in value for value in stage_inputs["rewriter"])
    assert rewriter_context is not None
    ranked_profile_ids = [
        item["id"]
        for source in rewriter_context["technical_profile_sources"]
        for item in source["evidence"]
    ][:2]
    assert len(ranked_profile_ids) == 2
    assert rewriter_context["preferred_technical_profile_claim_ids"] == ranked_profile_ids
    for sources in profile_prompt_contexts.values():
        evidence = [claim for source in sources for claim in source["evidence"]]
        assert evidence
        assert all(
            set(claim) == {
                "id", "evidence_excerpt", "source_excerpts", "related_cv_lines", "review_status"
            }
            for claim in evidence
        )
        assert any(claim["evidence_excerpt"] == technical_tail for claim in evidence)


def test_unreviewed_profile_excerpt_can_be_used_only_with_jev_source_evidence():
    source_id = "technical-profile-1"
    evidence_id = "profile-claim-1"
    excerpt = "Built AgentDock with immutable run versions and allowlisted tool calls."
    profile_sources = [
        {
            "source_id": source_id,
            "filename": "technical-profile.md",
            "text": f"Project notes\n{excerpt}\n",
            "evidence": [
                {
                    "id": evidence_id,
                    "claim_text": excerpt,
                    "evidence_excerpt": excerpt,
                    "review_status": "unreviewed",
                }
            ],
        }
    ]
    rewrite = {
        "selected_cv_id": "cv-synthetic",
        "resume_line_order": ["L0001"],
        "resume_bullet_edits": [
            {
                "line_id": "L0001",
                "original_text": "Project experience",
                "revised_text": "Built AgentDock with immutable run versions and allowlisted tool calls.",
                "claim_ids": [evidence_id],
                "cv_line_ids": ["L0001"],
            }
        ],
        "cover_letter_paragraphs": [],
        "application_answers": [],
        "outreach_drafts": [],
    }
    _validate_rewrite(
        rewrite,
        "cv-synthetic",
        {"L0001": "Project experience"},
        set(),
        {"findings": [], "contacts": []},
        [],
        "preserve",
        permitted_profile_claim_ids={evidence_id},
    )
    assertions = _grounding_assertions(
        rewrite,
        [],
        {"findings": [], "contacts": []},
        technical_profile_sources=profile_sources,
        cv_lines={"L0001": "Project experience"},
        cv_source_id="cv-synthetic",
    )
    assert assertions[0]["evidence"] == [
        {
            "source_type": "owner_provided_resume",
            "source_id": "cv-synthetic",
            "text": "Selected CV line L0001: Project experience",
        },
        {
            "source_type": "owner_provided_technical_profile",
            "source_id": source_id,
            "text": f"Technical profile excerpt (reference {evidence_id}; review status: unreviewed): {excerpt}",
        }
    ]

    unsupported = {
        **rewrite,
        "resume_bullet_edits": [{
            **rewrite["resume_bullet_edits"][0],
            "claim_ids": ["not-bound-to-this-profile"],
        }],
        "cover_letter_paragraphs": [{
            "text": "This paragraph cites an unapproved claim.",
            "claim_ids": ["not-bound-to-this-profile"],
            "source_urls": [],
            "research_finding_ids": [],
            "preference_source_ids": [],
        }],
    }
    _validate_rewrite(
        unsupported,
        "cv-synthetic",
        {"L0001": "Project experience"},
        set(),
        {"findings": [], "contacts": []},
        [],
        "preserve",
        permitted_profile_claim_ids={evidence_id},
    )
    assert unsupported["resume_bullet_edits"] == []
    assert unsupported["cover_letter_paragraphs"] == []
    assert unsupported["unresolved_questions"] == [
        "Clue removed 2 generated block(s) that cited evidence outside this request's approved set."
    ]


def test_mapped_technical_profile_passages_reach_rewriter_and_jev():
    source_id = "technical-profile-details"
    agent_claim_id = "profile-claim-agent-platform"
    vision_claim_id = "profile-claim-medical-registration"
    source_text = (
        "Ayni is a platform for creating and operating specialized AI-agent workflows.\n\n"
        "Current architecture includes:\n"
        "- structured output validation\n"
        "- bounded tool invocation\n"
        "- approval gates and evaluation suites\n\n"
        "The work evolves from a TensorFlow VoxelMorph pipeline deployed on the AMD Kria "
        "KV260 DPU for deformable MR-to-CT registration.\n"
    )
    source = {
        "source_id": source_id,
        "filename": "technical-profile.md",
        "text": source_text,
        "use": "Owner-provided technical source; Jev checks cited statements against source text.",
        "evidence": [
            {
                "id": agent_claim_id,
                "claim_text": "Agentic AI systems include Ayni and bounded agent architectures.",
                "evidence_excerpt": "Ayni, retrieval, tool execution, approval, and evaluation.",
                "review_status": "unreviewed",
            },
            {
                "id": vision_claim_id,
                "claim_text": "Medical image registration and FPGA/DPU deployment experience.",
                "evidence_excerpt": "Medical image registration, AMD Vitis AI, FPGA/DPU deployment.",
                "review_status": "unreviewed",
            },
            {
                "id": "unmapped-claim",
                "claim_text": "A role-unrelated project detail.",
                "evidence_excerpt": "Unrelated detail.",
                "review_status": "unreviewed",
            },
        ],
    }
    cv_lines = {
        "L0001": "Built an agent platform with model routing, tool execution, retrieval memory, and CRM API integration.",
        "L0002": "Trained and deployed a quantized medical-image registration model on FPGA hardware.",
    }
    requirements = [
        {
            "requirement_id": "REQ-01",
            "criterion": "Build agent-assisted workflows.",
            "claim_ids": [agent_claim_id],
            "cv_line_ids": ["L0001"],
        },
        {
            "requirement_id": "REQ-02",
            "criterion": "Deploy AI services in biomedical contexts.",
            "claim_ids": [vision_claim_id],
            "cv_line_ids": ["L0002"],
        },
    ]

    selected = _select_technical_profile_evidence(
        [source],
        {agent_claim_id, vision_claim_id},
        recruiter_requirements=requirements,
        cv_lines=cv_lines,
    )
    prompt_context = _technical_profile_prompt_context(selected)
    serialized_context = json.dumps(prompt_context, ensure_ascii=False)
    prompt_claims = prompt_context[0]["evidence"]
    agent_prompt_claim = next(item for item in prompt_claims if item["id"] == agent_claim_id)
    vision_prompt_claim = next(item for item in prompt_claims if item["id"] == vision_claim_id)
    assert any("Ayni is a platform" in item for item in agent_prompt_claim["source_excerpts"])
    assert any("bounded tool invocation" in item for item in agent_prompt_claim["source_excerpts"])
    assert any("TensorFlow VoxelMorph pipeline" in item for item in vision_prompt_claim["source_excerpts"])
    assert any("AMD Kria KV260 DPU" in item for item in vision_prompt_claim["source_excerpts"])
    assert "unmapped-claim" not in serialized_context
    assert source_text not in serialized_context

    rewrite = {
        "cover_letter_paragraphs": [
            {
                "text": "I built the Ayni platform for specialized agent workflows.",
                "claim_ids": [agent_claim_id],
                "cv_line_ids": ["L0001"],
                "source_urls": [JOB_URL],
                "research_finding_ids": [],
                "preference_source_ids": [],
            },
            {
                "text": "I deployed a TensorFlow VoxelMorph pipeline on the AMD Kria KV260 DPU.",
                "claim_ids": [vision_claim_id],
                "cv_line_ids": ["L0002"],
                "source_urls": [JOB_URL],
                "research_finding_ids": [],
                "preference_source_ids": [],
            },
        ],
    }
    assertions = _grounding_assertions(
        rewrite,
        [],
        {"findings": [], "contacts": []},
        job_context={"title": "AI Engineer", "company": "IFOM", "description": "agent-assisted workflows and AI-service deployment", "canonical_url": JOB_URL},
        technical_profile_sources=selected,
        cv_lines=cv_lines,
        cv_source_id="selected-cv",
    )
    assert any(
        "Ayni is a platform" in evidence["text"]
        for evidence in assertions[0]["evidence"]
    )
    assert any(
        "TensorFlow VoxelMorph pipeline" in evidence["text"]
        for evidence in assertions[1]["evidence"]
    )


def test_long_profile_evidence_uses_exact_excerpts_without_duplicate_source_text(settings, database):
    technical_text = (
        "TECHNICAL_PROFILE_START\n"
        + ("Project detail with implementation evidence.\n" * 1_100)
        + "Separately evaluated a local retrieval workflow using synthetic fixtures.\n"
    )
    descriptive_text = "DESCRIPTIVE_PROFILE_START\n" + ("Direct, modest writing preference.\n" * 1_000)
    record, _cv, _technical, _descriptive, *_ = _create_preparation(
        settings,
        database,
        technical_text=technical_text,
        descriptive_text=descriptive_text,
    )
    enabled = _enable_synthetic_openai(settings, database)
    FakePreparationClient.instances.clear()
    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )

    completed = get_preparation(database, record["id"])
    assert completed["state"] == "review"
    calls = FakePreparationClient.instances[-1].calls
    recruiter = next(
        payload for payload in calls
        if payload["text"]["format"]["name"] == "clue_recruiter_v3"
    )
    rewriter = next(
        payload for payload in calls
        if payload["text"]["format"]["name"] == "clue_rewriter_v14"
    )
    recruiter_context = json.loads(recruiter["input"][0]["content"])
    rewriter_context = json.loads(rewriter["input"][0]["content"])
    recruiter_source = recruiter_context["technical_profile_sources"][0]
    assert "text" not in recruiter_source
    mapped_technical_evidence = [
        evidence
        for profile in rewriter_context.get("technical_profile_sources", [])
        for evidence in profile.get("evidence", [])
    ]
    assert len(mapped_technical_evidence) == 1
    assert mapped_technical_evidence[0]["evidence_excerpt"] == (
        "Separately evaluated a local retrieval workflow using synthetic fixtures."
    )
    assert technical_text not in json.dumps(rewriter_context, ensure_ascii=False)
    assert rewriter_context["owner_writing_context"][0]["text"] == descriptive_text
    with connect(database) as db:
        assert db.execute("SELECT COUNT(*) FROM openai_usage WHERE status = 'settled'").fetchone()[0] == 6


def test_fake_full_packet_review_gmail_draft_and_no_automatic_application(settings, database):
    record, cv, _technical, _descriptive, _sample, _claim_id, original_resume = _create_preparation(
        settings, database, structure_policy="allow_improvements"
    )
    enabled = _enable_synthetic_openai(settings, database)
    FakePreparationClient.instances.clear()
    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )

    completed = get_preparation(database, record["id"])
    assert completed["state"] == "review", completed.get("status_message")
    with connect(database) as db:
        assert db.execute("SELECT COUNT(*) FROM applications").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM application_events").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM jev_usage").fetchone()[0] == 0
        usages = db.execute("SELECT stage, status FROM openai_usage ORDER BY rowid").fetchall()
    assert len(usages) == 6  # two research calls, three writer roles, and the letter-quality review
    assert all(row["status"] == "settled" for row in usages)
    assert [row["stage"] for row in usages].count("researcher") == 2

    with connect(database) as db:
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]
    packet = get_packet(database, packet_id)
    assert packet["output"]["jev_tailored_resume_review"]["status"] == "approved"
    assert packet["output"]["quality_check"]["status"] == "passed"
    assert packet["output"]["quality_check"]["body_paragraph_count"] == 5
    assert packet["output"]["quality_check"]["owner_review_required"] is True
    assert "paragraph arrangement are flexible" in packet["output"]["review_note"]
    assert packet["output"]["model"] == MODEL_ID
    assert packet["output"]["reasoning_effort"] == "high"
    assert packet["output"]["prompt_version"] == PROMPT_VERSION
    assert packet["output"]["output_schema_version"] == OUTPUT_SCHEMA_VERSION
    assert packet["output"]["research"]["contacts"][0]["public_email"] == "alex@example.org"
    assert Path(cv["file_path"]).read_text(encoding="utf-8") == original_resume
    resume_artifact = next(item for item in packet["artifacts"] if item["artifact_type"] == "resume")
    resume_doc = Document(resume_artifact["file_path"])
    paragraphs = [paragraph.text for paragraph in resume_doc.paragraphs]
    assert paragraphs[:2] == ["Experience", "Built reliable Python services for Example Labs."]
    assert paragraphs.index("Education") > paragraphs.index("Experience")
    assert resume_doc.styles["Normal"].font.name == "Arial"
    assert resume_doc.styles["Normal"].font.size.pt == 9
    experience_heading = next(p for p in resume_doc.paragraphs if p.text == "Experience")
    assert experience_heading.style.name == "Heading 1"
    assert experience_heading.paragraph_format.keep_with_next is True
    assert resume_doc.paragraphs[1].paragraph_format.keep_together is True
    bullet_paragraphs = [
        paragraph for paragraph in resume_doc.paragraphs if paragraph.style.name == "List Bullet"
    ]
    assert bullet_paragraphs
    assert all(
        paragraph.paragraph_format.left_indent.inches == pytest.approx(0.28, abs=0.001)
        and paragraph.paragraph_format.first_line_indent.inches == pytest.approx(-0.18, abs=0.001)
        for paragraph in bullet_paragraphs
    )
    cover_letter_artifact = next(
        item for item in packet["artifacts"] if item["artifact_type"] == "cover_letter"
    )
    cover_letter_doc = Document(cover_letter_artifact["file_path"])
    cover_letter_paragraphs = [paragraph.text for paragraph in cover_letter_doc.paragraphs]
    assert re.fullmatch(r"\d{1,2} \w+ \d{4}", cover_letter_paragraphs[0])
    assert cover_letter_paragraphs[1] == "Application for Software Engineer at Example Labs"
    assert cover_letter_paragraphs[2] == "Dear Hiring Team,"
    assert cover_letter_paragraphs[-1] == "Kind regards,"
    cover_letter_title_style = cover_letter_doc.styles["Clue Cover Letter Title"]
    assert cover_letter_doc.paragraphs[1].style.name == "Clue Cover Letter Title"
    assert cover_letter_title_style.base_style.name == "Normal"
    assert b"w:pBdr" not in cover_letter_title_style._element.xml.encode()
    assert cover_letter_title_style.font.color.rgb == (0, 0, 0)
    assert cover_letter_title_style.font.underline is False
    assert "Python service work" in cover_letter_paragraphs[3]
    assert "Ayni" in cover_letter_paragraphs[4]
    assert "VoxelMorph" in cover_letter_paragraphs[5]
    assert cover_letter_artifact["content_text"].startswith(cover_letter_paragraphs[0])
    assert cover_letter_artifact["content_text"].endswith("Kind regards,")
    outreach_artifact = next(
        item for item in packet["artifacts"] if item["artifact_type"] == "outreach_draft"
    )
    outreach_text = Path(outreach_artifact["file_path"]).read_text(encoding="utf-8")
    assert "To: Alex Rivera <alex@example.org>" in outreach_text
    assert "Subject: Question about the Software Engineer role" in outreach_text
    assert "Source: " + JOB_URL in outreach_text
    assert all(
        Path(item["file_path"]).resolve().is_relative_to(settings.data_dir.resolve())
        for item in packet["artifacts"]
    )
    client = TestClient(create_app(enabled), base_url="http://127.0.0.1")
    missing_confirmation = client.post(
        f"/preparations/{record['id']}/packets/{packet_id}/approve",
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert "Confirm+that+you+reviewed" in missing_confirmation.headers["location"]
    assert get_packet(database, packet_id)["status"] == "review"
    approved = client.post(
        f"/preparations/{record['id']}/packets/{packet_id}/approve",
        data={"confirm_review": "on"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert approved.status_code == 303
    assert get_packet(database, packet_id)["status"] == "approved"
    with connect(database) as db:
         assert db.execute("SELECT COUNT(*) FROM applications").fetchone()[0] == 0

    with pytest.raises(GmailDraftError, match="Confirm the exact recipient"):
        create_approved_packet_draft(
            database,
            enabled,
            record["id"],
            packet_id,
            0,
            confirmed=False,
            create_draft=lambda *_args: pytest.fail("No Gmail call without confirmation"),
        )
    calls = []
    row_id, gmail_id = create_approved_packet_draft(
        database,
        enabled,
        record["id"],
        packet_id,
        0,
        confirmed=True,
        create_draft=lambda *args: calls.append(args) or "synthetic-gmail-draft-id",
    )
    assert gmail_id == "synthetic-gmail-draft-id"
    assert calls[0][1:] == (
        packet["output"]["rewriter"]["outreach_drafts"][0]["recipient_email"],
        packet["output"]["rewriter"]["outreach_drafts"][0]["subject"],
        packet["output"]["rewriter"]["outreach_drafts"][0]["body"],
    )
    with connect(database) as db:
        draft = db.execute("SELECT state FROM gmail_drafts WHERE id = ?", (row_id,)).fetchone()
    assert draft["state"] == "created"
    with pytest.raises(GmailDraftError, match="already created or unresolved"):
        create_approved_packet_draft(
            database,
            enabled,
            record["id"],
            packet_id,
            0,
            confirmed=True,
            create_draft=lambda *_args: pytest.fail("Duplicate Gmail call is blocked"),
        )
    with Path(resume_artifact["file_path"]).open("ab") as output:
        output.write(b"tampered")
    with pytest.raises(ValueError, match="changed after approval"):
        verify_packet_approval(database, enabled, packet_id)


def test_resume_signature_name_uses_only_a_plain_cv_header_name():
    assert _resume_signature_name([{"text": "JUAN MARTIN SANCHEZ BARDELLINI"}]) == "JUAN MARTIN SANCHEZ BARDELLINI"
    assert _resume_signature_name([{"text": "JUAN MARTIN | AI/ML Engineer"}]) == ""
    assert _resume_signature_name([{"text": "Curriculum Vitae"}]) == ""


def test_review_snapshot_can_finish_with_jev_approval_without_mutating_match_status(settings, database):
    record, *_ = _create_preparation(
        settings,
        database,
        filter_status="review",
        additional_source_url=OFFICIAL_SOURCE_URL,
    )
    enabled = _enable_synthetic_openai(settings, database)
    FakePreparationClient.instances.clear()
    FakeCrawler.instances.clear()

    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )

    completed = get_preparation(database, record["id"])
    with connect(database) as db:
        saved_status = db.execute(
            "SELECT filter_status FROM search_results WHERE run_id = ? AND job_id = ?",
            ("run-workflow", record["job_id"]),
        ).fetchone()["filter_status"]
        application_count = db.execute("SELECT COUNT(*) FROM applications").fetchone()[0]
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]
    packet = get_packet(database, packet_id)

    assert completed["state"] == "review"
    assert completed["snapshot"]["filter_status"] == "review"
    assert saved_status == "review"
    assert packet["output"]["jev_tailored_resume_review"]["status"] == "approved"
    assert packet["output"]["quality_check"]["status"] == "passed"
    assert application_count == 0
    assert OFFICIAL_SOURCE_URL in FakeCrawler.instances[-1].allowed_urls


def test_owner_acknowledged_conflict_and_location_override_can_finish_all_document_gates(
    settings, database
):
    record, *_ = _create_preparation(
        settings,
        database,
        filter_status="conflict",
        eligibility_status="needs_verification",
        confirm_jev_conflict=True,
        confirm_location_eligibility=True,
    )
    enabled = _enable_synthetic_openai(settings, database)
    FakePreparationClient.instances.clear()
    FakeCrawler.instances.clear()

    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )

    completed = get_preparation(database, record["id"])
    with connect(database) as db:
        saved = db.execute(
            "SELECT filter_status, eligibility_status FROM search_results WHERE run_id = ? AND job_id = ?",
            ("run-workflow", record["job_id"]),
        ).fetchone()
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]
    packet = get_packet(database, packet_id)

    assert completed["state"] == "review"
    assert completed["snapshot"]["filter_status"] == "conflict"
    assert completed["snapshot"]["eligibility_status"] == "needs_verification"
    assert completed["snapshot"]["owner_preparation_overrides"] == {
        "jev_conflict_acknowledged": True,
        "location_eligibility_acknowledged": True,
    }
    assert saved["filter_status"] == "conflict"
    assert saved["eligibility_status"] == "needs_verification"
    assert packet["output"]["jev_tailored_resume_review"]["status"] == "approved"
    assert packet["output"]["quality_check"]["status"] == "passed"
    assert {artifact["artifact_type"] for artifact in packet["artifacts"]} >= {
        "resume",
        "cover_letter",
    }
    dossier = TestClient(create_app(enabled), base_url="http://127.0.0.1").get(
        f"/applications/{record['id']}"
    )
    assert dossier.status_code == 200
    assert "You chose to prepare this listing with a Jev warning" in dossier.text
    assert "Saved Jev filter status: <strong>conflict</strong>" in dossier.text
    assert "Saved Jev location eligibility: <strong>needs_verification</strong>" in dossier.text


@pytest.mark.parametrize(
    (
        "description_instruction",
        "page_instruction_quote",
        "expected_quote",
        "expected_email",
        "expected_source_url",
    ),
    [
        (
            "",
            "To apply, email your CV and cover letter to recruitment@example.org.",
            "To apply, email your CV and cover letter to recruitment@example.org.",
            "recruitment@example.org",
            APPLICATION_URL,
        ),
        (
            "",
            "To apply, email your CV and cover letter to the recruitment team.",
            "To apply, email your CV and cover letter to the recruitment team.",
            "",
            APPLICATION_URL,
        ),
        (
            "To apply, email your CV and cover letter to applications@example.org.",
            "Apply through our online portal.",
            "To apply, email your CV and cover letter to applications@example.org.",
            "applications@example.org",
            JOB_URL,
        ),
    ],
)
def test_linked_application_page_generates_grounded_form_answer_and_email_packet(
    settings,
    database,
    monkeypatch,
    description_instruction,
    page_instruction_quote,
    expected_quote,
    expected_email,
    expected_source_url,
):
    role_description = "Build Python services that support reliable model inference."
    if description_instruction:
        role_description += " " + description_instruction
    record, *_ = _create_preparation(
        settings, database, description=role_description
    )
    enabled = _enable_synthetic_openai(settings, database)
    monkeypatch.setattr(FakeCrawler, "inspect_application_page", True)
    monkeypatch.setattr(FakeCrawler, "application_instruction_quote", page_instruction_quote)
    monkeypatch.setattr(FakePreparationClient, "inspect_application_page", True)
    FakePreparationClient.instances.clear()
    FakeCrawler.instances.clear()

    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )
    researcher_input = json.loads(
        FakePreparationClient.instances[-1].calls[0]["input"][0]["content"]
    )
    assert researcher_input["job_and_jev_snapshot"]["description"] == role_description
    if description_instruction:
        assert "applications@example.org" in json.dumps(researcher_input)

    completed = get_preparation(database, record["id"])
    assert completed["state"] == "review", completed.get("status_message")
    with connect(database) as db:
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]
    packet = get_packet(database, packet_id)
    research = packet["output"]["research"]
    instruction = research["application_email_instruction"]
    assert instruction == {
        "required": True,
        "quote": expected_quote,
        "source_url": expected_source_url,
    }
    assert research["application_form"]["status"] == "observed"
    assert research["application_form"]["application_pages"][0]["url"] == APPLICATION_URL
    question = "Tell us about a Python service you have built."
    assert research["application_form"]["questions"] == [question]
    answer = packet["output"]["rewriter"]["application_answers"][0]
    assert answer["question"] == question
    assert answer["source_urls"] == [APPLICATION_URL]
    assert answer["claim_ids"]
    assert packet["output"]["quality_check"]["status"] == "passed"

    email = packet["output"]["rewriter"]["outreach_drafts"][0]
    assert email["purpose"] == "application_email"
    assert email["recipient_email"] == expected_email
    assert email["contact_source_url"] == expected_source_url
    assert expected_source_url in email["source_urls"]
    artifacts = {item["artifact_type"]: item for item in packet["artifacts"]}
    assert {"resume", "cover_letter", "application_answers", "outreach_draft"} <= set(artifacts)
    assert artifacts["outreach_draft"]["filename"] == "application-email-draft-v1.txt"
    assert question in artifacts["application_answers"]["content_text"]
    expected_to = (
        f"To: Recruitment team <{expected_email}>"
        if expected_email
        else "To: Recruitment team (verify address)"
    )
    assert expected_to in artifacts["outreach_draft"]["content_text"]
    assert "tailored CV and cover letter" in artifacts["outreach_draft"]["content_text"]
    dossier = TestClient(create_app(enabled), base_url="http://127.0.0.1").get(
        f"/applications/{record['id']}"
    )
    assert dossier.status_code == 200
    assert "Apply by email" in dossier.text
    assert expected_source_url in dossier.text
    assert question in dossier.text
    assert "Use this local preview in your email app" in dossier.text
    assert "Create this exact message as an unsent Gmail draft" not in dossier.text
    if expected_email:
        assert "Application recipient address" in dossier.text
        assert expected_email in dossier.text
    else:
        assert "No verified address for this application channel is available" in dossier.text
    with connect(database) as db:
        assert db.execute("SELECT COUNT(*) FROM applications").fetchone()[0] == 0


def test_unsupported_or_unresolved_blocks_are_omitted_before_documents(settings, database):
    record, _cv, _technical, _descriptive, _sample, _claim_id, _resume = _create_preparation(
        settings, database
    )
    enabled = _enable_synthetic_openai(settings, database)

    def reject_generated_text(_database, _settings, _request_id, assertions):
        assert assertions
        assert all(item["evidence"] for item in assertions)
        return {
            "status": "complete",
            "model": "jev-test",
            "actual_cost_usd": 0.0,
            "results": [
                {
                    "id": item["id"],
                    "status": "contradicted" if item["output_type"] == "resume_bullet_edit" else "unresolved",
                    "confidence": 0.9,
                }
                for item in assertions
            ],
        }

    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        claim_support_checker=reject_generated_text,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )

    completed = get_preparation(database, record["id"])
    assert completed["state"] == "failed"
    with connect(database) as db:
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]
    packet = get_packet(database, packet_id)
    assert packet["status"] == "obsolete"
    grounding = packet["output"]["grounding"]
    assert grounding["total_assertions"] == 7
    assert grounding["supported_count"] == 0
    assert grounding["omitted_count"] == 7
    assert all(not item["included_in_artifact"] for item in grounding["results"])
    assert packet["output"]["rewriter"]["resume_bullet_edits"] == []
    assert packet["output"]["rewriter"]["cover_letter_paragraphs"] == []
    assert packet["output"]["rewriter"]["outreach_drafts"] == []

    assert packet["artifacts"] == []
    assert packet["output"]["quality_check"]["status"] == "failed"
    assert any("No cover-letter body was generated" in issue for issue in packet["output"]["quality_check"]["issues"])
    assert all(item["artifact_type"] != "outreach_draft" for item in packet["artifacts"])


def test_quality_review_triggers_one_letter_repair_and_second_review(settings, database, monkeypatch):
    record, *_ = _create_preparation(settings, database)
    enabled = _enable_synthetic_openai(settings, database)
    monkeypatch.setattr(FakePreparationClient, "short_letter_mode", True)
    support_calls = []

    def support_all(_database, _settings, _request_id, assertions):
        support_calls.append(assertions)
        return {
            "status": "complete",
            "model": "jev-test",
            "actual_cost_usd": 0.0,
            "results": [
                {"id": item["id"], "status": "supported", "confidence": 0.9}
                for item in assertions
            ],
        }

    FakePreparationClient.instances.clear()
    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        claim_support_checker=support_all,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )

    completed = get_preparation(database, record["id"])
    assert completed["state"] == "review", completed["status_message"]
    with connect(database) as db:
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]
        reviewer_usage = db.execute(
            "SELECT COUNT(*) FROM openai_usage WHERE request_id = ? AND stage = 'letter_reviewer'",
            (record["id"],),
        ).fetchone()[0]
    packet = get_packet(database, packet_id)
    quality = packet["output"]["quality_check"]
    review_calls = [
        call for call in FakePreparationClient.instances[-1].calls
        if call["text"]["format"]["name"] == "clue_letter_reviewer_v2"
    ]
    review_contexts = [json.loads(call["input"][0]["content"]) for call in review_calls]

    assert len(review_calls) == reviewer_usage == 2
    assert len(review_contexts[0]["letter_paragraphs"]) == 2
    assert len(review_contexts[1]["letter_paragraphs"]) == 2
    assert len(review_contexts[0]["selected_profile_evidence_for_review"]) == quality["preferred_profile_claim_count"]
    assert all("claim_id" in item and "claim_text" in item for item in review_contexts[0]["selected_profile_evidence_for_review"])
    assert quality["status"] == "passed", quality["issues"]
    assert quality["writing_review"]["status"] == "ready"
    assert quality["revision"]["initial_writing_review"]["status"] == "needs_revision"
    assert quality["revision"]["status"] == "completed"
    assert len(packet["output"]["rewriter"]["cover_letter_paragraphs"]) == 2
    assert len(support_calls) == 2
    assert {item["artifact_type"] for item in packet["artifacts"]} == {
        "cover_letter", "resume", "outreach_draft"
    }


def test_jev_rejected_letter_paragraph_gets_one_targeted_repair_and_full_recheck(settings, database):
    record, _cv, _technical, _descriptive, _sample, _claim_id, _resume = _create_preparation(
        settings, database
    )
    enabled = _enable_synthetic_openai(settings, database)
    support_calls = 0

    def repair_then_support(_database, _settings, _request_id, assertions):
        nonlocal support_calls
        support_calls += 1
        if support_calls == 1:
            return {
                "status": "complete",
                "model": "jev-test",
                "actual_cost_usd": 0.0,
                "results": [
                    {
                        "id": item["id"],
                        "status": "unresolved"
                        if item["id"] == "cover_letter_paragraph_0"
                        else "supported",
                        "confidence": 0.9,
                    }
                    for item in assertions
                ],
            }
        assert support_calls == 2
        assert len(assertions) == 5
        assert all(item["output_type"] == "cover_letter_paragraph" for item in assertions)
        return {
            "status": "complete",
            "model": "jev-test",
            "actual_cost_usd": 0.0,
            "results": [
                {"id": item["id"], "status": "supported", "confidence": 0.9}
                for item in assertions
            ],
        }

    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        claim_support_checker=repair_then_support,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )

    completed = get_preparation(database, record["id"])
    assert completed["state"] == "review"
    with connect(database) as db:
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]
        usage = db.execute(
            "SELECT stage, status FROM openai_usage WHERE request_id = ? ORDER BY rowid",
            (record["id"],),
        ).fetchall()
    packet = get_packet(database, packet_id)
    assert packet["output"]["quality_check"]["status"] == "passed"
    revision = packet["output"]["quality_check"]["revision"]
    assert revision["attempted"] is True
    assert revision["status"] == "completed"
    assert revision["attempt_limit"] == 2
    assert revision["initial_paragraph_count"] == 5
    assert revision["initial_supported_count"] == 4
    assert revision["jev_rejected_paragraph_indices"] == [0]
    assert revision["style_revision_paragraph_indices"] == [0]
    assert revision["revised_paragraph_count"] == 1
    assert revision["final_supported_count"] == 5
    assert revision["initial_writing_review"]["status"] == "ready"
    assert revision["final_writing_review"]["status"] == "ready"
    paragraphs = packet["output"]["rewriter"]["cover_letter_paragraphs"]
    assert paragraphs[0]["section"] == "opening"
    assert paragraphs[0]["text"].startswith("I am applying for the Software Engineer role")
    assert paragraphs[0]["claim_ids"] == []
    assert paragraphs[1]["text"].startswith("I worked on Ayni")
    assert paragraphs[2]["text"].startswith("For Medical Image Registration")
    assert support_calls == 2
    assert sum(row["stage"] == "letter_reviser" for row in usage) == 1
    assert sum(row["stage"] == "letter_reviewer" for row in usage) == 2
    assert all(row["status"] == "settled" for row in usage)
    assert packet["artifacts"]
    assert not any(item.get("included_in_artifact") and item.get("status") != "supported" for item in packet["output"]["grounding"]["results"])



def test_letter_fails_closed_when_jev_cannot_support_repaired_project_paragraphs(settings, database):
    record, *_ = _create_preparation(settings, database)
    enabled = _enable_synthetic_openai(settings, database)
    support_calls = 0

    def keep_project_evidence_unresolved(_database, _settings, _request_id, assertions):
        nonlocal support_calls
        support_calls += 1
        return {
            "status": "complete",
            "model": "jev-test",
            "actual_cost_usd": 0.0,
            "results": [
                {
                    "id": item["id"],
                    "status": "unresolved"
                    if item["id"] in {"cover_letter_paragraph_1", "cover_letter_paragraph_2"}
                    else "supported",
                    "confidence": 0.9,
                }
                for item in assertions
            ],
        }

    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        claim_support_checker=keep_project_evidence_unresolved,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )

    completed = get_preparation(database, record["id"])
    assert completed["state"] == "failed"
    with connect(database) as db:
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]
    packet = get_packet(database, packet_id)
    quality = packet["output"]["quality_check"]
    assert support_calls == 2
    assert quality["status"] == "failed"
    assert quality["revision"]["final_supported_count"] == 3
    assert any("needs at least one concrete example" in issue for issue in quality["issues"])
    assert packet["artifacts"] == []


def test_jev_resume_rejection_persists_report_without_reviewable_artifacts(settings, database):
    record, *_ = _create_preparation(settings, database)
    enabled = _enable_synthetic_openai(settings, database)

    def reject_resume(*_args, **_kwargs):
        return {
            "rubric_version": "tailored-resume-fit-v1",
            "status": "revise",
            "model_status": "revise",
            "confidence": 0.88,
            "confidence_policy": "diagnostic only",
            "reason": "Jev did not approve the tailored resume; revise its role relevance.",
            "model": "jev-test",
            "actual_input_tokens": 0,
            "actual_cost_usd": 0.0,
        }

    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=reject_resume,
    )

    completed = get_preparation(database, record["id"])
    assert completed["state"] == "failed"
    with connect(database) as db:
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]
    packet = get_packet(database, packet_id)
    assert packet["status"] == "obsolete"
    assert packet["output"]["jev_tailored_resume_review"]["status"] == "revise"
    assert packet["output"]["quality_check"]["status"] == "failed"
    assert packet["artifacts"] == []


def test_packet_quality_gate_rejects_legacy_packet_without_jev_resume_approval():
    with pytest.raises(ValueError, match="Jev must explicitly approve"):
        _assert_packet_quality_gates({"output": {}, "artifacts": []})


def test_saved_job_listing_is_grounding_for_role_facts_only(settings):
    rewrite = {
        "selected_cv_id": "cv-synthetic",
        "resume_line_order": ["L0001"],
        "resume_bullet_edits": [],
        "cover_letter_paragraphs": [
            {
                "text": "The role asks for Python and PyTorch experience.",
                "claim_ids": [],
                "research_finding_ids": [],
                "source_urls": [JOB_URL],
            }
        ],
        "application_answers": [],
        "outreach_drafts": [],
    }
    research = {"findings": [], "contacts": []}
    _validate_rewrite(
        rewrite,
        "cv-synthetic",
        {"L0001": "Experience"},
        set(),
        research,
        [],
        "preserve",
        job_source_url=JOB_URL,
    )
    assertions = _grounding_assertions(
        rewrite,
        [],
        research,
        job_context={
            "canonical_url": JOB_URL,
            "title": "Software Engineer",
            "company": "Example Labs",
            "description": "Build typed Python and PyTorch services; contact jobs@example.org.",
        },
    )
    assert assertions[0]["evidence"] == [
        {
            "source_type": "saved_job_listing",
            "source_id": "selected_job",
            "source_url": JOB_URL,
            "text": (
                "Saved job listing facts:\ntitle: Software Engineer\ncompany: Example Labs\n"
                "description: Build typed Python and PyTorch services; contact [redacted email]."
            ),
        }
    ]

    unsupported = {**rewrite, "cover_letter_paragraphs": [{
        "text": "The role asks for Python and PyTorch experience.",
        "claim_ids": [],
        "research_finding_ids": [],
        "source_urls": ["https://unverified.example/role"],
    }]}
    with pytest.raises(OpenAIProviderError, match="exact research finding ID"):
        _validate_rewrite(
            unsupported,
            "cv-synthetic",
            {"L0001": "Experience"},
            set(),
            research,
            [],
            "preserve",
            job_source_url=JOB_URL,
        )


def test_resume_edit_grounding_checks_the_complete_reflowed_bullet():
    cv_lines = {
        "L0001": "EXPERIENCE",
        "L0043": "\ufffd Trained and deployed a quantized medical-image registration model on FPGA hardware. Cut inference and",
        "L0044": "image-warping time by 66% versus CPU while maintaining image-alignment quality.",
    }
    rewrite = {
        "resume_line_order": ["L0001", "L0043", "L0044"],
        "resume_bullet_edits": [{
            "line_id": "L0043",
            "original_text": cv_lines["L0043"],
            "revised_text": (
                "Trained and deployed a quantized VoxelMorph medical-image registration model "
                "on the AMD Kria KV260 DPU (FPGA). Cut inference and"
            ),
            "claim_ids": ["profile-voxel"],
            "cv_line_ids": ["L0043"],
        }],
        "cover_letter_paragraphs": [],
        "application_answers": [],
        "outreach_drafts": [],
    }
    technical_profiles = [{
        "source_id": "technical-profile",
        "evidence": [{
            "id": "profile-voxel",
            "claim_text": "Deployed a VoxelMorph registration pipeline on the AMD Kria KV260 DPU.",
            "evidence_excerpt": "I deployed a VoxelMorph pipeline on the AMD Kria KV260 DPU.",
            "source_excerpts": ["For Medical Image Registration, I deployed a VoxelMorph pipeline on the AMD Kria KV260 DPU."],
            "review_status": "unreviewed",
        }],
    }]

    assertions = _grounding_assertions(
        rewrite,
        [],
        {"findings": [], "contacts": []},
        technical_profile_sources=technical_profiles,
        cv_lines=cv_lines,
        cv_source_id="cv-synthetic",
    )
    edit_assertion = assertions[0]

    assert edit_assertion["text"] == (
        "Trained and deployed a quantized VoxelMorph medical-image registration model "
        "on the AMD Kria KV260 DPU (FPGA). Cut inference and image-warping time by 66% "
        "versus CPU while maintaining image-alignment quality."
    )
    assert [item["text"].split(":", 1)[0] for item in edit_assertion["evidence"][:2]] == [
        "Selected CV line L0043",
        "Selected CV line L0044",
    ]
    assert any(
        item["source_type"] == "owner_provided_technical_profile"
        for item in edit_assertion["evidence"]
    )


def test_cover_letter_grounding_binds_exact_research_findings_not_every_fact_on_page():
    research = {
        "findings": [
            {
                "topic": "unrelated public fact",
                "fact": "The employer operates several offices.",
                "source_url": JOB_URL,
                "quote": "Several offices",
            },
            {
                "topic": "relevant public fact",
                "fact": "The team builds Python services.",
                "source_url": JOB_URL,
                "quote": "builds Python services",
            },
        ],
        "contacts": [],
    }
    context_findings = _research_context(research)["findings"]
    assert [item["id"] for item in context_findings] == ["RF01", "RF02"]

    paragraph = {
        "text": "My documented Python service work is relevant to the team's Python services.",
        "claim_ids": ["profile-claim"],
        "research_finding_ids": ["RF02"],
        "source_urls": [JOB_URL],
        "preference_source_ids": [],
    }
    rewrite = {
        "selected_cv_id": "cv-synthetic",
        "resume_line_order": ["L0001"],
        "resume_bullet_edits": [],
        "cover_letter_paragraphs": [paragraph],
        "application_answers": [],
        "outreach_drafts": [],
    }
    _validate_rewrite(
        rewrite,
        "cv-synthetic",
        {"L0001": "Experience"},
        {"profile-claim"},
        research,
        [],
        "preserve",
        job_source_url=JOB_URL,
        permitted_profile_claim_ids={"profile-claim"},
    )
    assertions = _grounding_assertions(
        rewrite,
        [],
        research,
        job_context={"canonical_url": JOB_URL, "description": "Build Python services."},
        technical_profile_sources=[
            {
                "source_id": "technical-profile",
                "evidence": [
                    {
                        "id": "profile-claim",
                        "evidence_excerpt": "Documented Python service work.",
                        "review_status": "unreviewed",
                    }
                ],
            }
        ],
    )
    evidence = assertions[0]["evidence"]
    assert [item["source_type"] for item in evidence] == [
        "owner_provided_technical_profile",
        "verified_public_source",
        "saved_job_listing",
    ]
    assert evidence[1]["source_id"] == "RF02"
    assert "several offices" not in evidence[1]["text"].casefold()

    paragraph["research_finding_ids"] = ["RF99"]
    with pytest.raises(OpenAIProviderError, match="unknown research finding"):
        _validate_rewrite(
            rewrite,
            "cv-synthetic",
            {"L0001": "Experience"},
            {"profile-claim"},
            research,
            [],
            "preserve",
            job_source_url=JOB_URL,
            permitted_profile_claim_ids={"profile-claim"},
        )


def test_rewriter_missing_role_source_is_bound_to_selected_listing():
    rewrite = {
        "selected_cv_id": "cv-synthetic",
        "resume_line_order": ["L0001"],
        "resume_bullet_edits": [],
        "cover_letter_paragraphs": _complete_quality_letter()["cover_letter_paragraphs"],
        "application_answers": [],
        "outreach_drafts": [],
    }
    paragraph = rewrite["cover_letter_paragraphs"][1]
    paragraph["source_urls"] = []

    _validate_rewrite(
        rewrite,
        "cv-synthetic",
        {"L0001": "Experience"},
        {"claim-ayni", "claim-registration"},
        {"findings": [], "contacts": []},
        [],
        "preserve",
        job_source_url=JOB_URL,
    )

    assert paragraph["source_urls"] == [JOB_URL]
    assert paragraph["clue_source_url_binding"] == "selected_listing_fallback"

    quality = _cover_letter_quality(
        rewrite,
        _supported_letter_grounding(),
    )
    assert quality["status"] == "passed", quality["issues"]


def test_cover_letter_accepts_exact_candidate_evidence_when_recruiter_map_is_incomplete():
    rewrite = {
        "selected_cv_id": "cv-synthetic",
        "resume_line_order": ["L0001"],
        "resume_bullet_edits": [],
        "cover_letter_paragraphs": [
            {
                "text": "I built a Python service for local evaluation.",
                "claim_ids": ["claim-1"],
                "requirement_ids": ["REQ-01"],
                "source_urls": [JOB_URL],
                "research_finding_ids": [],
                "preference_source_ids": [],
            },
            {
                "text": "I tested a retrieval prototype with local fixtures.",
                "claim_ids": ["claim-2"],
                "requirement_ids": ["REQ-02"],
                "source_urls": [JOB_URL],
                "research_finding_ids": [],
                "preference_source_ids": [],
            },
        ],
        "application_answers": [],
        "outreach_drafts": [],
    }
    requirements = [
        {"requirement_id": "REQ-01", "claim_ids": ["claim-2"]},
        {"requirement_id": "REQ-02", "claim_ids": ["claim-1"]},
    ]

    _validate_rewrite(
        rewrite,
        "cv-synthetic",
        {"L0001": "Experience"},
        {"claim-1", "claim-2"},
        {"findings": [], "contacts": []},
        [],
        "preserve",
        job_source_url=JOB_URL,
        recruiter_requirements=requirements,
    )

    assert len(rewrite["cover_letter_paragraphs"]) == 2


def test_cover_letter_still_requires_exact_candidate_evidence_ids():
    rewrite = {
        "selected_cv_id": "cv-synthetic",
        "resume_line_order": ["L0001"],
        "resume_bullet_edits": [],
        "cover_letter_paragraphs": [
            {
                "text": "The role asks for Python service work.",
                "claim_ids": [],
                "cv_line_ids": [],
                "requirement_ids": ["REQ-01"],
                "source_urls": [JOB_URL],
                "research_finding_ids": [],
                "preference_source_ids": [],
            }
        ],
        "application_answers": [],
        "outreach_drafts": [],
    }

    with pytest.raises(OpenAIProviderError, match="must cite exact candidate evidence"):
        _validate_rewrite(
            rewrite,
            "cv-synthetic",
            {"L0001": "Experience"},
            set(),
            {"findings": [], "contacts": []},
            [],
            "preserve",
            job_source_url=JOB_URL,
            recruiter_requirements=[{"requirement_id": "REQ-01", "claim_ids": []}],
        )


def test_role_relevant_profile_evidence_is_selected_when_recruiter_map_is_incomplete():
    source = {
        "source_id": "technical-profile",
        "filename": "technical-profile.md",
        "text": (
            "# Ayni\nAyni is a platform for creating and operating specialized AI-agent workflows.\n\n"
            "# Medical Image Registration\n"
            "The work evolves from a TensorFlow VoxelMorph pipeline deployed on the AMD Kria KV260 DPU.\n"
        ),
        "use": "Owner-provided technical source.",
        "evidence": [
            {
                "id": "profile-ayni",
                "claim_text": "Ayni is a platform for creating and operating specialized AI-agent workflows.",
                "evidence_excerpt": "Ayni is a platform for creating and operating specialized AI-agent workflows.",
                "review_status": "unreviewed",
            },
            {
                "id": "profile-registration",
                "claim_text": "Medical image registration and FPGA/DPU deployment.",
                "evidence_excerpt": "Medical image registration and FPGA/DPU deployment.",
                "review_status": "unreviewed",
            },
            {
                "id": "profile-unrelated",
                "claim_text": "A small unrelated game interface project.",
                "evidence_excerpt": "A small unrelated game interface project.",
                "review_status": "unreviewed",
            },
        ],
    }
    positioning_claims = [
        "His computer-vision work includes medical image registration and anomaly detection, while his edge-AI work covers FPGA deployment and quantization.",
        "His work spans deep learning, computer vision, agentic AI, and hardware-aware inference.",
        "On the LLM side, he builds bounded agent architectures using structured outputs, retrieval, tools and actions.",
        "Juan Martin is listed as a team member/collaborator on this Advanced Deep Learning project.",
        "This is relevant because the portfolio supports agent-assisted workflows and AI service deployment.",
        "This is a strong example of research-oriented thinking around model evaluation and sequential decisions.",
        "This supports positioning for deployment-focused AI engineering roles.",
    ]
    source["evidence"].extend(
        {
            "id": f"positioning-{index}",
            "claim_text": text,
            "evidence_excerpt": text,
            "review_status": "unreviewed",
        }
        for index, text in enumerate(positioning_claims)
    )
    requirements = [
        {"requirement_id": "REQ-01", "requirement": "Build agent-assisted workflows."},
        {"requirement_id": "REQ-02", "requirement": "Deploy AI models for scientific work."},
    ]
    cv_lines = {
        "L0001": "Built an agent platform with retrieval memory and approval workflows.",
        "L0002": "Trained and deployed a quantized medical-image registration model on FPGA hardware.",
    }

    selected = _select_technical_profile_evidence(
        [source],
        set(),
        recruiter_requirements=requirements,
        cv_lines=cv_lines,
    )
    selected_context = _technical_profile_prompt_context(selected)
    selected_ids = {item["id"] for item in selected_context[0]["evidence"]}

    assert selected_ids == {"profile-ayni", "profile-registration"}
    assert not any(item_id.startswith("positioning-") for item_id in selected_ids)
    assert len(selected_ids) <= 6
    assert any(
        "TensorFlow VoxelMorph" in excerpt and "AMD Kria KV260 DPU" in excerpt
        for item in selected_context[0]["evidence"]
        if item["id"] == "profile-registration"
        for excerpt in item["source_excerpts"]
    )


def test_named_profile_projects_rank_ahead_of_generic_mapped_skills():
    generic_claims = [
        {
            "id": f"mapped-generic-{index}",
            "claim_text": claim,
            "evidence_excerpt": claim,
            "review_status": "unreviewed",
        }
        for index, claim in enumerate(
            [
                "Built Kubernetes services with autoscaling, telemetry, and container orchestration.",
                "Maintained continuous-integration pipelines with reproducible model evaluations.",
                "Configured cloud infrastructure as code and automated environment provisioning.",
                "Reviewed distributed service logs and improved operational monitoring dashboards.",
            ]
        )
    ]
    source = {
        "source_id": "technical-profile",
        "filename": "technical-profile.md",
        "text": (
            "# Ayni Agentic Systems Platform\n"
            "Ayni is a platform for creating and operating specialized AI-agent workflows.\n\n"
            "# Medical Image Registration and FPGA Deployment\n"
            "The work evolves from a TensorFlow VoxelMorph pipeline deployed on the AMD Kria KV260 DPU.\n\n"
            "# General Engineering Skills\n"
            "Kubernetes, continuous integration, infrastructure as code, and operational monitoring.\n"
        ),
        "use": "Owner-provided technical source.",
        "evidence": [
            {
                "id": "profile-ayni",
                "claim_text": "Ayni is a platform for creating and operating specialized AI-agent workflows.",
                "evidence_excerpt": "Ayni is a platform for creating and operating specialized AI-agent workflows.",
                "review_status": "unreviewed",
            },
            {
                "id": "profile-registration",
                "claim_text": (
                    "Medical image registration and FPGA/DPU deployment: the work evolves from "
                    "a TensorFlow VoxelMorph pipeline deployed on the AMD Kria KV260 DPU."
                ),
                "evidence_excerpt": "The work evolves from a TensorFlow VoxelMorph pipeline deployed on the AMD Kria KV260 DPU.",
                "review_status": "unreviewed",
            },
            *generic_claims,
        ],
    }
    requirements = [
        {
            "requirement_id": "REQ-01",
            "requirement": (
                "Build LLM applications, RAG systems, and agent-assisted workflows; deploy AI "
                "services for medical-image registration; maintain Kubernetes, CI, infrastructure "
                "as code, and operational monitoring."
            ),
            "claim_ids": [item["id"] for item in generic_claims],
        }
    ]

    selected = _select_technical_profile_evidence(
        [source],
        {item["id"] for item in generic_claims},
        recruiter_requirements=requirements,
        cv_lines={},
    )
    selected_ids = [
        item["id"]
        for item in _technical_profile_prompt_context(selected)[0]["evidence"]
    ]

    assert set(selected_ids[:2]) == {"profile-ayni", "profile-registration"}
    assert len(selected_ids) == 6


def test_recruiter_coverage_without_exact_cv_line_is_downgraded_to_uncertain():
    requirements = _bind_recruiter_requirements(
        [
            {
                "requirement": "Build retrieval services.",
                "priority": "required",
                "document_coverage": "covered",
                "claim_ids": ["profile-only"],
                "cv_line_ids": ["L9999"],
                "notes": "The CV supports this requirement.",
            },
            {
                "requirement": "Use Python.",
                "priority": "required",
                "document_coverage": "partial",
                "claim_ids": [],
                "cv_line_ids": ["L0001"],
                "notes": "Python is listed.",
            },
        ],
        valid_cv_line_ids={"L0001"},
        valid_claim_ids={"approved-claim"},
    )

    assert requirements[0]["requirement_id"] == "REQ-01"
    assert requirements[0]["document_coverage"] == "uncertain"
    assert requirements[0]["claim_ids"] == []
    assert requirements[0]["cv_line_ids"] == []
    assert "could not be verified" in requirements[0]["notes"]
    assert requirements[1]["document_coverage"] == "partial"
    assert requirements[1]["cv_line_ids"] == ["L0001"]


def test_cover_letter_uses_exact_cv_lines_for_recruiter_mapping_and_jev_grounding():
    cv_lines = {
        "L0001": "Built a Python retrieval service with fixture-based evaluation and a local FastAPI endpoint, comparing query results across a repeatable set of local cases.",
        "L0002": "Deployed a VoxelMorph medical-image registration model on an edge accelerator by preparing it for target hardware, checking registration output, and assessing runtime constraints.",
    }
    requirements = [
        {"requirement_id": "REQ-01", "claim_ids": [], "cv_line_ids": ["L0001"]},
        {"requirement_id": "REQ-02", "claim_ids": [], "cv_line_ids": ["L0002"]},
    ]
    rewrite = {
        "selected_cv_id": "cv-synthetic",
        "resume_line_order": ["L0001", "L0002"],
        "resume_bullet_edits": [],
        "cover_letter_paragraphs": _complete_quality_letter(
            responsibilities=("Python retrieval services", "model deployment")
        )["cover_letter_paragraphs"],
        "application_answers": [],
        "outreach_drafts": [],
    }
    rewrite["cover_letter_paragraphs"][1].update(
        {
            "text": (
                "I built a Python retrieval service with fixture-based evaluation and a local FastAPI endpoint. "
                "I compared query results across a repeatable set of local cases, giving the evaluation a consistent "
                "basis for reviewing retrieval behavior across runs rather than relying on isolated examples. "
                "This ties a working service to a repeatable way of examining its outputs. The posting identifies "
                "Python retrieval services and evaluation as responsibilities."
            ),
            "claim_ids": [],
            "cv_line_ids": ["L0001"],
            "requirement_ids": ["REQ-01"],
        }
    )
    rewrite["cover_letter_paragraphs"][2].update(
        {
            "text": (
                "I deployed a VoxelMorph medical-image registration model on an edge accelerator. "
                "I prepared it for target hardware, checked registration output along the deployment path, and assessed "
                "runtime constraints. Those checks keep the model's behavior visible under the "
                "conditions of running it on the target device, so the deployment account includes both output review "
                "and practical limits. Model deployment for scientific applications appears among the listed "
                "responsibilities."
            ),
            "claim_ids": [],
            "cv_line_ids": ["L0002"],
            "requirement_ids": ["REQ-02"],
        }
    )
    rewrite["cover_letter_paragraphs"][0]["requirement_ids"] = ["REQ-01"]
    rewrite["cover_letter_paragraphs"][3].update(
        {"claim_ids": [], "cv_line_ids": ["L0001", "L0002"]}
    )
    research = {"findings": [], "contacts": []}
    _validate_rewrite(
        rewrite,
        "cv-synthetic",
        cv_lines,
        set(),
        research,
        [],
        "preserve",
        job_source_url=JOB_URL,
        recruiter_requirements=requirements,
    )

    assertions = _grounding_assertions(
        rewrite,
        [],
        research,
        job_context={"canonical_url": JOB_URL, "title": "AI Engineer", "description": "Build Python retrieval services and deploy models."},
        cv_lines=cv_lines,
        cv_source_id="cv-synthetic",
    )
    by_id = {item["id"]: item for item in assertions}
    assert by_id["cover_letter_paragraph_1"]["evidence"][0] == {
        "source_type": "owner_provided_resume",
        "source_id": "cv-synthetic",
        "text": "Selected CV line L0001: " + cv_lines["L0001"],
    }
    assert by_id["cover_letter_paragraph_2"]["evidence"][0] == {
        "source_type": "owner_provided_resume",
        "source_id": "cv-synthetic",
        "text": "Selected CV line L0002: " + cv_lines["L0002"],
    }
    quality = _cover_letter_quality(
        rewrite,
        {
            "results": [
                    {"id": "cover_letter_paragraph_0", "output_type": "cover_letter_paragraph", "status": "supported"},
                    {"id": "cover_letter_paragraph_1", "output_type": "cover_letter_paragraph", "status": "supported"},
                    {"id": "cover_letter_paragraph_2", "output_type": "cover_letter_paragraph", "status": "supported"},
                    {"id": "cover_letter_paragraph_3", "output_type": "cover_letter_paragraph", "status": "supported"},
                    {"id": "cover_letter_paragraph_4", "output_type": "cover_letter_paragraph", "status": "supported"},
            ]
        },
        recruiter_requirements=requirements,
    )
    assert quality["status"] == "passed", quality["issues"]
    assert quality["candidate_evidence_and_role_requirements_cited"] is True


def test_unbound_cv_line_citation_is_removed_before_jev():
    rewrite = {
        "selected_cv_id": "cv-synthetic",
        "resume_line_order": ["L0001"],
        "resume_bullet_edits": [],
        "cover_letter_paragraphs": [
            {
                "text": "I built a Python service.",
                "claim_ids": [],
                "cv_line_ids": ["L9999"],
                "requirement_ids": ["REQ-01"],
                "research_finding_ids": [],
                "source_urls": [],
                "preference_source_ids": [],
            }
        ],
        "application_answers": [],
        "outreach_drafts": [],
    }
    _validate_rewrite(
        rewrite,
        "cv-synthetic",
        {"L0001": "Experience"},
        set(),
        {"findings": [], "contacts": []},
        [],
        "preserve",
        job_source_url=JOB_URL,
        recruiter_requirements=[{"requirement_id": "REQ-01", "claim_ids": [], "cv_line_ids": ["L0001"]}],
    )
    assert rewrite["cover_letter_paragraphs"] == []
    assert "outside this request's approved set" in rewrite["unresolved_questions"][0]


def test_cover_letter_quality_flags_exact_repeated_sentences_for_one_revision():
    repeated_sentence = "The team values careful Python service work and disciplined operational testing across its production systems."
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][1]["text"] = rewrite["cover_letter_paragraphs"][1]["text"].replace(
        "Northstar calls for developing agent-assisted workflows in shared services.",
        repeated_sentence,
    )
    rewrite["cover_letter_paragraphs"][2]["text"] = rewrite["cover_letter_paragraphs"][2]["text"].replace(
        "Northstar names model deployment as an AI Engineer responsibility.",
        repeated_sentence,
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["revision_paragraph_indices"] == [2]
    assert quality["body_word_count"] == sum(quality["paragraph_word_counts"])
    assert quality["sections_are_flexible"] is True
    assert "repeats a full sentence" in " ".join(quality["issues"])
    assert quality["owner_review_required"] is True
    assert quality["owner_review_notes"] == []
    assert "not a gate" in _packet_review_note(quality)


def test_cover_letter_quality_revises_repeated_evidence_when_multiple_examples_reuse_it():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][2]["claim_ids"] = ["claim-ayni"]

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["revision_paragraph_indices"] == [2]
    assert quality["revision_reasons"] == {"2": "duplicate_candidate_evidence"}
    assert "repeating the same candidate proof point" in " ".join(quality["issues"])


def test_cover_letter_quality_revises_concrete_example_and_both_relate_bridges():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][1]["text"] = rewrite["cover_letter_paragraphs"][1]["text"].replace(
        "Northstar calls for developing agent-assisted workflows in shared services.",
        "The Ayni work provides a concrete example of Northstar's agent-assisted workflow responsibility.",
    )
    rewrite["cover_letter_paragraphs"][3]["text"] = (
        "Both relate to Northstar's responsibility for operating AI services for research teams. "
        "Together, the projects show attention to evaluation and operational controls across distinct systems."
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["boilerplate_bridge_paragraph_indices"] == [1, 3]
    assert quality["revision_reasons"]["1"] == "generic_relevance_bridge"
    assert quality["revision_reasons"]["3"] == "generic_relevance_bridge"


def test_cover_letter_quality_accepts_a_substantive_owner_specific_letter():
    rewrite = _complete_quality_letter()

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "passed"
    assert quality["rubric_version"] == "research-informed-cover-letter-argument-v28"
    assert quality["body_paragraph_count"] == 5
    assert len(quality["paragraph_word_counts"]) == 5
    assert sum(quality["paragraph_word_counts"]) == quality["body_word_count"]
    assert len(quality["paragraph_sentence_counts"]) == 5
    assert quality["section_labels_used"] == ["opening", "evidence", "evidence", "contribution", "closing"]
    assert quality["sections_are_flexible"] is True
    assert quality["one_page_is_practical_guidance_not_a_gate"] is True
    assert quality["boilerplate_bridge_paragraph_indices"] == []
    assert quality["role_duty_list_paragraph_indices"] == []


def test_cover_letter_quality_revises_paragraphs_that_refer_to_the_cv_as_a_document():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][2]["text"] = rewrite["cover_letter_paragraphs"][2]["text"].replace(
        "For Medical Image Registration,",
        "My CV notes that for Medical Image Registration,",
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["cv_reference_paragraph_indices"] == [2]
    assert quality["revision_paragraph_indices"] == [2]
    assert quality["revision_reasons"] == {"2": "cv_document_reference"}


def test_cover_letter_quality_does_not_use_word_or_paragraph_count_as_a_gate():
    rewrite = {
        "cover_letter_paragraphs": [
            {
                "section": "evidence",
                "text": (
                    "For Medical Image Registration, I deployed a TensorFlow VoxelMorph pipeline on an AMD Kria KV260 DPU. "
                    "Deploying and operating AI services are responsibilities in IFOM's role."
                ),
                "claim_ids": ["claim-registration"],
                "requirement_ids": ["REQ-01"],
                "source_urls": [JOB_URL],
            },
            {
                "section": "evidence",
                "text": (
                    "I built Ayni, a platform for creating and operating specialized AI-agent workflows. "
                    "Agent-assisted workflows are also an explicit part of IFOM's role."
                ),
                "claim_ids": ["claim-ayni"],
                "requirement_ids": ["REQ-02"],
                "source_urls": [JOB_URL],
            },
        ]
    }

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding(2))

    assert quality["status"] == "passed", quality["issues"]
    assert quality["body_paragraph_count"] == 2
    assert quality["body_word_count"] <= 60
    assert quality["sections_are_flexible"] is True
    assert quality["one_page_is_practical_guidance_not_a_gate"] is True

    reviewer = _validate_letter_review(
        {
            "decision": "needs_revision",
            "dimensions": [
                {
                    "name": name,
                    "score": 2 if name in {"letter_completeness", "candidate_evidence_development"} else 3,
                    "reason": "The synthetic sample is too note-like." if name == "letter_completeness" else "Criterion reviewed.",
                }
                for name in LETTER_REVIEW_DIMENSION_NAMES
            ],
            "issues": [{"issue": "Develop the application case beyond a note.", "paragraph_indices": [0, 1]}],
        },
        rewrite["cover_letter_paragraphs"],
    )
    assert reviewer["status"] == "needs_revision"
    assert reviewer["revision_paragraph_indices"] == [0, 1]


def test_letter_reviewer_flags_a_distinct_missing_profile_example_without_a_count_rule():
    paragraphs = [
        {
            "section": "evidence",
            "text": "I built Ayni for agent-assisted workflows.",
            "claim_ids": ["claim-ayni"],
            "requirement_ids": ["REQ-02"],
            "source_urls": [JOB_URL],
        }
    ]
    evidence = {
        "claim-ayni": {"claim_text": "Worked on Ayni agent workflows."},
        "claim-registration": {
            "claim_text": "Deployed a VoxelMorph pipeline on the AMD Kria KV260 DPU.",
            "related_cv_lines": [{"line_id": "L0004", "text": "Deployed a VoxelMorph pipeline."}],
        },
    }
    review = _validate_letter_review(
        {
            "decision": "ready",
            "dimensions": [
                {"name": name, "score": 3, "reason": "The body names relevant work."}
                for name in LETTER_REVIEW_DIMENSION_NAMES
            ],
            "issues": [],
            "profile_evidence_assessment": [
                {
                    "claim_id": "claim-ayni",
                    "status": "included",
                    "paragraph_index": 0,
                    "rationale": "The Ayni project is developed in paragraph one.",
                },
                {
                    "claim_id": "claim-registration",
                    "status": "material_gap",
                    "paragraph_index": None,
                    "rationale": "The separate edge deployment proof adds a distinct required model-deployment responsibility.",
                },
            ],
        },
        paragraphs,
        preferred_profile_claim_ids=["claim-ayni", "claim-registration"],
        preferred_profile_claim_evidence=evidence,
    )

    assert review["status"] == "needs_revision"
    assert review["revision_paragraph_indices"] == [1]
    assert review["revision_reasons"] == {"1": "missing_candidate_evidence"}
    assert [item["score"] for item in review["dimensions"] if item["name"] in {
        "letter_completeness", "candidate_evidence_development"
    }] == [2, 2]
    assert review["profile_evidence_assessment"][1]["status"] == "material_gap"


def test_letter_repair_context_can_append_a_missing_material_profile_example():
    paragraphs = [{
        "section": "evidence",
        "text": "I built Ayni for agent-assisted workflows.",
        "claim_ids": ["claim-ayni"],
        "cv_line_ids": ["L0001"],
        "requirement_ids": ["REQ-02"],
        "source_urls": [JOB_URL],
    }]
    technical_profile = [{
        "source_id": "technical-profile",
        "filename": "technical.md",
        "use": "cover_letter",
        "evidence": [
            {
                "id": "claim-ayni",
                "claim_text": "Worked on Ayni agent workflows.",
                "evidence_excerpt": "I worked on Ayni agent workflows.",
                "source_excerpts": ["I worked on Ayni agent workflows."],
                "related_cv_lines": [{"line_id": "L0001", "text": "Built Ayni."}],
                "review_status": "unreviewed",
            },
            {
                "id": "claim-registration",
                "claim_text": "Deployed a VoxelMorph pipeline on the AMD Kria KV260 DPU.",
                "evidence_excerpt": "I deployed a VoxelMorph pipeline on the AMD Kria KV260 DPU.",
                "source_excerpts": ["I deployed a VoxelMorph pipeline on the AMD Kria KV260 DPU."],
                "related_cv_lines": [{"line_id": "L0002", "text": "Deployed a medical image model."}],
                "review_status": "unreviewed",
            },
        ],
    }]

    context = _letter_repair_context(
        job={"title": "AI Engineer", "company": "Northstar", "canonical_url": JOB_URL},
        original_paragraphs=paragraphs,
        grounding=_supported_letter_grounding(1),
        failed_indices=[1],
        repair_reasons={1: "missing_candidate_evidence"},
        approved_claims=[],
        technical_profile_sources=technical_profile,
        preferred_profile_claim_ids=["claim-ayni", "claim-registration"],
        writing_preferences={},
        owner_writing_context=[],
        research={"findings": []},
        recruiter_requirements=[
            {"requirement_id": "REQ-01", "requirement": "Deploy AI models for scientific work."},
            {"requirement_id": "REQ-02", "requirement": "Build agent-assisted workflows."},
        ],
        cv_lines={"L0001": "Built Ayni.", "L0002": "Deployed a medical image model."},
        quality_issues=["The letter omits distinct model deployment evidence."],
    )

    append = context["paragraphs_for_repair"]
    assert len(append) == 1
    assert append[0]["paragraph_index"] == 1
    assert append[0]["section"] == "evidence"
    assert append[0]["repair_reason"] == "missing_candidate_evidence"
    assert append[0]["allowed_claim_ids"] == ["claim-registration"]
    assert append[0]["allowed_cv_line_ids"] == ["L0002"]
    assert context["preferred_technical_profile_claim_ids"] == ["claim-registration"]


def test_one_supported_example_and_requirement_can_pass_without_section_labels():
    rewrite = {
        "cover_letter_paragraphs": [
            {
                "section": None,
                "text": (
                    "I deployed a TensorFlow VoxelMorph pipeline on an AMD Kria KV260 DPU, "
                    "addressing this role's model-deployment responsibility."
                ),
                "claim_ids": ["claim-registration"],
                "requirement_ids": ["REQ-01"],
                "source_urls": [JOB_URL],
            }
        ]
    }
    quality = _cover_letter_quality(
        rewrite,
        _supported_letter_grounding(1),
        recruiter_requirements=[
            {"requirement_id": "REQ-01", "requirement": "Deploy AI models for scientific work."}
        ],
    )

    assert quality["status"] == "passed", quality["issues"]
    assert quality["body_paragraph_count"] == 1
    assert quality["distinct_role_requirement_count"] == 1
    assert quality["distinct_candidate_evidence_count"] == 1
    assert quality["section_labels_used"] == [None]


def test_omitted_unresolved_jev_paragraph_does_not_fail_the_final_letter():
    rewrite = {
        "cover_letter_paragraphs": [{
            "section": None,
            "text": (
                "I deployed a TensorFlow VoxelMorph pipeline on an AMD Kria KV260 DPU, "
                "addressing this role's model-deployment responsibility."
            ),
            "claim_ids": ["claim-registration"],
            "requirement_ids": ["REQ-01"],
            "source_urls": [JOB_URL],
        }]
    }
    grounding = _supported_letter_grounding(1)
    grounding["results"].append({
        "id": "cover_letter_paragraph_1",
        "output_type": "cover_letter_paragraph",
        "status": "unresolved",
        "included_in_artifact": False,
        "reason": "This unsupported block was removed before the final letter.",
    })

    quality = _cover_letter_quality(rewrite, grounding)

    assert quality["status"] == "passed", quality["issues"]
    assert quality["jev_rejected_paragraphs"] == []


def test_grounding_maps_supported_blocks_to_their_final_artifact_indexes():
    paragraph = {
        "section": None,
        "text": (
            "I deployed a TensorFlow VoxelMorph pipeline on an AMD Kria KV260 DPU, "
            "addressing this role's model-deployment responsibility."
        ),
        "claim_ids": ["claim-registration"],
        "requirement_ids": ["REQ-01"],
        "source_urls": [JOB_URL],
    }
    rewrite = {
        "cover_letter_paragraphs": [
            {**paragraph, "text": "An unsupported paragraph that Jev removes."},
            paragraph,
        ],
        "resume_bullet_edits": [
            {"line_id": "line-1", "revised_text": "Unsupported edit."},
            {"line_id": "line-2", "revised_text": "Supported edit."},
        ],
    }
    assertions = [
        {
            "id": "cover_letter_paragraph_0",
            "output_type": "cover_letter_paragraph",
            "text": rewrite["cover_letter_paragraphs"][0]["text"],
            "evidence": [{"source_type": "cv", "text": "Candidate evidence."}],
        },
        {
            "id": "cover_letter_paragraph_1",
            "output_type": "cover_letter_paragraph",
            "text": paragraph["text"],
            "evidence": [{"source_type": "cv", "text": "Candidate evidence."}],
        },
        {
            "id": "resume_bullet_edit_0",
            "output_type": "resume_bullet_edit",
            "text": "Unsupported edit.",
            "evidence": [{"source_type": "cv", "text": "Candidate evidence."}],
        },
        {
            "id": "resume_bullet_edit_1",
            "output_type": "resume_bullet_edit",
            "text": "Supported edit.",
            "evidence": [{"source_type": "cv", "text": "Candidate evidence."}],
        },
    ]
    result = {
        "status": "complete",
        "results": [
            {
                "id": item["id"],
                "status": "unresolved" if item["id"].endswith("_0") else "supported",
                "reason": "No support." if item["id"].endswith("_0") else "Supported.",
            }
            for item in assertions
        ],
    }

    grounding = _apply_grounding_results(rewrite, result, assertions)

    assert rewrite["cover_letter_paragraphs"] == [paragraph]
    assert rewrite["resume_bullet_edits"] == [{"line_id": "line-2", "revised_text": "Supported edit."}]
    grounding_by_source_id = {item["id"]: item for item in grounding["results"]}
    assert grounding_by_source_id["cover_letter_paragraph_0"]["final_artifact_id"] is None
    assert grounding_by_source_id["cover_letter_paragraph_1"]["final_artifact_id"] == "cover_letter_paragraph_0"
    assert grounding_by_source_id["resume_bullet_edit_1"]["final_artifact_id"] == "resume_bullet_edit_0"

    quality = _cover_letter_quality(
        rewrite,
        grounding,
        recruiter_requirements=[
            {"requirement_id": "REQ-01", "requirement": "Deploy AI models for scientific work."}
        ],
    )
    assert quality["status"] == "passed", quality["issues"]
    assert quality["jev_rejected_paragraphs"] == []


def test_cover_letter_reports_unused_preferred_profile_claims_without_a_count_gate():
    preferred = ["profile-project-a", "profile-project-b"]
    rewrite = _complete_quality_letter(preferred_profile_claim_ids=preferred)
    rewrite["cover_letter_paragraphs"][2]["claim_ids"] = ["claim-registration"]
    rewrite["cover_letter_paragraphs"][3]["claim_ids"] = ["claim-ayni", "claim-registration"]
    grounding = _supported_letter_grounding()

    missing = _cover_letter_quality(
        rewrite, grounding, preferred_profile_claim_ids=preferred
    )
    assert missing["status"] == "passed"
    assert missing["preferred_profile_claim_count"] == 2
    assert missing["preferred_profile_claims_used"] == 1
    assert missing["preferred_profile_claim_ids_not_used"] == [preferred[1]]
    assert missing["revision_paragraph_indices"] == []

    sourced = _complete_quality_letter(preferred_profile_claim_ids=preferred)
    passed = _cover_letter_quality(
        sourced, grounding, preferred_profile_claim_ids=preferred
    )
    assert passed["status"] == "passed"
    assert passed["preferred_profile_claims_used"] == 2
    assert passed["profile_claim_count_by_paragraph"] == [0, 1, 1, 2, 0]

    duplicated = _complete_quality_letter(preferred_profile_claim_ids=preferred)
    duplicated["cover_letter_paragraphs"][2]["claim_ids"] = [preferred[0]]
    duplicate_quality = _cover_letter_quality(
        duplicated, grounding, preferred_profile_claim_ids=preferred
    )
    assert duplicate_quality["status"] == "failed"
    assert duplicate_quality["revision_paragraph_indices"] == [2]

    swapped = _complete_quality_letter(preferred_profile_claim_ids=preferred)
    swapped["cover_letter_paragraphs"][1]["claim_ids"] = [preferred[1]]
    swapped["cover_letter_paragraphs"][2]["claim_ids"] = [preferred[0]]
    swapped_quality = _cover_letter_quality(
        swapped, grounding, preferred_profile_claim_ids=preferred
    )
    assert swapped_quality["status"] == "passed"
    assert swapped_quality["preferred_profile_claims_used"] == 2


def test_preferred_profile_letter_requires_owner_action_and_distinctive_source_detail():
    preferred = ["profile-registration", "profile-ayni"]
    evidence = {
        preferred[0]: {
            "claim_text": (
                "The work evolves from a TensorFlow VoxelMorph pipeline deployed on the "
                "AMD Kria KV260 DPU into a more advanced 2.5D PyTorch/Vitis AI registration pipeline."
            ),
            "evidence_excerpt": (
                "The work evolves from a TensorFlow VoxelMorph pipeline deployed on the "
                "AMD Kria KV260 DPU into a more advanced 2.5D PyTorch/Vitis AI registration pipeline."
            ),
        },
        preferred[1]: {
            "claim_text": "Ayni is a platform for creating and operating specialized AI-agent workflows.",
            "evidence_excerpt": "Ayni is a platform for creating and operating specialized AI-agent workflows.",
        },
    }
    direct_owner_letter = _complete_quality_letter(
        preferred_profile_claim_ids=preferred, registration_first=True
    )
    rejected = deepcopy(direct_owner_letter)
    rejected["cover_letter_paragraphs"][1]["text"] = rejected["cover_letter_paragraphs"][1]["text"].replace(
        "For Medical Image Registration, I deployed a VoxelMorph pipeline on the AMD Kria KV260 DPU.",
        "The Medical Image Registration project examined a model on edge hardware in considerable technical detail.",
    )
    rejected["cover_letter_paragraphs"][2]["text"] = rejected["cover_letter_paragraphs"][2]["text"].replace(
        "I worked on Ayni, a platform for creating and operating specialized AI-agent workflows.",
        "Ayni is a platform for creating and operating specialized AI-agent workflows.",
    )
    rejected_quality = _cover_letter_quality(
        rejected,
        _supported_letter_grounding(),
        preferred_profile_claim_ids=preferred,
        preferred_profile_claim_evidence=evidence,
    )
    assert rejected_quality["status"] == "failed"
    assert rejected_quality["direct_first_person_action_paragraph_indices"] == [1]
    assert rejected_quality["missing_preferred_profile_detail_paragraph_indices"] == [1]
    assert rejected_quality["revision_paragraph_indices"] == [1]

    direct_owner_letter["cover_letter_paragraphs"][2]["text"] = direct_owner_letter["cover_letter_paragraphs"][2]["text"].replace(
        "I worked on Ayni, a platform", "I worked on a platform"
    )
    accepted = _cover_letter_quality(
        direct_owner_letter,
        _supported_letter_grounding(),
        preferred_profile_claim_ids=preferred,
        preferred_profile_claim_evidence=evidence,
    )
    assert accepted["status"] == "passed"
    assert accepted["rubric_version"] == "research-informed-cover-letter-argument-v28"
    assert accepted["direct_first_person_action_paragraph_indices"] == []
    assert accepted["missing_preferred_profile_detail_paragraph_indices"] == []


def test_preferred_profile_claim_picker_uses_named_projects_before_generic_stack():
    sources = [{
        "evidence": [
            {
                "id": "software-stack",
                "claim_text": (
                    "He also develops the surrounding software stack: Python/FastAPI/Pydantic "
                    "services, React/TypeScript applications, persistence, testing, and deployment tooling."
                ),
                "evidence_excerpt": "He also develops the surrounding software stack.",
            },
            {
                "id": "registration",
                "claim_text": (
                    "The work evolves from a TensorFlow VoxelMorph pipeline deployed on the "
                    "AMD Kria KV260 DPU into a more advanced 2.5D PyTorch/Vitis AI registration pipeline."
                ),
                "evidence_excerpt": (
                    "The work evolves from a TensorFlow VoxelMorph pipeline deployed on the "
                    "AMD Kria KV260 DPU into a more advanced 2.5D PyTorch/Vitis AI registration pipeline."
                ),
                "source_excerpts": ["FPGA DPU deployment; CPU/GPU/FPGA comparisons; evaluation and benchmarking."],
            },
            {
                "id": "ayni",
                "claim_text": "Ayni is a platform for creating and operating specialized AI-agent workflows.",
                "evidence_excerpt": "Ayni is a platform for creating and operating specialized AI-agent workflows.",
                "source_excerpts": [
                    "Agent registration, bounded tool invocation, human approval, evaluation, and recovery.",
                ],
            },
        ],
    }]
    requirements = [
        {"requirement": "Design and deploy AI-enabled services for scientific work.", "priority": "required"},
        {"requirement": "Build LLM applications, RAG systems, and agent-assisted workflows.", "priority": "required"},
        {"requirement": "Deploy and operate AI services; support evaluation and maintainability.", "priority": "required"},
    ]

    preferred = _preferred_cover_letter_profile_claim_ids(sources, requirements)

    assert preferred == ["ayni", "registration"]


def test_preferred_profile_picker_keeps_one_relevant_project_without_a_two_claim_quota():
    evidence = [{
        "id": "ayni",
        "claim_text": "I worked on Ayni, a platform for creating and operating specialized AI-agent workflows.",
        "evidence_excerpt": "I worked on Ayni, a platform for creating and operating specialized AI-agent workflows.",
        "source_excerpts": ["Ayni supports bounded tool execution and workflow evaluation."],
        "related_cv_lines": [{"line_id": "L0034", "text": "Built a platform for agent workflows."}],
    }]
    requirements = [{
        "requirement": "Build and operate AI-agent workflows.",
        "priority": "required",
    }]

    assert _preferred_cover_letter_profile_claim_ids(
        [{"evidence": evidence}], requirements
    ) == ["ayni"]


def test_owner_written_example_order_wins_over_relevance_rank():
    evidence = [
        {
            "id": "registration",
            "claim_text": "I deployed a VoxelMorph pipeline on the AMD Kria KV260 DPU for medical image registration.",
            "evidence_excerpt": "I deployed a VoxelMorph pipeline on the AMD Kria KV260 DPU for medical image registration.",
            "source_id": "owner-examples",
            "owner_written_source": True,
            "source_order": 110,
            "related_cv_lines": [{"line_id": "L0042", "text": "Deployed a VoxelMorph pipeline on AMD Kria KV260 hardware."}],
        },
        {
            "id": "ayni",
            "claim_text": "I worked on Ayni, a platform for creating and operating specialized AI-agent workflows.",
            "evidence_excerpt": "I worked on Ayni, a platform for creating and operating specialized AI-agent workflows.",
            "source_id": "owner-examples",
            "owner_written_source": True,
            "source_order": 20,
            "related_cv_lines": [{"line_id": "L0034", "text": "Built an agent platform with model routing and retrieval memory."}],
        },
    ]
    requirements = [
        {"requirement": "Deploy AI models and services for medical image registration.", "priority": "required"},
        {"requirement": "Build agent-assisted workflows and LLM applications.", "priority": "required"},
    ]

    assert _preferred_cover_letter_profile_claim_ids(
        [{"evidence": evidence}], requirements
    ) == ["ayni", "registration"]


def test_preferred_letter_projects_must_match_the_selected_cv_when_two_are_available():
    source = {
        "source_id": "technical-profile",
        "filename": "technical-profile.md",
        "use": "Owner-provided technical source.",
        "text": (
            "# Ayni Agentic Systems Platform\n"
            "Ayni is a platform for creating and operating specialized AI-agent workflows.\n\n"
            "# Medical Image Registration and FPGA Deployment\n"
            "The work evolves from a TensorFlow VoxelMorph pipeline deployed on the AMD Kria KV260 DPU.\n\n"
            "# Project Marx\n"
            "Project Marx evaluates a separate research problem with machine-learning methods.\n"
        ),
        "evidence": [
            {
                "id": "ayni",
                "claim_text": "Ayni is a platform for creating and operating specialized AI-agent workflows.",
                "evidence_excerpt": "Ayni is a platform for creating and operating specialized AI-agent workflows.",
                "review_status": "unreviewed",
            },
            {
                "id": "registration",
                "claim_text": "The work evolves from a TensorFlow VoxelMorph pipeline deployed on the AMD Kria KV260 DPU.",
                "evidence_excerpt": "The work evolves from a TensorFlow VoxelMorph pipeline deployed on the AMD Kria KV260 DPU.",
                "review_status": "unreviewed",
            },
            {
                "id": "marx",
                "claim_text": "Project Marx evaluates a separate research problem with machine-learning methods.",
                "evidence_excerpt": "Project Marx evaluates a separate research problem with machine-learning methods.",
                "review_status": "unreviewed",
            },
        ],
    }
    cv_lines = {
        "L0032": "Projects",
        "L0033": "AI Agent Orchestration Platform | Independent project 2026",
        "L0034": "Built an agent platform with model routing, tool execution, retrieval memory, and CRM API integration.",
        "L0035": "Added evaluation, approval workflows, and recovery safeguards.",
        "L0041": "Medical Image Registration | Independent project 2025",
        "L0042": "Trained and deployed a quantized TensorFlow VoxelMorph pipeline on AMD Kria KV260 FPGA hardware.",
        "L0043": "Achieved 66% faster inference and warping than CPU while maintaining quality.",
    }
    requirements = [
        {
            "requirement_id": "REQ-01",
            "requirement": "Build LLM applications and agent-assisted workflows.",
            "priority": "required",
        },
        {
            "requirement_id": "REQ-02",
            "requirement": "Deploy AI models and services for medical image registration.",
            "priority": "required",
        },
        {
            "requirement_id": "REQ-03",
            "requirement": "Apply machine-learning methods to scientific research.",
            "priority": "required",
        },
    ]

    selected = _select_technical_profile_evidence(
        [source],
        set(),
        recruiter_requirements=requirements,
        cv_lines=cv_lines,
    )
    evidence = _technical_profile_prompt_context(selected)[0]["evidence"]
    preferred = _preferred_cover_letter_profile_claim_ids(selected, requirements)

    assert preferred == ["ayni", "registration"]
    assert all(
        item["related_cv_lines"]
        for item in evidence
        if item["id"] in preferred
    )
    assert "marx" not in preferred


def test_preferred_profile_anchor_picker_skips_sentence_pronouns_and_keeps_project_names():
    assert _preferred_profile_letter_anchors({
        "claim_text": "He also develops the surrounding software stack: Python/FastAPI/Pydantic services."
    }) == ["Python/FastAPI/Pydantic"]
    assert _preferred_profile_letter_anchors({
        "claim_text": "Ayni is a platform for creating and operating specialized AI-agent workflows."
    }) == ["Ayni"]
    assert _preferred_profile_letter_anchors({
        "claim_text": (
            "The work evolves from a TensorFlow VoxelMorph pipeline deployed on the "
            "AMD Kria KV260 DPU into a more advanced registration pipeline."
        )
    }) == ["VoxelMorph", "KV260"]
    assert _preferred_profile_letter_anchors({
        "claim_text": "Separately evaluated a local retrieval workflow using synthetic fixtures."
    }) == []


def test_letter_requirement_ids_are_normalized_or_bound_to_paragraph_content():
    requirements = [
        {"requirement_id": "REQ-01", "requirement": "Use Python for API development.", "priority": "required"},
        {"requirement_id": "REQ-02", "requirement": "Build LLM applications and agent-assisted workflows.", "priority": "required"},
        {"requirement_id": "REQ-03", "requirement": "Deploy and evaluate VoxelMorph registration models.", "priority": "required"},
    ]
    paragraphs = [
        {
            "section": "opening",
            "text": "I am applying for the AI Engineer role, which calls for agent-assisted workflows.",
            "requirement_ids": ["req-2"],
        },
        {
            "section": "evidence",
            "text": "I built Ayni for agent-assisted workflows and LLM applications.",
            "requirement_ids": ["REQ-01"],
        },
        {
            "section": "evidence",
            "text": "I deployed and evaluated a VoxelMorph registration model on FPGA hardware.",
            "requirement_ids": ["unmapped-role-reference"],
        },
    ]

    repairs = _repair_letter_requirement_ids(paragraphs, requirements)

    assert [paragraph["requirement_ids"] for paragraph in paragraphs] == [
        ["REQ-02"], ["REQ-02"], ["REQ-03"],
    ]
    assert [repair["paragraph_index"] for repair in repairs] == [0, 1, 2]


def test_profile_evidence_quality_failure_is_repaired_and_rechecked(
    settings, database, monkeypatch
):
    record, _cv, _technical, _descriptive, _sample, _claim_id, _resume = (
        _create_preparation(
            settings,
            database,
            technical_text=(
                "Built reliable Python services and a typed local workflow.\n"
                "Developed a Python workflow-testing harness for local synthetic reports.\n"
                "Tested a simulated maintenance workflow with synthetic reports to measure duplicate handling.\n"
            ),
        )
    )
    enabled = _enable_synthetic_openai(settings, database)
    monkeypatch.setattr(FakePreparationClient, "omit_preferred_profile_claims", True)
    monkeypatch.setattr(FakePreparationClient, "force_letter_review", True)
    FakePreparationClient.instances.clear()
    support_calls = []

    def support_all(_database, _settings, _request_id, assertions):
        support_calls.append(assertions)
        return {
            "status": "complete",
            "model": "jev-test",
            "actual_cost_usd": 0.0,
            "results": [
                {"id": item["id"], "status": "supported", "confidence": 0.9}
                for item in assertions
            ],
        }

    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        claim_support_checker=support_all,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )

    completed = get_preparation(database, record["id"])
    assert completed["state"] == "review", completed["status_message"]
    with connect(database) as db:
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]
    packet = get_packet(database, packet_id)
    output = packet["output"]
    paragraphs = output["rewriter"]["cover_letter_paragraphs"]
    fake_calls = FakePreparationClient.instances[-1].calls
    rewriter_context = json.loads(
        next(
            call["input"][0]["content"]
            for call in fake_calls
            if call["text"]["format"]["name"].startswith("clue_rewriter_")
        )
    )
    preferred_ids = rewriter_context["preferred_technical_profile_claim_ids"]

    assert len(preferred_ids) == 2, rewriter_context["technical_profile_sources"]
    reviser_context = json.loads(
        next(
            call["input"][0]["content"]
            for call in fake_calls
            if call["text"]["format"]["name"] == "clue_letter_reviser_v12"
        )
    )
    assert reviser_context["cover_letter_quality_issues"]
    assert all("text" in item and "paragraph_index" in item for item in reviser_context["paragraphs_for_repair"])
    assert reviser_context["research"]["contacts"] == []
    assert all(
        item["id"] in set(preferred_ids)
        for source in reviser_context["technical_profile_sources"]
        for item in source["evidence"]
    )
    assert len(support_calls) == 2
    assert output["quality_check"]["status"] == "passed"
    assert output["quality_check"]["revision"]["status"] == "completed"
    assert output["quality_check"]["revision"]["style_revision_paragraph_indices"]
    assert [item["section"] for item in paragraphs] == [
        "opening", "evidence", "evidence", "contribution", "closing"
    ]
    assert any(
        set(paragraph["claim_ids"]) & set(preferred_ids)
        for paragraph in paragraphs[1:3]
    )
    assert output["quality_check"]["preferred_profile_claims_used"] >= 1
    assert output["jev_tailored_resume_review"]["status"] == "approved"
    assert packet["artifacts"]


def test_cover_letter_quality_revises_formulaic_role_includes_lists():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][1]["text"] = rewrite["cover_letter_paragraphs"][1]["text"].replace(
        "Northstar calls for developing agent-assisted workflows in shared services.",
        "Northstar's role includes agent workflows, retrieval features, and technical operations.",
    )
    rewrite["cover_letter_paragraphs"][2]["text"] = rewrite["cover_letter_paragraphs"][2]["text"].replace(
        "Northstar names model deployment as an AI Engineer responsibility.",
        "The AI Engineer position includes model deployment, infrastructure operations, and monitoring.",
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["role_duty_list_paragraph_indices"] == [1, 2]
    assert quality["revision_paragraph_indices"] == [1, 2]
    assert quality["revision_reasons"] == {"1": "role_duty_list", "2": "role_duty_list"}


def test_cover_letter_quality_allows_one_concrete_role_includes_responsibility():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][1]["text"] = rewrite["cover_letter_paragraphs"][1]["text"].replace(
        "Northstar calls for developing agent-assisted workflows in shared services.",
        "Northstar's role includes developing agent-assisted workflows for shared services.",
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert _contains_role_duty_inventory("IFOM's role includes deploying AI-enabled services for science.") is False
    assert _contains_role_duty_inventory("The position includes deploying services and evaluating their performance.") is False
    assert _contains_role_duty_inventory("The role includes agent workflows, model evaluation, and deployment.") is True
    assert quality["status"] == "passed"
    assert quality["role_duty_list_paragraph_indices"] == []


def test_cover_letter_quality_flags_repeated_role_duty_sentence_frame():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][1]["text"] = rewrite["cover_letter_paragraphs"][1]["text"].replace(
        "Northstar calls for developing agent-assisted workflows in shared services.",
        "Northstar's role includes developing agent-assisted workflows in shared services.",
    )
    rewrite["cover_letter_paragraphs"][2]["text"] = rewrite["cover_letter_paragraphs"][2]["text"].replace(
        "Northstar names model deployment as an AI Engineer responsibility.",
        "Northstar's role includes model deployment for AI services.",
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["role_duty_list_paragraph_indices"] == []
    assert quality["repeated_role_duty_frame_paragraph_indices"] == [1, 2]
    assert quality["revision_reasons"] == {
        "1": "repeated_role_duty_frame",
        "2": "repeated_role_duty_frame",
    }


def test_cover_letter_quality_flags_generic_opening_preview_and_example_bridge():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][0]["text"] += (
        " My relevant project work has covered agent workflows and model deployment."
    )
    rewrite["cover_letter_paragraphs"][2]["text"] += (
        " This project is a concrete example of AI deployment, relevant to Northstar's responsibility."
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["boilerplate_bridge_paragraph_indices"] == [0, 2]
    assert quality["revision_paragraph_indices"] == [0, 2]
    assert quality["revision_reasons"] == {
        "0": "generic_relevance_bridge",
        "2": "generic_relevance_bridge",
    }


def test_cover_letter_quality_rejects_malformed_replacement_character():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][2]["text"] = rewrite["cover_letter_paragraphs"][2]["text"].replace(
        "VoxelMorph pipeline", "VoxelMorph� pipeline", 1
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["replacement_character_paragraph_indices"] == [2]
    assert quality["revision_paragraph_indices"] == [2]
    assert quality["revision_reasons"] == {"2": "encoding_garble"}


def test_cover_letter_quality_rejects_padded_relevance_bridge_without_length_limit():
    rewrite = _complete_quality_letter()
    paragraph = rewrite["cover_letter_paragraphs"][1]
    paragraph["text"] = paragraph["text"].replace(
        "Northstar calls for developing agent-assisted workflows in shared services.",
        "This example connects directly to Northstar's listed evaluation work. "
        "Northstar's AI Engineer role calls for developing and operating agent-assisted workflows.",
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["boilerplate_bridge_paragraph_indices"] == [1]
    assert quality["revision_paragraph_indices"] == [1]
    assert quality["revision_reasons"] == {"1": "generic_relevance_bridge"}


def test_cover_letter_quality_catches_subjectless_matches_formula():
    rewrite = _complete_quality_letter()
    for index, project_sentence in ((1, "This matches Northstar's responsibility to build agent-assisted workflows."),
                                    (2, "This matches Northstar's responsibility to deploy AI services.")):
        text_value = rewrite["cover_letter_paragraphs"][index]["text"]
        final_sentence = text_value.split(". ")[-1]
        rewrite["cover_letter_paragraphs"][index]["text"] = text_value.replace(final_sentence, project_sentence)

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["boilerplate_bridge_paragraph_indices"] == [1, 2]
    assert quality["revision_reasons"] == {
        "1": "generic_relevance_bridge",
        "2": "generic_relevance_bridge",
    }


def test_cover_letter_quality_rejects_repeated_role_link_verb():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][1]["text"] = rewrite["cover_letter_paragraphs"][1]["text"].replace(
        "Northstar calls for developing agent-assisted workflows in shared services.",
        "Its use overlaps with Northstar's agent-assisted workflow responsibility.",
    )
    rewrite["cover_letter_paragraphs"][2]["text"] = rewrite["cover_letter_paragraphs"][2]["text"].replace(
        "Northstar names model deployment as an AI Engineer responsibility.",
        "Its deployment overlaps with Northstar's model deployment responsibility.",
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["boilerplate_bridge_paragraph_indices"] == [1, 2]


def test_cover_letter_quality_rejects_direct_connection_noun_and_verb_bridges():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][1]["text"] = rewrite["cover_letter_paragraphs"][1]["text"].replace(
        "Northstar calls for developing agent-assisted workflows in shared services.",
        "That work is a direct connection to Northstar's agent-workflow responsibility.",
    )
    rewrite["cover_letter_paragraphs"][2]["text"] = rewrite["cover_letter_paragraphs"][2]["text"].replace(
        "Northstar names model deployment as an AI Engineer responsibility.",
        "Deploying it connects with Northstar's model-service responsibility.",
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["boilerplate_bridge_paragraph_indices"] == [1, 2]


def test_cover_letter_quality_catches_candidate_work_matches_role_formula():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][1]["text"] = rewrite["cover_letter_paragraphs"][1]["text"].replace(
        "Northstar calls for developing agent-assisted workflows in shared services.",
        "Ayni work matches Northstar's agent-workflow needs in shared services.",
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["boilerplate_bridge_paragraph_indices"] == [1]
    assert quality["revision_reasons"] == {"1": "generic_relevance_bridge"}


def test_cover_letter_quality_allows_a_single_semicolon():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][1]["text"] = rewrite["cover_letter_paragraphs"][1]["text"].replace(
        "I contributed to workflow design in Python services, including model routing",
        "I contributed to workflow design in Python services; I also worked on model routing",
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "passed"
    assert quality["stacked_detail_paragraph_indices"] == []
    assert quality["revision_paragraph_indices"] == []


def test_cover_letter_quality_revises_multiple_semicolon_stacked_details():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][1]["text"] = rewrite["cover_letter_paragraphs"][1]["text"].replace(
        "I contributed to workflow design in Python services, including model routing",
        "I contributed to workflow design in Python services; I also worked on model routing; "
        "I tested recovery behavior",
    )

    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())

    assert quality["status"] == "failed"
    assert quality["stacked_detail_paragraph_indices"] == [1]
    assert quality["revision_paragraph_indices"] == [1]
    assert quality["revision_reasons"] == {"1": "stacked_detail"}


def test_letter_revision_preserves_original_sources_and_filters_unapproved_ids():
    paragraphs = _complete_quality_letter()["cover_letter_paragraphs"]
    original = paragraphs[1]
    repair = {
        "revisions": [
            {
                **original,
                "paragraph_index": 1,
                "text": "I revised the project example while retaining its source evidence.",
                "claim_ids": ["fabricated-claim"],
                "cv_line_ids": ["L0001"],
                "requirement_ids": ["fabricated-requirement"],
                "research_finding_ids": ["RF99"],
                "source_urls": ["https://untrusted.example/page"],
                "preference_source_ids": ["unapproved-style-source"],
            }
        ],
        "unresolved_questions": [],
    }
    context = {
        "preferred_technical_profile_claim_ids": ["claim-ayni", "claim-registration"],
        "paragraphs_for_repair": [
            {
                "paragraph_index": 1,
                "allowed_claim_ids": ["claim-ayni"],
                "allowed_cv_line_ids": ["L0001"],
                "allowed_requirement_ids": ["REQ-02"],
                "allowed_research_finding_ids": [],
                "allowed_source_urls": [JOB_URL],
                "allowed_preference_source_ids": [],
            }
        ],
    }

    revised, events = _merge_letter_revisions(
        paragraphs,
        _supported_letter_grounding(),
        repair,
        [1],
        repair_context=context,
        repair_reasons={1: "jev_unsupported"},
    )

    assert revised[1]["text"] == repair["revisions"][0]["text"]
    assert revised[1]["claim_ids"] == ["claim-ayni"]
    assert revised[1]["cv_line_ids"] == ["L0001"]
    assert revised[1]["requirement_ids"] == ["REQ-02"]
    assert revised[1]["research_finding_ids"] == []
    assert revised[1]["source_urls"] == [JOB_URL]
    assert revised[1]["preference_source_ids"] == []
    assert events == [
        {
            "paragraph_index": 1,
            "fields_preserved_or_filtered": [
                "claim_ids", "requirement_ids", "research_finding_ids",
                "source_urls", "preference_source_ids",
            ],
        }
    ]


def test_letter_revision_fails_closed_when_a_required_paragraph_is_omitted():
    paragraphs = _complete_quality_letter()["cover_letter_paragraphs"]

    with pytest.raises(OpenAIProviderError, match="omitted required paragraph"):
        _merge_letter_revisions(
            paragraphs,
            _supported_letter_grounding(),
            {"revisions": [], "unresolved_questions": []},
            [1, 2],
            repair_context={"paragraphs_for_repair": []},
            repair_reasons={1: "jev_unsupported", 2: "jev_unsupported"},
        )


def test_cover_letter_contact_header_uses_only_contact_lines_before_cv_sections():
    lines = [
        {"text": "Alex Rivera"},
        {"text": "Riverdale, US | +1 555 123 4567 | alex@example.test"},
        {"text": "github.com/alexrivera | US citizen | English and Spanish (fluent)"},
        {"text": "Experience"},
        {"text": "Built a reporting service handling 1234567890 records."},
    ]

    assert _resume_signature_name(lines) == "Alex Rivera"
    assert _resume_contact_lines(lines) == [
        "Riverdale, US | +1 555 123 4567 | alex@example.test",
        "github.com/alexrivera",
    ]


def test_resume_artifact_bullets_keep_wrapped_lines_inside_hanging_indent(settings, database):
    artifacts = _render_artifacts(
        settings,
        "synthetic-request",
        {},
        [
            {
                "id": "L0001",
                "text": (
                    "- Built a small service with a deliberately long bullet that describes validation "
                    "across local integrations, process recovery, documented deployment checks, and "
                    "measured reductions in processing time for partner-facing workflows."
                ),
            }
        ],
        {
            "resume_line_order": ["L0001"],
            "resume_bullet_edits": [],
            "cover_letter_paragraphs": [],
        },
    )

    document = Document(artifacts[0]["file_path"])
    bullet = next(paragraph for paragraph in document.paragraphs if paragraph.style.name == "List Bullet")

    assert bullet.paragraph_format.left_indent.inches == pytest.approx(0.28, abs=0.001)
    assert bullet.paragraph_format.first_line_indent.inches == pytest.approx(-0.18, abs=0.001)


def test_rewriter_preserve_mode_omits_redundant_line_order_and_bounds_edits():
    prompt = " ".join(ROLE_PROMPTS["rewriter"].split())
    rewrite = {
        "selected_cv_id": "cv1",
        "resume_line_order": [],
        "resume_bullet_edits": [],
        "cover_letter_paragraphs": [],
        "application_answers": [],
        "outreach_drafts": [],
    }

    _validate_rewrite(
        rewrite,
        "cv1",
        {"L0001": "Experience", "L0002": "Built a Python service."},
        set(),
        {"findings": [], "contacts": []},
        [],
        "preserve",
    )

    assert rewrite["resume_line_order"] == ["L0001", "L0002"]
    assert "return an empty `resume_line_order`" in prompt
    assert "at most eight high-value" in prompt
    assert OUTPUT_SCHEMAS["rewriter"]["properties"]["resume_bullet_edits"]["maxItems"] == 8
    assert MAX_OUTPUT_TOKENS["rewriter"] == 20_000
    assert MAX_OUTPUT_TOKENS["letter_reviser"] == 12_000
    paragraph_schema = OUTPUT_SCHEMAS["rewriter"]["properties"]["cover_letter_paragraphs"]
    paragraph_schema = paragraph_schema["items"]
    assert "section" in paragraph_schema["required"]
    assert paragraph_schema["properties"]["section"] == {
        "type": ["string", "null"],
        "enum": ["opening", "evidence", "contribution", "closing", None],
    }
    revision_schema = OUTPUT_SCHEMAS["letter_reviser"]["properties"]["revisions"]["items"]
    assert "section" in revision_schema["required"]
    assert revision_schema["properties"]["section"]["type"] == ["string", "null"]

    def assert_strict_objects_have_required_properties(schema):
        if schema.get("type") == "object":
            properties = schema.get("properties", {})
            assert set(schema.get("required", [])) == set(properties)
            assert schema.get("additionalProperties") is False
            for child in properties.values():
                assert_strict_objects_have_required_properties(child)
        if schema.get("type") == "array":
            assert_strict_objects_have_required_properties(schema["items"])

    assert_strict_objects_have_required_properties(OUTPUT_SCHEMAS["rewriter"])
    assert_strict_objects_have_required_properties(OUTPUT_SCHEMAS["letter_reviser"])
    assert "research_finding_ids" in paragraph_schema["required"]
    assert paragraph_schema["properties"]["research_finding_ids"]["maxItems"] == 2
    assert "requirement_ids" in paragraph_schema["required"]
    assert OUTPUT_SCHEMAS["rewriter"]["properties"]["cover_letter_paragraphs"]["maxItems"] == 8
    assert PROMPT_VERSION == "2026-10-12.3"
    assert OUTPUT_SCHEMA_VERSION == "application-output-v19"
    assert "unsupported equivalence claims" in prompt
    assert "A project or technology name alone does not establish" in prompt
    assert "one well-developed example can be enough" in " ".join(prompt.split())
    assert "supports another important responsibility" in prompt
    assert "Make one bounded repair" in " ".join(ROLE_PROMPTS["letter_reviser"].split())
    assert "multiple stacked semicolon details" in " ".join(ROLE_PROMPTS["letter_reviser"].split())
    assert "Apply every item in `cover_letter_quality_issues`" in " ".join(ROLE_PROMPTS["letter_reviser"].split())
    assert "There is no required word count" in " ".join(ROLE_PROMPTS["rewriter"].split())
    assert "Use `section` only as optional internal metadata" in " ".join(ROLE_PROMPTS["rewriter"].split())
    assert "one page is a practical norm" in " ".join(ROLE_PROMPTS["letter_reviewer"].split())
    assert "Do not summarize or inventory several job requirements" in " ".join(ROLE_PROMPTS["rewriter"].split())
    assert "remove the inventory or preview and keep only supported application context" in " ".join(ROLE_PROMPTS["letter_reviser"].split())
    assert MAX_OUTPUT_TOKENS["letter_reviewer"] == 8_000
    assert "public_email_verified=true" in " ".join(ROLE_PROMPTS["rewriter"].split())
    assert "70–120 words" in " ".join(ROLE_PROMPTS["rewriter"].split())
    assert "avoid generic bridges" in " ".join(ROLE_PROMPTS["rewriter"].split()).casefold()
    assert "replacement characters" in " ".join(ROLE_PROMPTS["rewriter"].split())
    assert "do not weaken \"built\" to \"worked on\"" in " ".join(ROLE_PROMPTS["rewriter"].split())
    assert "For separate examples, vary the sentence structure" in " ".join(ROLE_PROMPTS["rewriter"].split())
    assert "Do not print candidate project brands or repository names" in " ".join(ROLE_PROMPTS["rewriter"].split())
    assert "Do not write \"the role includes\"" in " ".join(ROLE_PROMPTS["rewriter"].split())
    assert "Do not add a generic preview such as \"my relevant project work has covered" in " ".join(ROLE_PROMPTS["rewriter"].split())
    assert "Do not treat a generic opening that previews later examples" in " ".join(ROLE_PROMPTS["letter_reviewer"].split())
    assert "strict fact boundary" in " ".join(ROLE_PROMPTS["rewriter"].split())
    assert "strict fact boundary" in " ".join(ROLE_PROMPTS["letter_reviser"].split())
    rewriter_prompt = " ".join(ROLE_PROMPTS["rewriter"].split())
    assert "When an exact owner-attributed profile claim adds a material role-relevant" in rewriter_prompt
    assert "consider incorporating it into that bullet" in rewriter_prompt
    assert "Do not summarize or inventory several job requirements" in rewriter_prompt
    assert "a few abbreviated notes as a completed letter" in rewriter_prompt
    assert "what this experience prepares the candidate to contribute" in rewriter_prompt
    assert "integrate it naturally and avoid repeating the same action" in rewriter_prompt
    assert "A factual project sentence followed by a job-duty sentence is only an outline" in rewriter_prompt


def test_invalid_optional_outreach_is_omitted_without_blocking_letter():
    rewrite = {
        "selected_cv_id": "cv-synthetic",
        "resume_line_order": ["L0001"],
        "resume_bullet_edits": [],
        "cover_letter_paragraphs": [
            {
                "section": "evidence",
                "text": "I built a Python service for Example Labs.",
                "claim_ids": [],
                "cv_line_ids": ["L0001"],
                "requirement_ids": [],
                "research_finding_ids": [],
                "source_urls": [JOB_URL],
                "preference_source_ids": [],
            }
        ],
        "application_answers": [],
        "outreach_drafts": [
            {
                "contact_source_url": "https://example.org/team",
                "recipient_name": "Recruitment team",
                "subject": "Question about the role",
                "body": "Could you share more about the role?",
                "claim_ids": [],
                "cv_line_ids": [],
                "source_urls": ["https://example.org/team"],
            }
        ],
        "unresolved_questions": [],
    }

    _validate_rewrite(
        rewrite,
        "cv-synthetic",
        {"L0001": "Built a Python service."},
        set(),
        {
            "findings": [],
            "contacts": [{
                "source_url": "https://example.org/team",
                "name": "Recruitment team",
                "public_email": "recruiting@example.org",
                "suppressed": False,
            }],
        },
        [],
        "preserve",
        job_source_url=JOB_URL,
    )

    assert len(rewrite["cover_letter_paragraphs"]) == 1
    assert rewrite["outreach_drafts"] == []
    assert rewrite["unresolved_questions"] == [
        "Clue omitted outreach draft 1 because it has no approved candidate evidence.",
        "A verified public recruitment address was available, but no supported outreach draft was generated.",
    ]


def test_invalid_optional_resume_edit_is_omitted_without_blocking_letter():
    rewrite = {
        "selected_cv_id": "cv-synthetic",
        "resume_line_order": ["L0001"],
        "resume_bullet_edits": [{
            "line_id": "L9999",
            "original_text": "A line from another CV.",
            "revised_text": "An unsupported rewrite.",
            "claim_ids": ["claim-synthetic"],
            "cv_line_ids": ["L0001"],
        }],
        "cover_letter_paragraphs": [{
            "section": "evidence",
            "text": "I built a Python service for the role's API work.",
            "claim_ids": ["claim-synthetic"],
            "cv_line_ids": [],
            "requirement_ids": [],
            "research_finding_ids": [],
            "source_urls": [JOB_URL],
            "preference_source_ids": [],
        }],
        "application_answers": [],
        "outreach_drafts": [],
        "unresolved_questions": [],
    }

    _validate_rewrite(
        rewrite,
        "cv-synthetic",
        {"L0001": "Built a Python service."},
        {"claim-synthetic"},
        {"findings": [], "contacts": []},
        [],
        "preserve",
        job_source_url=JOB_URL,
    )

    assert rewrite["resume_bullet_edits"] == []
    assert len(rewrite["cover_letter_paragraphs"]) == 1
    assert rewrite["unresolved_questions"] == [
        "Clue omitted 1 optional resume edit(s) because the exact selected CV source line was not verified; the original line was retained."
    ]


def test_supported_outreach_draft_requires_and_accepts_verified_contact_source():
    contact_url = "https://example.org/team"
    rewrite = {
        "selected_cv_id": "cv-synthetic",
        "resume_line_order": ["L0001"],
        "resume_bullet_edits": [],
        "cover_letter_paragraphs": [],
        "application_answers": [],
        "outreach_drafts": [{
            "contact_source_url": contact_url,
            "recipient_name": "Recruitment team",
            "subject": "Question about AI Engineer",
            "body": (
                "Hello Recruitment Team, I am writing about the AI Engineer opening and its work on model deployment. "
                "I built a Python retrieval service and am interested in how the team evaluates services used by researchers. "
                "If you are able to share guidance, I would appreciate learning which technical priorities are most important for this role. "
                "A short pointer to the team's public information would be helpful as I prepare my application. "
                "Thank you for your time."
            ),
            "claim_ids": [],
            "cv_line_ids": ["L0001"],
            "source_urls": [contact_url, JOB_URL],
        }],
        "unresolved_questions": [],
    }

    _validate_rewrite(
        rewrite,
        "cv-synthetic",
        {"L0001": "Built a Python retrieval service."},
        set(),
        {
            "findings": [],
            "contacts": [{
                "source_url": contact_url,
                "name": "Recruitment team",
                "public_email": "recruiting@example.org",
                "suppressed": False,
            }],
        },
        [],
        "preserve",
        job_source_url=JOB_URL,
    )

    assert len(rewrite["outreach_drafts"]) == 1
    assert rewrite["outreach_drafts"][0]["recipient_email"] == "recruiting@example.org"
    assert rewrite["unresolved_questions"] == []


def test_outreach_draft_rejects_replacement_characters_short_bodies_and_body_addresses():
    contact_url = "https://example.org/team"
    base = {
        "selected_cv_id": "cv-synthetic",
        "resume_line_order": ["L0001"],
        "resume_bullet_edits": [],
        "cover_letter_paragraphs": [],
        "application_answers": [],
        "unresolved_questions": [],
    }
    valid_body = (
        "Hello Recruitment Team, I am writing about the AI Engineer opening and its work on model deployment. "
        "I built a Python retrieval service and am interested in how the team evaluates services used by researchers. "
        "If you are able to share guidance, I would appreciate learning which technical priorities are most important for this role. "
        "A short pointer to the team's public information would be helpful as I prepare my application. Thank you for your time."
    )
    invalid_drafts = [
        {"subject": "Question � about the AI Engineer role", "body": valid_body},
        {"subject": "Question about the AI Engineer role", "body": "Could you share more about the role?"},
        {"subject": "Question about the AI Engineer role", "body": valid_body + " Please email me at juan@example.com."},
    ]
    expected_reasons = [
        "replacement character",
        "70-to-120-word range",
        "email address",
    ]

    for draft, expected_reason in zip(invalid_drafts, expected_reasons, strict=True):
        rewrite = {
            **base,
            "outreach_drafts": [{
                "contact_source_url": contact_url,
                "recipient_name": "Recruitment team",
                **draft,
                "claim_ids": [],
                "cv_line_ids": ["L0001"],
                "source_urls": [contact_url, JOB_URL],
            }],
        }
        _validate_rewrite(
            rewrite,
            "cv-synthetic",
            {"L0001": "Built a Python retrieval service."},
            set(),
            {
                "findings": [],
                "contacts": [{
                    "source_url": contact_url,
                    "name": "Recruitment team",
                    "public_email": "recruiting@example.org",
                    "suppressed": False,
                }],
            },
            [],
            "preserve",
            job_source_url=JOB_URL,
        )

        assert rewrite["outreach_drafts"] == []
        assert any(expected_reason in question for question in rewrite["unresolved_questions"])


def test_resume_document_heading_detection_covers_common_section_and_project_titles():
    assert _looks_like_heading("Professional Experience")
    assert _looks_like_heading("Selected AI and ML Projects")
    assert _looks_like_subheading("Transformer Model | Independent project")
    assert not _looks_like_subheading("name@example.org | +39 123 456")


def test_resume_reflow_repairs_pdf_wraps_bullets_and_inline_skill_categories():
    lines = [
        {"id": "L0001", "text": "JUAN MARTIN SANCHEZ-BARDELLINI"},
        {"id": "L0002", "text": "AI/ML Engineer"},
        {"id": "L0003", "text": "Professional Profile"},
        {"id": "L0004", "text": "Engineer using simulation      -generated systems across"},
        {"id": "L0005", "text": "large-scale test cases."},
        {"id": "L0006", "text": "Technical Skills"},
        {
            "id": "L0007",
            "text": "APIs and applications: FastAPI, React Databases: PostgreSQL, SQLite",
        },
        {"id": "L0008", "text": "\ufffd Built a retrieval service with routing"},
        {"id": "L0009", "text": "and continued retrieval behavior."},
    ]
    rewrite = {
        "resume_line_order": [item["id"] for item in lines],
        "resume_bullet_edits": [
            {"line_id": "L0008", "revised_text": "Built a retrieval service with routing"}
        ],
    }

    blocks = _resume_render_blocks(lines, rewrite)

    assert [block["type"] for block in blocks] == [
        "header", "header", "heading", "plain", "heading", "skill", "skill", "bullet"
    ]
    assert blocks[3]["text"] == "Engineer using simulation-generated systems across large-scale test cases."
    assert blocks[5]["text"] == "APIs and applications: FastAPI, React"
    assert blocks[6]["text"] == "Databases: PostgreSQL, SQLite"
    assert blocks[7]["text"] == "Built a retrieval service with routing and continued retrieval behavior."
    assert blocks[7]["type"] == "bullet"
    assert lines[3]["text"] == "Engineer using simulation      -generated systems across"


def test_descriptive_profile_supports_only_close_cover_letter_preferences():
    source_id = "descriptive-profile-1"
    profile = {
        "source_id": source_id,
        "source_type": "descriptive_profile",
        "authorship_label": "owner_written",
        "text": "I prefer direct, modest writing and collaborative technical teams.",
    }
    rewrite = {
        "selected_cv_id": "cv-synthetic",
        "resume_line_order": ["L0001"],
        "resume_bullet_edits": [
            {
                "line_id": "L0001",
                "original_text": "Experience",
                "revised_text": "Built reliable Python services.",
                "claim_ids": ["claim-1"],
                "cv_line_ids": ["L0001"],
            }
        ],
        "cover_letter_paragraphs": [
            {
                "text": "I prefer direct, modest writing and collaborative technical teams.",
                "claim_ids": [],
                "research_finding_ids": [],
                "source_urls": [],
                "preference_source_ids": [source_id],
            },
            {
                "text": "I value careful evaluation, useful documentation, and collaborative technical work.",
                "claim_ids": [],
                "research_finding_ids": [],
                "source_urls": [],
                "preference_source_ids": [source_id],
            },
        ],
        "application_answers": [],
        "outreach_drafts": [],
    }
    _validate_rewrite(
        rewrite,
        "cv-synthetic",
        {"L0001": "Experience"},
        {"claim-1"},
        {"findings": [], "contacts": []},
        [],
        "preserve",
        permitted_preference_source_ids={source_id},
    )

    assertions = _grounding_assertions(
        rewrite,
        [{"id": "claim-1", "claim_text": "Built a Python service", "evidence_excerpt": "Service implementation."}],
        {"findings": [], "contacts": []},
        descriptive_profiles=[profile],
        cv_lines={"L0001": "Experience"},
        cv_source_id="cv-synthetic",
    )

    assert [item["source_type"] for item in assertions[0]["evidence"]] == [
        "owner_provided_resume",
        "approved_claim",
    ]
    assert assertions[1]["evidence"] == [
        {
            "source_type": "owner_stated_preference",
            "source_id": source_id,
            "text": "Owner-stated preference; not career evidence: " + profile["text"],
        }
    ]
    assert assertions[2]["evidence"] == []  # Loose topic overlap cannot support extra preferences.

    inferred_profile = {**profile, "authorship_label": "ai_assisted"}
    inferred_assertions = _grounding_assertions(
        rewrite,
        [{"id": "claim-1", "claim_text": "Built a Python service", "evidence_excerpt": "Service implementation."}],
        {"findings": [], "contacts": []},
        descriptive_profiles=[inferred_profile],
    )
    assert inferred_assertions[1]["evidence"] == []
    with pytest.raises(OpenAIProviderError, match="unapproved writing preference source"):
        _validate_rewrite(
            rewrite,
            "cv-synthetic",
            {"L0001": "Experience"},
            {"claim-1"},
            {"findings": [], "contacts": []},
            [],
            "preserve",
            permitted_preference_source_ids=set(),
        )

    unapproved = {**rewrite, "cover_letter_paragraphs": [{
        **rewrite["cover_letter_paragraphs"][0],
        "preference_source_ids": ["writing-sample-1"],
    }]}
    with pytest.raises(OpenAIProviderError, match="unapproved writing preference source"):
        _validate_rewrite(
            unapproved,
            "cv-synthetic",
            {"L0001": "Experience"},
            {"claim-1"},
            {"findings": [], "contacts": []},
            [],
            "preserve",
            permitted_preference_source_ids={source_id},
        )


def test_researcher_prompt_requires_useful_facts_even_without_public_contacts():
    prompt = " ".join(ROLE_PROMPTS["researcher"].split())

    assert "return at most eight distinct concise atomic findings" in prompt
    assert MAX_OUTPUT_TOKENS["researcher"] == 4_000
    assert "A lack of contact details does not make useful company or role facts irrelevant" in prompt
    assert "Never guess a contact, email address, or relationship" in prompt
    assert "organization-published general recruitment inbox" in prompt
    assert "do not imply that it belongs to a named person" in prompt
    contact_types = OUTPUT_SCHEMAS["researcher"]["properties"]["contacts"]["items"]["properties"]["contact_type"]["enum"]
    assert "general_recruitment_inbox" in contact_types


def test_cover_letter_prompt_and_review_use_research_based_flexible_quality_criteria():
    prompt = " ".join(ROLE_PROMPTS["rewriter"].split())
    reviser_prompt = " ".join(ROLE_PROMPTS["letter_reviser"].split())
    reviewer_prompt = " ".join(ROLE_PROMPTS["letter_reviewer"].split())
    recruiter_prompt = " ".join(ROLE_PROMPTS["recruiter"].split())

    assert "complete, role-specific application letter, not a brief note" in prompt
    assert "These are content functions, not mandatory separate sections" in prompt
    assert "Keep the owner's direct, modest voice" in prompt
    assert "Keep the owner's direct, modest voice" in reviser_prompt
    assert "use this as a practical guide, not a quota" in prompt
    assert "one well-developed example can be enough" in prompt
    assert "complementary evidence can make a stronger case" in prompt
    assert "State a personal reason, enthusiasm, or employer preference only when" in prompt
    assert "strict fact boundary" in prompt
    assert "reproduce its `original_text` exactly" in prompt
    assert "requirement IDs are provenance metadata and do not count as a written connection" in prompt
    assert "do not wrap it in phrases such as" in prompt
    assert "Jev checks each complete paragraph" in prompt
    assert "direct, modest voice" in prompt
    assert "There is no fixed word count, paragraph count, section order, or sentence count" in reviser_prompt
    assert "When Jev removes a whole example" in reviser_prompt
    assert "replace the paragraph with a directly supported example when one is available" in reviser_prompt
    assert "Preserve the original citation and source ID arrays exactly" in reviser_prompt
    assert "Do not require a specific number of words, paragraphs, or labeled sections" in reviewer_prompt
    assert "professional conclusion" in reviewer_prompt
    assert "the title, salutation, and signature do not do this work" in prompt
    assert "one project fact plus one job duty" in reviser_prompt
    assert "A contribution sentence plus a parallel job-duty sentence is an outline" in reviewer_prompt
    assert prompt.count("Connect each candidate example") <= 1
    assert "Do not print candidate project brands or repository names" in prompt
    assert "A focused body can be complete when it identifies the application" in reviewer_prompt
    assert "body much shorter than the ordinary one-page norm deserves close scrutiny" in reviewer_prompt
    assert "do not count as body substance" in reviewer_prompt
    assert "A short letter passes only when its prose still makes a developed, role-specific case" in reviewer_prompt
    assert "A 3 is an affirmative quality judgment" in reviewer_prompt
    assert "few compressed project statements do not become a complete letter" in reviewer_prompt
    assert "Recommend `ready` only when every dimension scores at least 3" in reviewer_prompt
    assert "selected_profile_evidence_for_review" in reviewer_prompt
    assert "material_gap" in reviewer_prompt
    assert "Do not require every selected claim to appear" in reviewer_prompt
    assert "Do not approve a short body only because it has an application title" in reviewer_prompt
    assert "`paragraph_index` equal to the current body paragraph count" in reviser_prompt
    assert "cite at least one exact selected-CV line ID" in recruiter_prompt
    assert "a profile-only claim cannot establish CV coverage" in recruiter_prompt

    paragraph_schema = OUTPUT_SCHEMAS["rewriter"]["properties"]["cover_letter_paragraphs"]
    assert paragraph_schema["maxItems"] == 8
    assert "minItems" not in paragraph_schema
    assert paragraph_schema["items"]["properties"]["section"]["enum"] == [
        "opening", "evidence", "contribution", "closing", None
    ]
    revision_schema = OUTPUT_SCHEMAS["letter_reviser"]["properties"]["revisions"]
    assert revision_schema["maxItems"] == 8
    assert "minItems" not in revision_schema
    assert revision_schema["items"]["properties"]["paragraph_index"]["maximum"] == 7
    review_schema = OUTPUT_SCHEMAS["letter_reviewer"]
    assert review_schema["properties"]["dimensions"]["minItems"] == 5
    assert review_schema["properties"]["dimensions"]["maxItems"] == 5
    assert review_schema["properties"]["issues"]["maxItems"] == 8
    assert review_schema["properties"]["profile_evidence_assessment"]["maxItems"] == 4
    assert set(review_schema["required"]) == {
        "decision", "dimensions", "issues", "profile_evidence_assessment"
    }

def test_diagnoser_prompt_requires_verbatim_source_and_line_ids():
    prompt = " ".join(ROLE_PROMPTS["diagnoser"].split())

    assert "Copy the exact source ID and line ID" in prompt
    assert "never invent, infer, or combine IDs" in prompt
    assert "eight highest-priority concrete issues" in prompt
    assert OUTPUT_SCHEMAS["diagnoser"]["properties"]["diagnostics"]["maxItems"] == 8
    assert MAX_OUTPUT_TOKENS["diagnoser"] == 2_800


def test_recruiter_prompt_is_concise_and_stage_budget_covers_measured_incomplete_response():
    prompt = " ".join(ROLE_PROMPTS["recruiter"].split())

    assert "one short evidence note per criterion" in prompt
    assert MAX_OUTPUT_TOKENS["recruiter"] == 9_000


def test_diagnoser_filter_drops_unknown_or_mismatched_line_references():
    filtered = _filter_diagnoser_references(
        {
            "overall_note": "One visible heading issue.",
            "diagnostics": [
                {"source_id": "cv-1", "line_id": "L0001", "issue": "Heading is split."},
                {"source_id": "invented", "line_id": "L0001", "issue": "Wrong source."},
                {"source_id": "cv-1", "line_id": "L0099", "issue": "Unknown line."},
                "malformed",
            ],
        },
        source_id="cv-1",
        valid_line_ids={"L0001", "L0002"},
    )

    assert filtered["diagnostics"] == [
        {"source_id": "cv-1", "line_id": "L0001", "issue": "Heading is split."}
    ]
    assert filtered["rejected_invalid_references"] == 3


def test_hiring_manager_prompt_requires_verbatim_question_and_answer_copy():
    prompt = " ".join(ROLE_PROMPTS["hiring_manager"].split())

    assert "copy the supplied question list in its exact order" in prompt
    assert "copy each matching question and answer verbatim" in prompt
    assert "Return one assessment per supplied answer" in prompt


def test_hiring_manager_practice_is_separate_and_has_no_tools(settings, database):
    record, _cv, _technical, _descriptive, _sample, _claim_id, _text = _create_preparation(
        settings, database
    )
    enabled = _enable_synthetic_openai(settings, database)
    FakePreparationClient.instances.clear()
    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )
    session = create_practice_session(database, record["id"])
    run_interview_practice_questions(
        database, enabled, session["id"], client_factory=FakePreparationClient
    )
    with connect(database) as db:
        question_row = db.execute(
            "SELECT * FROM practice_sessions WHERE id = ?", (session["id"],)
        ).fetchone()
    questions = json.loads(question_row["output_json"])["questions"]
    assert question_row["state"] == "questions"
    queue_practice_answers(
        database,
        session["id"],
        json.dumps(
            [{"question": question, "answer": "I would measure latency and error rates."} for question in questions]
        ),
    )
    run_interview_practice_assessment(
        database, enabled, session["id"], client_factory=FakePreparationClient
    )
    with connect(database) as db:
        assessed = db.execute(
            "SELECT * FROM practice_sessions WHERE id = ?", (session["id"],)
        ).fetchone()
    assert assessed["state"] == "complete"
    assert len(json.loads(assessed["output_json"])["answer_assessments"]) == 3
    practice_calls = [
        call
        for instance in FakePreparationClient.instances
        for call in instance.calls
        if call["text"]["format"]["name"] == "clue_hiring_manager_v3"
    ]
    assert len(practice_calls) == 2
    assert all("tools" not in call for call in practice_calls)


def test_owner_reported_outcomes_and_receipts_remain_distinct(settings, database):
    record, _cv, _technical, _descriptive, _sample, _claim_id, _text = _create_preparation(
        settings, database
    )
    enabled = _enable_synthetic_openai(settings, database)
    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )
    with connect(database) as db:
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]

    client = TestClient(create_app(enabled), base_url="http://127.0.0.1")
    approved = client.post(
        f"/preparations/{record['id']}/packets/{packet_id}/approve",
        data={"confirm_review": "on"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert approved.status_code == 303
    submitted = client.post(
        f"/preparations/{record['id']}/record-stage",
        data={"stage": "submitted", "owner_attestation": "on"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert submitted.status_code == 303
    with connect(database) as db:
        owner_event = db.execute(
            "SELECT * FROM application_events WHERE event_type = 'owner_submission_attestation'"
        ).fetchone()
        assert owner_event["evidence_level"] == "owner_reported"
        assert db.execute("SELECT COUNT(*) FROM applications").fetchone()[0] == 1

    feedback = client.post(
        f"/preparations/{record['id']}/interview-feedback",
        data={
            "stage": "technical",
            "self_assessment": "I explained the retry behavior clearly.",
            "employer_feedback": "Strong debugging example.",
            "gap_tags": "system design, tradeoffs",
        },
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert feedback.status_code == 303
    receipt = client.post(
        f"/preparations/{record['id']}/receipt",
        files={"receipt_file": ("confirmation.txt", b"Synthetic receipt", "text/plain")},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert receipt.status_code == 303
    with connect(database) as db:
        receipt_event = db.execute(
            "SELECT * FROM application_events WHERE event_type = 'receipt_imported'"
        ).fetchone()
        interview = db.execute("SELECT * FROM interview_feedback").fetchone()
    assert receipt_event["evidence_level"] == "receipt_imported"
    assert interview["self_assessment"] == "I explained the retry behavior clearly."
    assert interview["employer_feedback"] == "Strong debugging example."
    outcomes = client.get("/outcomes")
    assert outcomes.status_code == 200
    assert "Owner Reported" in outcomes.text
    assert "Receipt Imported" in outcomes.text


def test_owner_followup_reminders_are_visible_and_cancel_on_reply_or_outcome(settings, database):
    record, _cv, _technical, _descriptive, _sample, _claim_id, _text = _create_preparation(
        settings, database
    )
    enabled = _enable_synthetic_openai(settings, database)
    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )
    with connect(database) as db:
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]
    client = TestClient(create_app(enabled), base_url="http://127.0.0.1")
    approved = client.post(
        f"/preparations/{record['id']}/packets/{packet_id}/approve",
        data={"confirm_review": "on"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert approved.status_code == 303

    current_day = datetime.now().astimezone().date()
    due_on = (current_day + timedelta(days=1)).isoformat()
    scheduled = client.post(
        f"/preparations/{record['id']}/follow-up",
        data={"kind": "outreach", "due_on": due_on, "note": "Check for a reply"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert scheduled.status_code == 303
    listing = client.get("/preparations")
    assert "Follow-up reminders" in listing.text
    assert "Check for a reply" in listing.text
    assert "This is an in-app reminder only" in listing.text
    assert list_followups(database, today=current_day + timedelta(days=1))[0]["due"]

    replied = client.post(
        f"/preparations/{record['id']}/follow-up/finish",
        data={"outcome": "reply"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert replied.status_code == 303
    with connect(database) as db:
        reminder = db.execute(
            "SELECT state, resolution FROM application_followups WHERE request_id = ?",
            (record["id"],),
        ).fetchone()
    assert reminder["state"] == "cancelled"
    assert "reply" in reminder["resolution"]

    schedule_followup(
        database, record["id"], kind="application", due_on=due_on, today=current_day
    )
    submitted = client.post(
        f"/preparations/{record['id']}/record-stage",
        data={"stage": "submitted", "owner_attestation": "on"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert submitted.status_code == 303
    with connect(database) as db:
        assert db.execute(
            "SELECT state FROM application_followups WHERE request_id = ? AND state = 'scheduled'",
            (record["id"],),
        ).fetchone()
    response = client.post(
        f"/preparations/{record['id']}/record-stage",
        data={"stage": "screen"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert response.status_code == 303
    with connect(database) as db:
        assert db.execute(
            "SELECT state FROM application_followups WHERE request_id = ? ORDER BY updated_at DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["state"] == "cancelled"

    schedule_followup(
        database, record["id"], kind="outreach", due_on=due_on, today=current_day
    )
    with connect(database) as db:
        db.execute("UPDATE jobs SET is_active = 0 WHERE id = ?", (record["job_id"],))
    client.get("/preparations")
    with connect(database) as db:
        assert db.execute(
            "SELECT state FROM application_followups WHERE request_id = ? ORDER BY updated_at DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["state"] == "cancelled"


def test_outcomes_exclude_pending_from_resolved_denominator_and_report_packet_effort(settings, database):
    record, _cv, _technical, _descriptive, _sample, _claim_id, _text = _create_preparation(
        settings, database
    )
    enabled = _enable_synthetic_openai(settings, database)
    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )
    with connect(database) as db:
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]
    client = TestClient(create_app(enabled), base_url="http://127.0.0.1")
    approved = client.post(
        f"/preparations/{record['id']}/packets/{packet_id}/approve",
        data={"confirm_review": "on"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert approved.status_code == 303
    submitted = client.post(
        f"/preparations/{record['id']}/record-stage",
        data={"stage": "submitted", "owner_attestation": "on"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert submitted.status_code == 303
    pending_summary = client.get("/outcomes")
    assert "PENDING · EXCLUDED" in pending_summary.text
    assert "No tracked application has an explicitly resolved owner-reported outcome yet." in pending_summary.text

    feedback_response = client.post(
        f"/preparations/{record['id']}/packets/{packet_id}/feedback",
        data={
            "owner_minutes": "27",
            "quality_rating": "4",
            "factual_corrections": "2",
            "owner_note": "Two dates needed owner correction.",
        },
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert feedback_response.status_code == 303
    assert "feedback+saved" in feedback_response.headers["location"]
    rejected = client.post(
        f"/preparations/{record['id']}/record-stage",
        data={"stage": "rejected", "details": "Synthetic outcome"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert rejected.status_code == 303
    outcomes = client.get("/outcomes")
    assert "Resolved application outcomes" in outcomes.text
    assert "1 of 1 resolved applications" in outcomes.text
    assert "100.0% of resolved owner-reported outcomes" in outcomes.text
    assert "27 min" in outcomes.text
    assert "4/5" in outcomes.text
    assert "2 factual corrections" in outcomes.text

def test_contact_suppression_invalidates_packet_and_survives_manual_retry(settings, database):
    record, _cv, _technical, _descriptive, _sample, _claim, _text = _create_preparation(
        settings, database
    )
    enabled = _enable_synthetic_openai(settings, database)
    FakePreparationClient.instances.clear()
    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )
    with connect(database) as db:
        contact = db.execute(
            "SELECT id FROM researched_contacts WHERE request_id = ? LIMIT 1", (record["id"],)
        ).fetchone()
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]
    client = TestClient(create_app(enabled), base_url="http://127.0.0.1")
    response = client.post(
        f"/preparations/{record['id']}/contacts/{contact['id']}/suppress",
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert get_packet(database, packet_id)["status"] == "obsolete"
    with connect(database) as db:
        saved = db.execute(
            "SELECT suppressed FROM researched_contacts WHERE id = ?", (contact["id"],)
        ).fetchone()
        assert db.execute("SELECT COUNT(*) FROM contact_suppressions").fetchone()[0] == 1
        assert db.execute(
            "SELECT state FROM preparation_requests WHERE id = ?", (record["id"],)
        ).fetchone()["state"] == "blocked"
    assert saved["suppressed"] == 1
    with pytest.raises(ValueError, match="Approve the complete packet"):
        create_approved_packet_draft(
            database,
            enabled,
            record["id"],
            packet_id,
            0,
            confirmed=True,
            create_draft=lambda *_args: pytest.fail("Suppressed contact cannot be staged"),
        )

    retry, created = retry_preparation(database, record["id"])
    assert created is True
    run_preparation(
        database,
        enabled,
        retry["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )
    refreshed = get_preparation(database, retry["id"])
    assert refreshed["state"] == "review"
    with connect(database) as db:
        new_contact = db.execute(
            "SELECT suppressed FROM researched_contacts WHERE request_id = ?", (retry["id"],)
        ).fetchone()
        new_packet = db.execute(
            "SELECT output_json FROM preparation_packets WHERE request_id = ?", (retry["id"],)
        ).fetchone()
    assert new_contact["suppressed"] == 1
    output = json.loads(new_packet["output_json"])
    assert output["rewriter"]["outreach_drafts"] == []
    assert output["research"]["contacts"][0]["suppressed"] is True


def test_crawler_can_follow_exact_external_source_links_but_not_unlinked_domains(settings, monkeypatch):
    from clue_ai import scrapling_research

    fetched = []

    def fake_fetch(url, host, _settings):
        fetched.append((url, host))
        if host == "jobs.example.org":
            links = [
                {"label": "Company team", "url": "https://company.example.org/team"},
                {"label": "Unrelated", "url": "https://unrelated.example.net/profile"},
            ]
        else:
            links = [
                {"label": "More team info", "url": "https://company.example.org/about"},
                {"label": "Next host", "url": "https://third.example.net/next"},
            ]
        return {"url": url, "title": "Synthetic", "text": "Synthetic page", "links": links, "observed_at": "2026-10-06"}

    monkeypatch.setattr(scrapling_research, "_scrapling_fetch", fake_fetch)
    monkeypatch.setattr(scrapling_research, "_require_public_resolution", lambda _host: None)
    crawler = BoundedResearchCrawler(settings, [JOB_URL])
    crawler.crawl(JOB_URL, "research team")
    external = crawler.crawl("https://company.example.org/team", "find a public team lead")
    assert external["url"] == "https://company.example.org/team"
    assert crawler.crawl("https://company.example.org/about", "more team details")["url"].endswith("/about")
    with pytest.raises(ResearchCrawlError, match="not supplied or linked"):
        crawler.crawl("https://third.example.net/next", "follow a second external domain")
    with pytest.raises(ResearchCrawlError, match="not supplied or linked"):
        crawler.crawl("https://unlinked.example.net/person", "guess a contact")
    assert len(fetched) == 3


def test_research_crawler_does_not_request_a_host_again_after_access_denial(settings, monkeypatch):
    from clue_ai import scrapling_research

    requests = []
    blocked_host = (urlsplit(JOB_URL).hostname or "").casefold()

    def deny_once(url, host, _settings):
        requests.append((url, host))
        if host == blocked_host:
            raise scrapling_research.ResearchCrawlBlockedError(
                "The site denied the public-page request.", request_sent=True
            )
        return {
            "url": url,
            "title": "Alternate allowed source",
            "text": "Public information from another already allowed host.",
            "links": [],
            "form_fields": [],
            "form_fields_truncated": False,
            "observed_at": "2026-10-12",
        }

    monkeypatch.setattr(scrapling_research, "_scrapling_fetch", deny_once)
    monkeypatch.setattr(scrapling_research, "_require_public_resolution", lambda _host: None)
    same_host_url = JOB_URL.rstrip("/") + "/company"
    alternate_url = "https://company.example.org/jobs/role"
    crawler = BoundedResearchCrawler(settings, [JOB_URL, same_host_url, alternate_url])

    with pytest.raises(ResearchCrawlError, match="denied") as first_error:
        crawler.crawl(JOB_URL, "Find public role details")
    assert first_error.value.request_sent is True

    with pytest.raises(ResearchCrawlError, match="No request was sent") as skipped_error:
        crawler.crawl(same_host_url, "Try another allowed page on the same host")
    assert skipped_error.value.request_sent is False
    alternate_page = crawler.crawl(alternate_url, "Read another allowed source")
    assert alternate_page["title"] == "Alternate allowed source"
    assert len(requests) == 2


def test_public_form_parser_keeps_labels_and_choices_but_never_input_values():
    from clue_ai.scrapling_research import _PublicFormParser

    parser = _PublicFormParser()
    parser.feed(
        '<form><label for="motivation">Why this role?</label>'
        '<textarea id="motivation" name="motivation">PRIVATE ANSWER</textarea>'
        '<label for="country">Country</label><select id="country" required>'
        '<option value="it">Italy</option><option value="fr">France</option></select>'
        '<label for="email">Email</label><input id="email" type="email" value="PRIVATE EMAIL">'
        '<input type="hidden" value="SECRET TOKEN">'
        '<div hidden><label for="ghost">Hidden question</label>'
        '<input id="ghost" type="text"></div>'
        '<input id="css-hidden" style="display: none" type="text">'
        '</form>'
    )

    fields = parser.result()
    assert [(item["label"], item["control_type"]) for item in fields] == [
        ("Why this role?", "textarea"),
        ("Country", "select"),
        ("Email", "email"),
    ]
    assert fields[0]["is_open_question"] is True
    assert fields[1]["options"] == ["Italy", "France"]
    assert fields[1]["required"] is True
    assert "PRIVATE ANSWER" not in repr(fields)
    assert "PRIVATE EMAIL" not in repr(fields)
    assert "SECRET TOKEN" not in repr(fields)
    assert "Hidden question" not in repr(fields)


def test_research_reports_only_application_form_observations_with_sources():
    posting_url = "https://jobs.example.org/careers/role"
    application_url = "https://jobs.example.org/apply/role"
    pages = {
        posting_url: {
            "url": posting_url,
            "text": "Job details",
            "links": [{"label": "Apply now", "url": application_url}],
            "form_fields": [
                {"field_id": "F001", "label": "Newsletter email", "control_type": "email"}
            ],
            "observed_at": "2026-10-11T10:00:00Z",
        },
        application_url: {
            "url": application_url,
            "text": "Public application form",
            "links": [],
            "form_fields": [
                {
                    "field_id": "F001", "label": "Email", "control_type": "email",
                    "required": True, "options": [], "is_open_question": False,
                },
                {
                    "field_id": "F002", "label": "Why are you applying?", "control_type": "textarea",
                    "required": True, "options": [], "is_open_question": True,
                },
            ],
            "observed_at": "2026-10-11T10:01:00Z",
        },
    }
    checked = _validate_research(
        {"findings": [], "contacts": [], "no_contact_found_reason": "No contact."}, pages
    )

    form = checked["application_form"]
    assert form["status"] == "observed"
    assert [field["label"] for field in form["fields"]] == ["Email", "Why are you applying?"]
    assert form["questions"] == ["Why are you applying?"]
    assert form["application_pages"][0]["url"] == application_url
    assert form["fields"][1]["source_url"] == application_url
    assert form["fields"][1]["observed_at"] == "2026-10-11T10:01:00Z"
    context = _research_context(checked)
    assert context["application_form"]["questions"] == ["Why are you applying?"]


def test_crawler_rejects_private_dns_and_research_requires_source_supported_contact(settings, monkeypatch):
    from clue_ai import scrapling_research

    monkeypatch.setattr(
        scrapling_research.socket,
        "getaddrinfo",
        lambda *_args, **_kwargs: [(None, None, None, None, ("127.0.0.1", 443))],
    )
    with pytest.raises(ResearchCrawlError, match="Private and local"):
        scrapling_research._require_public_resolution("jobs.example.org")

    page = {
        "url": JOB_URL,
        "text": "Example Labs builds Python services. Alex Rivera, Engineering Manager at Example Labs, alex@example.org.",
        "observed_at": "2026-10-06T10:00:00Z",
    }
    research = {
        "findings": [],
        "contacts": [
            {
                "name": "Alex Rivera",
                "role": "Engineering Manager",
                "organization": "Example Labs",
                "contact_type": "likely_team_lead",
                "source_url": JOB_URL,
                "quote": "Alex Rivera, Engineering Manager at Example Labs",
                "public_email": "alex@example.org",
                "confidence": "high",
                "function_match": "high",
            },
            {
                "name": "Fake Contact",
                "role": "Chief Engineer",
                "organization": "Example Labs",
                "contact_type": "unknown",
                "source_url": JOB_URL,
                "quote": "Example Labs builds Python services",
                "public_email": "guessed@example.org",
                "confidence": "high",
                "function_match": "high",
            },
        ],
        "no_contact_found_reason": "",
    }
    checked = _validate_research(research, {JOB_URL: page})
    assert len(checked["contacts"]) == 1
    assert checked["contacts"][0]["name"] == "Alex Rivera"
    writer_context = _research_context(checked)
    assert '"public_email":' not in json.dumps(writer_context)
    assert "alex@example.org" not in json.dumps(writer_context)
    assert writer_context["contacts"][0]["public_email_verified"] is True
    assert "alex@example.org" not in json.dumps(writer_context)


def test_research_accepts_only_explicit_general_recruitment_inboxes():
    page = {
        "url": JOB_URL,
        "text": (
            "Example Labs builds Python services. If you are interested in this position, "
            "please submit your CV and letter of interest to careers@example.org. "
            "For general questions, email info@example.org."
        ),
        "observed_at": "2026-10-06T10:00:00Z",
    }
    research = {
        "findings": [],
        "contacts": [
            {
                "name": "Invented Person",
                "role": "Talent Director",
                "organization": "Example Labs",
                "contact_type": "general_recruitment_inbox",
                "source_url": JOB_URL,
                "quote": (
                    "please submit your CV and letter of interest to careers@example.org"
                ),
                "public_email": "careers@example.org",
                "confidence": "medium",
                "function_match": "unknown",
            },
            {
                "name": "Info inbox",
                "role": "General recruitment inbox",
                "organization": "Example Labs",
                "contact_type": "general_recruitment_inbox",
                "source_url": JOB_URL,
                "quote": "For general questions, email info@example.org.",
                "public_email": "info@example.org",
                "confidence": "high",
                "function_match": "high",
            },
        ],
        "no_contact_found_reason": "No contact was found.",
    }

    checked = _validate_research(research, {JOB_URL: page})

    assert len(checked["contacts"]) == 1
    contact = checked["contacts"][0]
    assert contact["name"] == "Recruitment team"
    assert contact["role"] == "General recruitment inbox"
    assert contact["contact_type"] == "general_recruitment_inbox"
    assert contact["confidence"] == "high"
    assert checked["no_contact_found_reason"] == ""
    writer_context = _research_context(checked)
    assert "careers@example.org" not in json.dumps(writer_context)
    assert "Invented Person" not in json.dumps(writer_context)


def test_document_metadata_uses_saved_listing_and_final_grounded_counts():
    assert _cover_letter_title({"title": "AI Engineer", "company": "IFOM"}) == (
        "Application for AI Engineer at IFOM"
    )
    assert _cover_letter_title({"title": "Broken\ufffd Role", "company": "IFOM"}) == (
        "Application for Broken Role at IFOM"
    )
    rewrite = {
        "resume_bullet_edits": [],
        "cover_letter_paragraphs": [{"text": "supported"}, {"text": "supported"}],
        "application_answers": [],
        "outreach_drafts": [],
        "change_summary": "The model can claim anything here.",
    }
    assert _rewrite_change_summary(rewrite) == (
        "Jev-supported draft content: 0 resume bullet edit(s); 2 cover-letter paragraph(s); "
        "0 application answer(s); 0 outreach draft(s)."
    )


def test_openai_reservation_is_separate_and_requires_fresh_consent(settings, database):
    gated = replace(settings, openai_api_key="synthetic-openai-key")
    with pytest.raises(OpenAIConfigurationError, match="disabled"):
        reserve_usage(
            database,
            gated,
            get_settings(database),
            request_id="synthetic-request",
            job_id="synthetic-job",
            stage="rewriter",
            payload={"input": "synthetic"},
            input_token_count=200,
            max_output_tokens=100,
        )
    with connect(database) as db:
        assert db.execute("SELECT COUNT(*) FROM openai_usage").fetchone()[0] == 0

    _enable_synthetic_openai(settings, database)
    payload = {"input": "Synthetic request with bounded local context."}
    usage_id = reserve_usage(
        database,
        gated,
        get_settings(database),
        request_id="synthetic-request",
        job_id="synthetic-job",
        stage="rewriter",
        payload=payload,
        input_token_count=400,
        max_output_tokens=100,
    )
    with connect(database) as db:
        row = db.execute("SELECT * FROM openai_usage WHERE id = ?", (usage_id,)).fetchone()
    payload_bytes = len(json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode())
    assert row["estimated_input_bytes"] == payload_bytes
    assert row["reserved_input_tokens"] == 400
    assert row["model"] == MODEL_ID
    settled = settle_usage(database, usage_id, input_tokens=390, output_tokens=50)
    assert settled > 0

    save_openai_controls(
        database,
        consent=True,
        monthly_cap_usd=0.00001,
        opportunity_cap_usd=0.00001,
        input_usd_per_million=0.1,
        output_usd_per_million=0.5,
        rate_card_revision=datetime.now(timezone.utc).date().isoformat(),
    )
    with pytest.raises(OpenAIBudgetError, match="exceed"):
        reserve_usage(
            database,
            gated,
            get_settings(database),
            request_id="synthetic-request-2",
            job_id="synthetic-job-2",
            stage="rewriter",
            payload=payload,
            input_token_count=400,
            max_output_tokens=100,
        )


def test_openai_long_context_reservation_uses_current_model_surcharge(settings, database):
    gated = _enable_synthetic_openai(settings, database)
    token_count = LONG_CONTEXT_SURCHARGE_THRESHOLD_TOKENS + 1
    usage_id = reserve_usage(
        database,
        gated,
        get_settings(database),
        request_id="synthetic-long-context",
        job_id="synthetic-long-context-job",
        stage="rewriter",
        payload={"input": "Synthetic long-context pricing boundary."},
        input_token_count=token_count,
        max_output_tokens=100,
    )
    with connect(database) as db:
        row = db.execute("SELECT * FROM openai_usage WHERE id = ?", (usage_id,)).fetchone()
    assert row["input_usd_per_million"] == 0.1 * LONG_CONTEXT_INPUT_RATE_MULTIPLIER
    assert row["output_usd_per_million"] == 0.5 * LONG_CONTEXT_OUTPUT_RATE_MULTIPLIER
    expected = (
        token_count * row["input_usd_per_million"] / 1_000_000
        + 100 * row["output_usd_per_million"] / 1_000_000
    )
    assert row["reserved_usd"] == pytest.approx(expected)
    settled = settle_usage(database, usage_id, input_tokens=token_count, output_tokens=100)
    assert settled == pytest.approx(expected)


def test_responses_client_retains_only_safe_http_error_diagnostics():
    from email.parser import BytesParser

    client = ResponsesClient("synthetic-key")
    headers = BytesParser(policy=policy.default).parsebytes(
        b"X-Request-ID: req_synthetic123\r\n\r\n"
    )
    body = json.dumps(
        {
            "error": {
                "type": "invalid_request_error",
                "code": "unsupported_parameter",
                "message": "PRIVATE_PROMPT_SENTINEL",
            }
        }
    ).encode()

    class FailingOpener:
        def open(self, _request, timeout):
            assert timeout == DEFAULT_GENERATION_TIMEOUT_SECONDS
            raise urllib.error.HTTPError(
                "https://api.openai.com/v1/responses",
                400,
                "Bad Request",
                headers,
                io.BytesIO(body),
            )

    client._opener = FailingOpener()
    payload = {
        "model": MODEL_ID,
        "reasoning": {"effort": REASONING_EFFORT},
        "store": False,
        "input": "Synthetic diagnostic only.",
    }
    with pytest.raises(OpenAIProviderError) as captured:
        client.create(payload)

    assert captured.value.status_code == 400
    assert captured.value.error_type == "invalid_request_error"
    assert captured.value.error_code == "unsupported_parameter"
    assert captured.value.request_id == "req_synthetic123"
    assert "HTTP 400" in str(captured.value)
    assert "PRIVATE_PROMPT_SENTINEL" not in str(captured.value)


def test_responses_client_counts_the_same_model_input_and_preserves_schema_and_tools():
    client = ResponsesClient("synthetic-key")
    payload = {
        "model": MODEL_ID,
        "reasoning": {"effort": REASONING_EFFORT},
        "instructions": "Synthetic instructions",
        "input": [{"role": "user", "content": "Synthetic request"}],
        "max_output_tokens": 300,
        "store": False,
        "text": {"format": {"type": "json_schema", "name": "synthetic", "schema": {"type": "object"}}},
        "tools": [CRAWL_TOOL],
        "tool_choice": "auto",
        "parallel_tool_calls": False,
    }

    class CountingOpener:
        request = None

        def open(self, request, timeout):
            self.request = request
            assert timeout == INPUT_TOKEN_COUNT_TIMEOUT_SECONDS
            return io.BytesIO(b'{"object":"response.input_tokens","input_tokens":321}')

    opener = CountingOpener()
    client._opener = opener
    assert client.count_input_tokens(payload) == 321
    assert opener.request.full_url == "https://api.openai.com/v1/responses/input_tokens"
    body = json.loads(opener.request.data)
    assert body == {
        "model": MODEL_ID,
        "input": payload["input"],
        "instructions": payload["instructions"],
        "tools": [CRAWL_TOOL],
        "text": payload["text"],
        "reasoning": payload["reasoning"],
        "tool_choice": "auto",
        "parallel_tool_calls": False,
    }
    assert "max_output_tokens" not in body
    assert "store" not in body


@pytest.mark.parametrize(
    "count_response",
    [
        {"object": "response.input_tokens", "input_tokens": True},
        {"object": "response.input_tokens", "input_tokens": -1},
        {"object": "response.input_tokens", "input_tokens": "321"},
        {"object": "response.input_tokens"},
    ],
)
def test_responses_client_rejects_invalid_input_token_counts(count_response):
    client = ResponsesClient("synthetic-key")

    class CountingOpener:
        def open(self, _request, timeout):
            assert timeout == INPUT_TOKEN_COUNT_TIMEOUT_SECONDS
            return io.BytesIO(json.dumps(count_response).encode())

    client._opener = CountingOpener()
    payload = {
        "model": MODEL_ID,
        "reasoning": {"effort": REASONING_EFFORT},
        "store": False,
        "input": "Synthetic token-count response.",
    }
    with pytest.raises(OpenAIProviderError, match="invalid input-token count"):
        client.count_input_tokens(payload)


def test_responses_client_uses_long_generation_timeout_and_preserves_response_id_on_read_timeout():
    client = ResponsesClient("synthetic-key")
    headers = BytesParser(policy=policy.default).parsebytes(
        b"X-Request-ID: req_synthetic_timeout\r\n\r\n"
    )

    class SlowResponse:
        def __init__(self):
            self.headers = headers

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def read(self, _limit):
            raise TimeoutError("synthetic timeout")

    class SlowOpener:
        def open(self, _request, timeout):
            assert timeout == DEFAULT_GENERATION_TIMEOUT_SECONDS
            return SlowResponse()

    client._opener = SlowOpener()
    payload = {
        "model": MODEL_ID,
        "reasoning": {"effort": REASONING_EFFORT},
        "store": False,
        "input": "Synthetic slow response.",
    }
    with pytest.raises(OpenAIProviderError, match="usage is unresolved") as captured:
        client.create(payload)

    assert captured.value.failure_kind == "timeout"
    assert captured.value.request_id == "req_synthetic_timeout"


def test_context_budget_uses_counted_input_plus_configured_output_allowance():
    ensure_context_fits(MODEL_CONTEXT_WINDOW_TOKENS - 300, 300)
    with pytest.raises(OpenAIProviderError, match="above GPT-6 Luna's 1,050,000-token context window"):
        ensure_context_fits(MODEL_CONTEXT_WINDOW_TOKENS - 299, 300)


def test_context_overflow_stops_before_usage_reservation_or_generation(settings, database):
    enabled = _enable_synthetic_openai(settings, database)

    class OversizedRequestClient:
        def count_input_tokens(self, _payload):
            return MODEL_CONTEXT_WINDOW_TOKENS

        def create(self, _payload):
            raise AssertionError("An over-context request must not reach generation.")

    with pytest.raises(OpenAIProviderError, match="above GPT-6 Luna's 1,050,000-token context window"):
        _api_call(
            database,
            enabled,
            {"id": "synthetic-request-context", "job_id": "synthetic-job-context"},
            "rewriter",
            [{"role": "user", "content": "Synthetic context boundary."}],
            OversizedRequestClient(),
            usage_ids=[],
        )
    with connect(database) as db:
        assert db.execute("SELECT COUNT(*) FROM openai_usage").fetchone()[0] == 0


def test_token_count_failure_stops_before_usage_reservation_or_generation(settings, database):
    enabled = _enable_synthetic_openai(settings, database)

    class CountFailureClient:
        def count_input_tokens(self, _payload):
            raise OpenAIProviderError("Synthetic token-count failure.", status_code=503)

        def create(self, _payload):
            raise AssertionError("A failed input-token count must stop generation.")

    with pytest.raises(OpenAIProviderError, match="Synthetic token-count failure"):
        _api_call(
            database,
            enabled,
            {"id": "synthetic-request-count-failure", "job_id": "synthetic-job-count-failure"},
            "rewriter",
            [{"role": "user", "content": "Synthetic count failure."}],
            CountFailureClient(),
            usage_ids=[],
        )
    with connect(database) as db:
        assert db.execute("SELECT COUNT(*) FROM openai_usage").fetchone()[0] == 0


def test_counted_input_reservation_still_enforces_local_spend_limit(settings, database):
    enabled = _enable_synthetic_openai(settings, database)
    save_openai_controls(
        database,
        consent=True,
        monthly_cap_usd=0.00001,
        opportunity_cap_usd=0.00001,
        input_usd_per_million=0.1,
        output_usd_per_million=0.5,
        rate_card_revision=datetime.now(timezone.utc).date().isoformat(),
    )

    class OverBudgetClient:
        def count_input_tokens(self, _payload):
            return 100

        def create(self, _payload):
            raise AssertionError("An over-budget generation must not reach the provider.")

    with pytest.raises(OpenAIBudgetError, match="exceed"):
        _api_call(
            database,
            enabled,
            {"id": "synthetic-request-budget", "job_id": "synthetic-job-budget"},
            "rewriter",
            [{"role": "user", "content": "Synthetic spend test."}],
            OverBudgetClient(),
            usage_ids=[],
        )
    with connect(database) as db:
        assert db.execute("SELECT COUNT(*) FROM openai_usage").fetchone()[0] == 0


def test_openai_http_diagnostic_reaches_workflow_and_marks_usage_unknown(settings, database):
    enabled = _enable_synthetic_openai(settings, database)

    class RejectedClient:
        def count_input_tokens(self, _payload):
            return 100

        def create(self, _payload):
            raise OpenAIProviderError(
                "OpenAI rejected the request; the provider error body was omitted.",
                status_code=503,
                error_type="server_error",
                error_code="server_error",
                request_id="req_synthetic123",
            )

    with pytest.raises(OpenAIProviderError, match="HTTP 503.*server_error"):
        _api_call(
            database,
            enabled,
            {"id": "synthetic-request-diagnostics", "job_id": "synthetic-job-diagnostics"},
            "rewriter",
            [{"role": "user", "content": "Synthetic diagnostic only."}],
            RejectedClient(),
            usage_ids=[],
        )

    with connect(database) as db:
        row = db.execute(
            "SELECT status, error_summary FROM openai_usage WHERE request_id = ?",
            ("synthetic-request-diagnostics",),
        ).fetchone()
    assert row["status"] == "unknown"
    assert "HTTP 503" in row["error_summary"]
    assert "server_error" in row["error_summary"]


def test_openai_response_model_mismatch_is_rejected_after_settling_usage(settings, database):
    enabled = _enable_synthetic_openai(settings, database)

    class MismatchedClient:
        def count_input_tokens(self, _payload):
            return 100

        def create(self, _payload):
            return {
                "model": "gpt-6-astra",
                "status": "completed",
                "output": [],
                "usage": {"input_tokens": 40, "output_tokens": 10},
            }

    with pytest.raises(OpenAIProviderError, match="did not use the configured model"):
        _api_call(
            database,
            enabled,
            {"id": "synthetic-request-model-mismatch", "job_id": "synthetic-job-model-mismatch"},
            "rewriter",
            [{"role": "user", "content": "Synthetic diagnostic only."}],
            MismatchedClient(),
            usage_ids=[],
        )

    with connect(database) as db:
        row = db.execute(
            "SELECT status, actual_usd, reserved_usd FROM openai_usage WHERE request_id = ?",
            ("synthetic-request-model-mismatch",),
        ).fetchone()
    assert row["status"] == "settled"
    assert row["actual_usd"] > 0
    assert row["reserved_usd"] > 0


@pytest.mark.parametrize(
    ("response", "message"),
    [
        (
            {
                "status": "incomplete",
                "output_text": "{}",
            },
            "did not complete",
        ),
        (
            {
                "status": "incomplete",
                "incomplete_details": {"reason": "max_output_tokens"},
                "output_text": "{}",
            },
            "output-token allowance",
        ),
        (
            {
                "status": "completed",
                "output": [
                    {"type": "message", "content": [{"type": "refusal", "refusal": "Synthetic"}]}
                ],
            },
            "returned a refusal",
        ),
        (
            {"status": "completed", "output_text": "not-json"},
            "invalid structured output",
        ),
        (
            {"status": "completed", "output_text": "[]"},
            "invalid output object",
        ),
    ],
)
def test_openai_stage_parser_rejects_incomplete_refusal_and_malformed_outputs(response, message):
    with pytest.raises(OpenAIProviderError, match=message):
        _parse_stage_result(response, "rewriter")


def test_openai_missing_usage_receipt_is_unknown_and_blocks_retry(settings, database):
    record = _create_preparation(settings, database)[0]
    enabled = _enable_synthetic_openai(settings, database)

    class MissingUsageClient:
        def count_input_tokens(self, _payload):
            return 100

        def create(self, _payload):
            return {"model": MODEL_ID, "status": "completed", "output": [], "usage": None}

    with pytest.raises(OpenAIProviderError, match="usage receipt; retry is blocked"):
        _api_call(
            database,
            enabled,
            {"id": record["id"], "job_id": record["job_id"]},
            "rewriter",
            [{"role": "user", "content": "Synthetic diagnostic only."}],
            MissingUsageClient(),
            usage_ids=[],
        )

    with connect(database) as db:
        row = db.execute(
            "SELECT status, reserved_usd FROM openai_usage WHERE request_id = ?",
            (record["id"],),
        ).fetchone()
    assert row["status"] == "unknown"
    assert row["reserved_usd"] > 0

    with connect(database) as db:
        db.execute(
            "UPDATE preparation_requests SET state = 'failed' WHERE id = ?",
            (record["id"],),
        )
    with pytest.raises(ValueError, match="previous API charge.*unresolved"):
        retry_preparation(database, record["id"])
    with pytest.raises(ValueError, match="previous API charge.*unresolved"):
        request_preparation(
            database,
            record["job_id"],
            record["run_id"],
            record["snapshot"]["inputs"]["selected_cv_id"],
        )


def test_openai_auth_rejection_releases_reservation_as_not_sent(settings, database):
    enabled = _enable_synthetic_openai(settings, database)

    class RejectedClient:
        def count_input_tokens(self, _payload):
            return 100

        def create(self, _payload):
            raise OpenAIProviderError(
                "OpenAI rejected the request; the provider error body was omitted.",
                status_code=401,
                error_type="invalid_request_error",
                error_code="invalid_api_key",
                request_id="req_synthetic123",
            )

    with pytest.raises(OpenAIProviderError, match="HTTP 401.*invalid_api_key"):
        _api_call(
            database,
            enabled,
            {"id": "synthetic-request-auth", "job_id": "synthetic-job-auth"},
            "rewriter",
            [{"role": "user", "content": "Synthetic diagnostic only."}],
            RejectedClient(),
            usage_ids=[],
        )

    with connect(database) as db:
        row = db.execute(
            "SELECT status, reserved_usd, actual_usd, error_summary FROM openai_usage WHERE request_id = ?",
            ("synthetic-request-auth",),
        ).fetchone()
    assert row["status"] == "released"
    assert row["reserved_usd"] == 0
    assert row["actual_usd"] == 0
    assert "invalid_api_key" in row["error_summary"]


def test_openai_local_preflight_releases_reservation_without_claiming_api_usage(settings, database):
    enabled = _enable_synthetic_openai(settings, database)

    class LocallyRejectedClient:
        def count_input_tokens(self, _payload):
            return 100

        def create(self, _payload):
            raise OpenAIConfigurationError("Synthetic local request validation failed.")

    with pytest.raises(OpenAIConfigurationError, match="local request validation"):
        _api_call(
            database,
            enabled,
            {"id": "synthetic-request-preflight", "job_id": "synthetic-job-preflight"},
            "rewriter",
            [{"role": "user", "content": "Synthetic diagnostic only."}],
            LocallyRejectedClient(),
            usage_ids=[],
        )

    with connect(database) as db:
        row = db.execute(
            "SELECT status, reserved_usd, actual_usd FROM openai_usage WHERE request_id = ?",
            ("synthetic-request-preflight",),
        ).fetchone()
    assert row["status"] == "released"
    assert row["reserved_usd"] == 0
    assert row["actual_usd"] == 0


def test_payloads_fix_model_storage_and_limit_crawler_to_researcher():
    plain = _payload("rewriter", {"test": "synthetic"})
    letter_reviser = _payload("letter_reviser", {"paragraphs_for_repair": [{"paragraph_index": 1}, {"paragraph_index": 2}]})
    researcher = _payload("researcher", {"test": "synthetic"}, tools=[CRAWL_TOOL])
    assert plain["model"] == researcher["model"] == "gpt-6-luna"
    assert plain["reasoning"] == letter_reviser["reasoning"] == researcher["reasoning"] == {"effort": "high"}
    assert plain["store"] is letter_reviser["store"] is researcher["store"] is False
    assert plain["text"]["format"]["name"] == "clue_rewriter_v14"
    assert letter_reviser["text"]["format"]["name"] == "clue_letter_reviser_v12"
    revision_array = letter_reviser["text"]["format"]["schema"]["properties"]["revisions"]
    assert revision_array["minItems"] == revision_array["maxItems"] == 2
    assert researcher["text"]["format"]["name"] == "clue_researcher_v4"
    assert "tools" not in plain
    assert researcher["tools"] == [CRAWL_TOOL]
    assert researcher["parallel_tool_calls"] is False
    with pytest.raises(OpenAIProviderError, match="must preserve every original CV line"):
        _validate_rewrite(
            {
                "selected_cv_id": "cv1",
                "resume_line_order": ["L0001"],
                "resume_bullet_edits": [],
                "cover_letter_paragraphs": [],
                "application_answers": [],
                "outreach_drafts": [],
            },
            "cv1",
            {"L0001": "Experience", "L0002": "Built a Python service."},
            {"claim1"},
            {"findings": [], "contacts": []},
            [],
            "allow_improvements",
        )
    with pytest.raises(OpenAIProviderError, match="reordered a CV"):
        _validate_rewrite(
            {
                "selected_cv_id": "cv1",
                "resume_line_order": ["L0002", "L0001"],
                "resume_bullet_edits": [],
                "cover_letter_paragraphs": [],
                "application_answers": [],
                "outreach_drafts": [],
            },
            "cv1",
            {"L0001": "Experience", "L0002": "Built a Python service."},
            {"claim1"},
            {"findings": [], "contacts": []},
            [],
            "preserve",
        )


def test_gmail_adapter_builds_only_a_create_draft_request(settings, monkeypatch):
    from clue_ai import gmail_drafts

    class TokenStore:
        def __init__(self, _data_dir):
            pass

        def load(self):
            return "synthetic-refresh-token"

    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def read(self, _limit):
            return json.dumps({"id": "synthetic-draft-id"}).encode()

    class FakeOpener:
        def open(self, request, timeout):
            captured["request"] = request
            captured["timeout"] = timeout
            return FakeResponse()

    captured = {}
    monkeypatch.setattr(gmail_drafts, "WindowsDPAPITokenStore", TokenStore)
    monkeypatch.setattr(gmail_drafts, "_refresh_access_token", lambda *_args: "synthetic-access-token")
    monkeypatch.setattr(gmail_drafts.urllib.request, "build_opener", lambda *_args: FakeOpener())
    configured = replace(settings, gmail_oauth_client_id="synthetic-google-client")
    result = create_unsent_draft(
        configured,
        "alex@example.org",
        "Question about the role",
        "Hello Alex, thank you for sharing the opening.",
    )
    assert result == "synthetic-draft-id"
    request = captured["request"]
    assert request.full_url == GMAIL_DRAFTS_URL
    assert request.get_method() == "POST"
    body = json.loads(request.data)
    mime_bytes = __import__("base64").urlsafe_b64decode(body["message"]["raw"] + "==")
    mime = BytesParser(policy=policy.default).parsebytes(mime_bytes)
    assert mime["To"] == "alex@example.org"
    assert mime["Subject"] == "Question about the role"
    assert "thank you for sharing" in mime.get_body(preferencelist=("plain",)).get_content()
    with pytest.raises(GmailDraftError, match="public recipient"):
        create_unsent_draft(configured, "person@@example.org", "Role", "Hello")


def test_gmail_disconnect_removes_local_token_and_stale_create_is_reconciled(settings, database):
    client = TestClient(create_app(settings), base_url="http://127.0.0.1")
    token_path = settings.data_dir / "gmail" / "refresh-token.bin"
    token_path.parent.mkdir(parents=True)
    token_path.write_bytes(b"synthetic-encrypted-token")
    set_gmail_connection(database, True)
    response = client.post("/settings/gmail/disconnect", headers=ORIGIN, follow_redirects=False)
    assert response.status_code == 303
    assert not token_path.exists()
    assert get_settings(database)["gmail_oauth_connected_at"] == ""

    record, _cv, _technical, _descriptive, _sample, _claim, _text = _create_preparation(
        settings, database
    )
    with connect(database) as db:
        db.execute(
            """INSERT INTO preparation_packets
               (id, request_id, revision, input_revision_sha256, output_json, status, created_at)
               VALUES ('packet-recovery', ?, 1, 'synthetic', '{}', 'review', '2026-10-06')""",
            (record["id"],),
        )
        db.execute(
            """INSERT INTO gmail_drafts
               (id, request_id, packet_id, recipient, subject, body, approved_sha256,
                state, created_at, updated_at)
               VALUES ('draft-recovery', ?, 'packet-recovery', 'alex@example.org', 'Role',
                       'Synthetic message', 'synthetic-hash', 'creating', '2026-10-06', '2026-10-06')""",
            (record["id"],),
        )
    assert recover_interrupted_gmail_drafts(database) == 1
    with connect(database) as db:
        draft = db.execute("SELECT state, error_summary FROM gmail_drafts WHERE id = 'draft-recovery'").fetchone()
    assert draft["state"] == "unknown"
    assert "Check Gmail Drafts before retrying" in draft["error_summary"]


def test_search_reset_retains_application_packets_and_provider_spend(settings, database):
    record, _cv, _technical, _descriptive, _sample, _claim, _text = _create_preparation(
        settings, database
    )
    _enable_synthetic_openai(settings, database)
    reserve_usage(
        database,
        replace(settings, openai_api_key="synthetic-openai-key"),
        get_settings(database),
        request_id=record["id"],
        job_id=record["job_id"],
        stage="researcher",
        payload={"input": "synthetic"},
        input_token_count=20,
        max_output_tokens=10,
    )
    result = reset_search_data(database)
    assert result.backup_path.is_file()
    with connect(database) as db:
        assert db.execute("SELECT COUNT(*) FROM jobs").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM search_runs").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM preparation_requests").fetchone()[0] == 1
        assert db.execute("SELECT COUNT(*) FROM openai_usage").fetchone()[0] == 1
    assert Path(record["snapshot"]["inputs"]["selected_cv_filename"])


def test_full_personal_data_deletion_removes_preparation_and_provider_records(settings, database):
    record, _cv, _technical, _descriptive, _sample, _claim, _text = _create_preparation(
        settings, database
    )
    _enable_synthetic_openai(settings, database)
    set_gmail_connection(database, True)
    with connect(database) as db:
        db.execute(
            """INSERT INTO preparation_packets
               (id, request_id, revision, input_revision_sha256, output_json, status, created_at)
               VALUES ('synthetic-packet', ?, 1, 'synthetic-revision', '{}', 'approved', '2026-10-06T00:00:00Z')""",
            (record["id"],),
        )
        db.execute(
            """INSERT INTO application_followups
               (id, request_id, kind, due_on, state, created_at, updated_at)
               VALUES ('synthetic-followup', ?, 'outreach', '2026-10-07', 'scheduled',
                       '2026-10-06T00:00:00Z', '2026-10-06T00:00:00Z')""",
            (record["id"],),
        )
        db.execute(
            """INSERT INTO packet_feedback
               (id, packet_id, owner_minutes, quality_rating, factual_corrections, created_at, updated_at)
               VALUES ('synthetic-feedback', 'synthetic-packet', 10, 4, 0,
                       '2026-10-06T00:00:00Z', '2026-10-06T00:00:00Z')"""
        )
    for dirname in ("cv", "application-packets", "gmail"):
        path = settings.data_dir / dirname / "synthetic"
        path.mkdir(parents=True, exist_ok=True)
        (path / "synthetic.txt").write_text("synthetic private data", encoding="utf-8")
    delete_personal_data(database, None, settings.data_dir)
    with connect(database) as db:
        assert db.execute("SELECT COUNT(*) FROM preparation_sources").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM candidate_claims").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM preparation_requests").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM preparation_packets").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM application_followups").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM packet_feedback").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM gmail_drafts").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM openai_usage").fetchone()[0] == 0
        setting = db.execute(
            "SELECT openai_consent_at, gmail_consent_at, gmail_oauth_connected_at FROM app_settings WHERE id = 1"
        ).fetchone()
    assert setting["openai_consent_at"] == ""
    assert setting["gmail_consent_at"] == ""
    assert setting["gmail_oauth_connected_at"] == ""
    assert not (settings.data_dir / "cv").exists()
    assert not (settings.data_dir / "application-packets").exists()
    assert not (settings.data_dir / "gmail").exists()
    assert record["id"]


def test_repository_template_keeps_openai_key_blank():
    template = Path(".env.example").read_text(encoding="utf-8")
    assert "OPENAI_API_KEY=" in template
    assert not any(line.strip().startswith("OPENAI_API_KEY=") and line.partition("=")[2].strip() for line in template.splitlines())


def test_application_workspace_groups_packet_versions_and_submission_receipts(settings, database):
    record, _cv, _technical, _descriptive, _sample, _claim_id, _resume = _create_preparation(settings, database)
    enabled = _enable_synthetic_openai(settings, database)
    run_preparation(
        database,
        enabled,
        record["id"],
        client_factory=FakePreparationClient,
        crawler_factory=FakeCrawler,
        tailored_resume_reviewer=_approve_synthetic_tailored_resume,
    )
    with connect(database) as db:
        packet_id = db.execute(
            "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
            (record["id"],),
        ).fetchone()["id"]

    client = TestClient(create_app(enabled), base_url="http://127.0.0.1")
    approved = client.post(
        f"/preparations/{record['id']}/packets/{packet_id}/approve",
        data={"confirm_review": "on"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert approved.status_code == 303
    submitted = client.post(
        f"/preparations/{record['id']}/record-stage",
        data={"stage": "submitted", "owner_attestation": "on"},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert submitted.status_code == 303
    uploaded = client.post(
        f"/preparations/{record['id']}/receipt",
        files={"receipt_file": ("confirmation.txt", b"Synthetic receipt", "text/plain")},
        headers=ORIGIN,
        follow_redirects=False,
    )
    assert uploaded.status_code == 303

    older_bytes = b"Older synthetic application file"
    older_path = settings.data_dir / "application-packets" / record["id"] / "v0" / "older-resume.txt"
    older_path.parent.mkdir(parents=True, exist_ok=True)
    older_path.write_bytes(older_bytes)
    with connect(database) as db:
        db.execute(
            """INSERT INTO preparation_packets
               (id, request_id, revision, input_revision_sha256, output_json, status, created_at)
               VALUES ('historical-packet', ?, 0, 'older-input-revision', ?, 'obsolete', '2026-10-05T00:00:00Z')""",
            (record["id"], json.dumps({"research": {}, "diagnoser": {}, "recruiter": {}, "rewriter": {}})),
        )
        db.execute(
            """INSERT INTO packet_artifacts
               (id, packet_id, artifact_type, filename, file_path, content_sha256, created_at)
               VALUES ('historical-artifact', 'historical-packet', 'resume', 'older-resume.txt', ?, ?, '2026-10-05T00:00:00Z')""",
            (str(older_path), sha256(older_bytes).hexdigest()),
        )

    index = client.get("/applications")
    assert index.status_code == 200
    assert "Applications" in index.text
    assert "Choose a role to generate a packet" in index.text
    assert f"/applications/{record['id']}" in index.text
    assert "Application stage · Submitted" in index.text

    detail = client.get(f"/applications/{record['id']}")
    assert detail.status_code == 200
    assert "Files for this application" in detail.text
    assert "Current packet · version 1 · Approved" in detail.text
    assert "Older packet · version 0 · Obsolete · 1 file" in detail.text
    assert "confirmation.txt" in detail.text
    assert "alex@example.org" in detail.text
    assert "/outcomes/receipts/" in detail.text
    historical_download = client.get("/packets/historical-packet/artifacts/historical-artifact")
    assert historical_download.status_code == 200
    assert historical_download.content == older_bytes
    assert client.get(f"/preparations/{record['id']}").status_code == 200
    assert client.get("/applications/not-a-request").status_code == 404


def test_cover_letter_body_cannot_duplicate_rendered_salutation_and_signoff():
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][0]["text"] = (
        "Dear Hiring Team,\n\n" + rewrite["cover_letter_paragraphs"][0]["text"]
    )
    rewrite["cover_letter_paragraphs"][-1]["text"] += "\n\nYours sincerely,\nCandidate"
    quality = _cover_letter_quality(rewrite, _supported_letter_grounding())
    assert quality["status"] == "failed"
    assert "Return body prose only" in " ".join(quality["issues"])
    assert 0 in quality["revision_paragraph_indices"]
    assert len(rewrite["cover_letter_paragraphs"]) - 1 in quality["revision_paragraph_indices"]


def test_letter_profile_evidence_in_introduction_counts_without_a_section_quota():
    preferred = ["intro-profile"]
    rewrite = _complete_quality_letter()
    rewrite["cover_letter_paragraphs"][0]["claim_ids"] = preferred
    quality = _cover_letter_quality(
        rewrite, _supported_letter_grounding(), preferred_profile_claim_ids=preferred
    )
    assert quality["preferred_profile_claims_used"] == 1
    assert "use_technical_profile_evidence" not in quality["revision_reasons"].values()


def test_second_letter_repair_rebases_surviving_indices_and_stops_after_two(monkeypatch):
    from clue_ai import application_workflow as workflow

    paragraphs = [{"section": "evidence", "text": "Supported body."}]
    grounding = {"results": [
        {"id": "cover_letter_paragraph_0", "output_type": "cover_letter_paragraph", "final_artifact_id": None},
        {"id": "cover_letter_paragraph_1", "output_type": "cover_letter_paragraph", "final_artifact_id": "cover_letter_paragraph_0"},
    ]}
    calls = []

    def repair(**kwargs):
        calls.append(kwargs)
        return {"cover_letter_paragraphs": paragraphs}, grounding, {
            "status": "completed", "attempted": True,
            "final_writing_review": {"status": "needs_revision"},
        }

    monkeypatch.setattr(workflow, "_repair_cover_letter_once", repair)
    monkeypatch.setattr(workflow, "_cover_letter_quality", lambda *_args, **_kwargs: {"status": "failed"})
    _, _, report = workflow._repair_cover_letter(
        original_paragraphs=[], recruiter_requirements=[],
        preferred_profile_claim_ids=[], technical_profile_sources=[],
    )
    assert len(calls) == 2
    assert calls[1]["original_paragraphs"] == paragraphs
    assert [item["id"] for item in calls[1]["grounding"]["results"]] == ["cover_letter_paragraph_0"]
    assert report["attempt_limit"] == 2
    assert report["attempt_count"] == 2
    assert len(report["attempt_reports"]) == 2


def test_missing_profile_repair_does_not_reassign_evidence_already_in_introduction():
    from clue_ai import application_workflow as workflow

    context = workflow._letter_repair_context(
        job={"title": "AI Engineer", "canonical_url": "https://example.org/job"},
        original_paragraphs=[
            {"section": "opening", "text": "I built an agent workflow system.", "claim_ids": ["agent"]},
            {"section": "closing", "text": "Thank you for considering my application."},
        ],
        grounding={"results": []}, failed_indices=[2],
        repair_reasons={2: "missing_candidate_evidence"}, approved_claims=[],
        technical_profile_sources=[{"source_id": "profile", "filename": "profile.md", "use": "cover_letter", "evidence": [
            {"id": "agent", "claim_text": "I built an agent workflow system.", "evidence_excerpt": "I built an agent workflow system.", "review_status": "unreviewed"},
            {"id": "deployment", "claim_text": "I deployed a model on FPGA hardware.", "evidence_excerpt": "I deployed a model on FPGA hardware.", "review_status": "unreviewed"},
        ]}], preferred_profile_claim_ids=["agent", "deployment"],
        writing_preferences={}, owner_writing_context=[], research={"findings": []},
        recruiter_requirements=[], cv_lines={}, quality_issues=[],
    )
    assert context["paragraphs_for_repair"][0]["allowed_claim_ids"] == ["deployment"]


def test_local_caps_can_be_disabled_without_disabling_usage_tracking(settings, database):
    from clue_ai.repository import reserve_jev_budget

    gated = replace(_enable_synthetic_openai(settings, database), enforce_local_spend_caps=False)
    save_openai_controls(
        database, consent=True, monthly_cap_usd=0.000001, opportunity_cap_usd=0.000001,
        input_usd_per_million=0.1, output_usd_per_million=0.5,
        rate_card_revision=datetime.now(timezone.utc).date().isoformat(),
    )
    usage_id = reserve_usage(
        database, gated, get_settings(database), request_id="no-local-cap", job_id="job",
        stage="rewriter", payload={"input": "Synthetic request"},
        input_token_count=10000, max_output_tokens=1000,
    )
    jev_id, _ = reserve_jev_budget(
        database, None, "jev-test", 80000, 0.042, 0.000001, enforce_cap=False
    )
    with connect(database) as db:
        assert db.execute("SELECT status FROM openai_usage WHERE id=?", (usage_id,)).fetchone()["status"] == "reserved"
        assert db.execute("SELECT status FROM jev_usage WHERE id=?", (jev_id,)).fetchone()["status"] == "reserved"
    assert settle_usage(database, usage_id, input_tokens=10, output_tokens=5) > 0
