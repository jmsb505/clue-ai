from __future__ import annotations

import io
import struct
import zipfile
import zlib
from datetime import datetime, timezone
from xml.etree import ElementTree

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from clue_ai.domain import SearchCriteria
from clue_ai.filters import filter_jobs
from clue_ai.jobs import canonical_url, classify_location, plain_text
from clue_ai.resume import (
    ResumeError,
    extract_resume_text,
    infer_profile_language,
    parse_candidate_profile,
    suggest_profile_sections,
)


def make_docx(text: str) -> bytes:
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    return save_docx(document)


def save_docx(document: Document) -> bytes:
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_png() -> bytes:
    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">2I5B", 1, 1, 8, 6, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(b"\x00\x00\x00\x00\x00"))
        + chunk(b"IEND", b"")
    )


def make_pdf(text: str, page_count: int = 1) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    for _ in range(page_count - 1):
        writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
    )
    stream = DecodedStreamObject()
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream.set_data(f"BT /F1 12 Tf 40 700 Td ({escaped}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_docx_text_is_extracted_for_review_and_section_suggestions(settings):
    content = make_docx(
        "SUMMARY\nProduct engineer with experience building reliable applications.\n"
        "SKILLS\nPython, SQL, APIs, and accessibility.\n"
        "EXPERIENCE\nBuilt local-first products for distributed teams."
    )

    text = extract_resume_text("resume.docx", content, settings)

    assert "Product engineer" in text
    assert suggest_profile_sections(text)["skills"] == "Python, SQL, APIs, and accessibility."


def test_docx_extracts_headers_footers_tables_and_document_order(settings):
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Alex Rivera | alex@example.test"
    document.add_paragraph("SUMMARY\nProduct engineer focused on accessible software systems.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "COLUMN ONE PROJECT"
    table.cell(0, 1).text = "COLUMN TWO PROJECT"
    table.cell(1, 0).text = "Built a service with careful evaluation."
    table.cell(1, 1).text = "Added a typed local API."
    document.add_paragraph("EDUCATION\nSynthetic computer science coursework.")
    document.sections[0].footer.paragraphs[0].text = "Portfolio: https://example.test/alex"

    text = extract_resume_text("resume.docx", save_docx(document), settings)

    expected_order = [
        "Alex Rivera",
        "Product engineer",
        "COLUMN ONE PROJECT",
        "COLUMN TWO PROJECT",
        "Synthetic computer science coursework.",
        "Portfolio: https://example.test/alex",
    ]
    positions = [text.index(part) for part in expected_order]
    assert positions == sorted(positions)


def test_docx_extracts_text_boxes_and_warns_when_images_are_not_ocrd(settings):
    document = Document()
    document.add_paragraph("SUMMARY\nSynthetic engineer with enough selectable text for review.")
    document.add_paragraph().add_run().add_picture(io.BytesIO(make_png()))
    content = save_docx(document)
    rewritten_parts: list[tuple[str, bytes]] = []
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        for item in archive.infolist():
            data = archive.read(item.filename)
            if item.filename == "word/document.xml":
                root = ElementTree.fromstring(data)
                body = root.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body")
                textbox_paragraph = ElementTree.Element(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
                )
                run_with_shape = ElementTree.SubElement(
                    textbox_paragraph,
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r",
                )
                drawing = ElementTree.SubElement(
                    run_with_shape,
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing",
                )
                text_box = ElementTree.SubElement(
                    drawing,
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}txbxContent",
                )
                paragraph = ElementTree.SubElement(
                    text_box,
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p",
                )
                run = ElementTree.SubElement(
                    paragraph,
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r",
                )
                box_text = ElementTree.SubElement(
                    run,
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t",
                )
                box_text.text = "Selected callout text from a Word text box."
                body.insert(len(body) - 1, textbox_paragraph)
                data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            rewritten_parts.append((item.filename, data))

    rewritten = io.BytesIO()
    with zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as archive:
        for filename, data in rewritten_parts:
            archive.writestr(filename, data)

    text = extract_resume_text("resume.docx", rewritten.getvalue(), settings)

    assert "Selected callout text from a Word text box." in text
    assert "Text box content (visual reading order may differ)" in text
    assert "Embedded image present" in text
    assert "OCR is not performed" in text


def test_resume_extraction_rejects_silent_truncation(settings):
    limited = settings.__class__(
        data_dir=settings.data_dir,
        api_key="",
        model=settings.model,
        monthly_jev_budget_usd=4.0,
        max_extracted_chars=60,
    )

    with pytest.raises(ResumeError, match="extracted text exceeds the 60-character local limit"):
        extract_resume_text("resume.docx", make_docx("A" * 100), limited)


def test_candidate_profile_parser_derives_editable_english_role_and_sections():
    text = """Alex Rivera
alex@example.test
SUMMARY
Product designer focused on accessible software and clear customer experiences.
SKILLS
Figma, user research, prototyping, design systems, accessibility
EXPERIENCE
Senior Product Designer | Northstar Studio | 2022 – Present
Led product design for a remote team, improving onboarding and usability.
EDUCATION
MSc Human Computer Interaction
LANGUAGES
English C1, Italian B2
"""

    profile = parse_candidate_profile(text)

    assert profile["target_roles"] == "Senior Product Designer"
    assert "accessible software" in profile["summary"]
    assert "Figma" in profile["skills"]
    assert "Northstar Studio" in profile["experience"]
    assert profile["profile_language"] == "en"
    assert "alex@example.test" not in " ".join(profile.values())


def test_candidate_profile_parser_uses_explicit_target_roles_and_conservative_language():
    text = """Obiettivo
Cerco una posizione professionale in un team italiano che collabora con clienti e aziende in tutta Europa.
TARGET ROLES
Product Designer, UX Researcher
ESPERIENZA
Ho progettato servizi digitali e coordinato la ricerca con utenti e gruppi di lavoro.
"""

    profile = parse_candidate_profile(text)

    assert profile["target_roles"] == "Product Designer, UX Researcher"
    assert profile["profile_language"] == "it"
    assert infer_profile_language("Short CV") == "unknown"


def test_candidate_profile_parser_does_not_turn_generic_experience_bullets_into_roles():
    profile = parse_candidate_profile(
        """SUMMARY
        Reliable teammate who improved processes across an international organization.
        EXPERIENCE
        Built better handoffs and worked with colleagues on internal improvements.
        """
    )

    assert profile["target_roles"] == ""


def test_pdf_text_is_extracted_locally_and_page_limit_is_enforced(settings):
    text = "Synthetic candidate profile with selectable text for the local PDF extraction flow."

    extracted = extract_resume_text("resume.pdf", make_pdf(text), settings)

    assert text in extracted
    page_limited = settings.__class__(
        data_dir=settings.data_dir,
        api_key="",
        model=settings.model,
        monthly_jev_budget_usd=4.0,
        max_cv_pages=1,
    )
    with pytest.raises(ResumeError, match="up to 1 pages"):
        extract_resume_text("many-pages.pdf", make_pdf(text, page_count=2), page_limited)


def test_scanned_pdf_with_no_selectable_text_is_rejected_locally(settings):
    with pytest.raises(ResumeError, match="no usable selectable text"):
        extract_resume_text("scanned-resume.pdf", make_pdf(""), settings)


def test_resume_rejects_wrong_extension_invalid_file_and_oversize(settings):
    with pytest.raises(ResumeError, match="PDF or DOCX"):
        extract_resume_text("resume.txt", b"x" * 80, settings)
    with pytest.raises(ResumeError, match="not a PDF"):
        extract_resume_text("resume.pdf", b"not pdf" * 20, settings)
    small_limit = settings.__class__(
        data_dir=settings.data_dir,
        api_key="",
        model=settings.model,
        monthly_jev_budget_usd=4.0,
        max_cv_bytes=20,
    )
    with pytest.raises(ResumeError, match="12 MB or smaller"):
        extract_resume_text("resume.docx", b"x" * 21, small_limit)


def test_plain_text_drops_script_and_style_content():
    assert plain_text("<p>Remote role</p><script>secret()</script><style>.x{}</style>") == "Remote role"


@pytest.mark.parametrize(
    ("location", "description", "expected"),
    [
        ("Italy", "Fully remote role", "eligible"),
        ("Europe", "Fully remote role", "eligible"),
        ("Remote", "Work from anywhere in the world", "eligible"),
        ("Remote", "Remote team; location requirements are not listed", "needs_verification"),
        ("United States only", "Fully remote role", "not_eligible"),
    ],
)
def test_location_classification_requires_explicit_geographic_evidence(
    location, description, expected
):
    status, evidence = classify_location(location, description, "Italy")

    assert status == expected
    assert evidence


def test_canonical_url_removes_tracking_and_rejects_private_targets():
    assert canonical_url("HTTPS://Jobs.Example.org/role?utm_source=mail&ref=board&id=7#apply") == (
        "https://jobs.example.org/role?id=7"
    )
    assert canonical_url("https://user:pass@jobs.example.org/role") == ""
    assert canonical_url("http://127.0.0.1/private") == ""
    assert canonical_url("https://internal.example.local/jobs") == ""


def test_remote_preferred_scope_drops_unverified_remote_location():
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    common = {
        "title": "Software Engineer",
        "company": "Example Labs",
        "description": "Remote Python role. Build APIs with Python and SQL.",
        "workplace_type": "remote",
        "employment_type": "full-time",
        "posted_at": now,
        "last_checked_at": now,
        "salary_min": None,
        "salary_max": None,
        "salary_currency": "",
        "salary_period": "",
        "visa_sponsorship": "unknown",
    }
    eligible = {**common, "id": "eligible", "location_raw": "Europe"}
    uncertain = {**common, "id": "uncertain", "location_raw": "Remote"}
    irrelevant = {**common, "id": "irrelevant", "title": "Designer", "location_raw": "Italy"}

    results = filter_jobs(
        [eligible, uncertain, irrelevant],
        SearchCriteria(roles="Software Engineer", must_have="Python", include_unknown_location=True),
    )

    assert [item["id"] for item in results] == ["eligible"]
    assert results[0]["eligibility_status"] == "eligible"
    assert results[0]["eligibility_evidence"]
    assert {item["freshness_status"] for item in results} == {"recent"}


def test_unknown_location_can_be_excluded_by_explicit_user_choice():
    item = {
        "id": "uncertain",
        "title": "Software Engineer",
        "company": "Example Labs",
        "description": "Remote software engineering position.",
        "location_raw": "Remote",
        "workplace_type": "remote",
        "posted_at": "",
        "last_checked_at": "",
    }

    assert filter_jobs(
        [item], SearchCriteria(workplace="any", include_unknown_location=False)
    ) == []


def test_results_checked_over_24_hours_ago_are_marked_stale():
    item = {
        "id": "old-check",
        "title": "Software Engineer",
        "company": "Example Labs",
        "description": "Fully remote role open to candidates in Italy.",
        "location_raw": "Italy",
        "workplace_type": "remote",
        "posted_at": "",
        "last_checked_at": "2020-01-01T00:00:00+00:00",
    }

    results = filter_jobs([item], SearchCriteria())

    assert len(results) == 1
    assert results[0]["freshness_status"] == "stale"
