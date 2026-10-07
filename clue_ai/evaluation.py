from __future__ import annotations

import argparse
import json
import math
import shutil
import tempfile
from collections.abc import Callable
from dataclasses import dataclass, replace
from datetime import datetime, timedelta, timezone
from math import isfinite
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

from clue_ai.config import Settings
from clue_ai.database import connect, initialize, save_profile, save_search_run, set_jev_consent
from clue_ai.domain import CandidateProfile, NormalizedJob, SearchCriteria
from clue_ai.filters import filter_jobs
from clue_ai.jev import score_run
from clue_ai.jobs import canonical_url
from clue_ai.repository import (
    all_active_jobs,
    get_run_results,
    monthly_jev_usage,
    save_jobs_with_report,
    save_run_results,
)


@dataclass(frozen=True)
class BenchmarkCase:
    job: NormalizedJob
    relevance: int
    should_pass_hard_filters: bool
    expected_jev_filter_status: str
    expected_jev_filter_checks: dict[str, str]
    rationale: str


@dataclass(frozen=True)
class GroundingCase:
    case_id: str
    assertion: str
    evidence: str
    expected: str
    rationale: str


MAX_EVALUATION_BUDGET_USD = 0.05
GROUNDING_LABELS = ("supported", "contradicted", "unresolved")
JEV_FILTER_LABELS = ("match", "review", "conflict", "unassessed")
JEV_GROUNDING_RUBRIC_VERSION = "jev-grounding-v1.3"

WRITING_RUBRIC = {
    "evidence_traceability": "Every factual sentence maps to approved evidence or a verified public source.",
    "factual_accuracy": "No invented number, date, name, credential, scope, responsibility, or outcome.",
    "role_specificity": "Draft uses a real role requirement and a relevant owner-approved proof point.",
    "tailored_resume_jev_approval": "Jev explicitly approves the final tailored resume against the selected listing and saved match.",
    "packet_quality": "A packet is reviewable only when it has a supported, role-specific cover letter and all required artifacts.",
    "voice": "Direct, modest, specific, and consistent with owner-confirmed writing preferences.",
    "xyz_integrity": "XYZ is used only when X, Y, and Z are all supported; qualitative impact is retained when no metric exists.",
    "question_adherence": "Only supplied application questions are answered; unknown attestations remain unresolved.",
    "research_provenance": "Contact, company, and role statements have exact public source provenance.",
    "concision": "Each output is brief enough for its intended reader and channel.",
}

SYNTHETIC_PREPARATION_PROJECTS = (
    {
        "name": "SignalForge",
        "claim": (
            "SignalForge: built a Python/PyTorch RAG service with FastAPI and citation validation; "
            "on a synthetic 240-question set it reached recall@5 of 0.82 and citation precision of 0.94."
        ),
        "resume_line": (
            "- SignalForge: built a Python/PyTorch RAG service with FastAPI and citation validation; "
            "on a synthetic 240-question set it reached recall@5 of 0.82 and citation precision of 0.94."
        ),
        "evidence_markers": ("recall@5 of 0.82", "citation precision of 0.94"),
        "role_families": ("llm", "evaluation", "platform"),
    },
    {
        "name": "EdgeReg",
        "claim": (
            "EdgeReg: implemented 2.5D PyTorch MR-to-CT registration and deployed INT8 inference "
            "to a ZCU104 DPU; on 9 held-out synthetic cases, mean Dice was 0.7647 and label-centroid TRE was 2.025 mm."
        ),
        "resume_line": (
            "- EdgeReg: implemented 2.5D PyTorch MR-to-CT registration and deployed INT8 inference "
            "to a ZCU104 DPU; on 9 held-out synthetic cases, mean Dice was 0.7647 and label-centroid TRE was 2.025 mm."
        ),
        "evidence_markers": ("mean Dice was 0.7647", "TRE was 2.025 mm"),
        "role_families": ("computer_vision", "edge", "medical_ai"),
    },
    {
        "name": "AgentDock",
        "claim": (
            "AgentDock: built a typed run-based agent system with immutable versions, allowlisted "
            "tool calls, approval-gated actions, deterministic regression comparisons, and idempotent recovery."
        ),
        "resume_line": (
            "- AgentDock: built a typed run-based agent system with immutable versions, allowlisted "
            "tool calls, approval-gated actions, deterministic regression comparisons, and idempotent recovery."
        ),
        "evidence_markers": ("approval-gated actions", "idempotent recovery"),
        "role_families": ("agents", "platform", "security"),
    },
    {
        "name": "FilmLab",
        "claim": (
            "FilmLab: trained a deterministic neural renderer from scratch, served PyTorch inference "
            "through a local FastAPI service, and connected it to a React/Tauri application."
        ),
        "resume_line": (
            "- FilmLab: trained a deterministic neural renderer from scratch, served PyTorch inference "
            "through a local FastAPI service, and connected it to a React/Tauri application."
        ),
        "evidence_markers": ("deterministic neural renderer", "React/Tauri application"),
        "role_families": ("computer_vision", "product"),
    },
    {
        "name": "QueryBench",
        "claim": (
            "QueryBench: compared dense, lexical, and hybrid retrieval on 240 synthetic queries; "
            "the hybrid configuration reached recall@10 of 0.88."
        ),
        "resume_line": (
            "- QueryBench: compared dense, lexical, and hybrid retrieval on 240 synthetic queries; "
            "the hybrid configuration reached recall@10 of 0.88."
        ),
        "evidence_markers": ("240 synthetic queries", "recall@10 of 0.88"),
        "role_families": ("llm", "evaluation"),
    },
    {
        "name": "PromptGuard",
        "claim": (
            "PromptGuard: built a Python prompt-injection test harness with 180 synthetic attack "
            "cases; it blocked 91% of the attacks while retaining all benign control prompts."
        ),
        "resume_line": (
            "- PromptGuard: built a Python prompt-injection test harness with 180 synthetic attack "
            "cases; it blocked 91% of the attacks while retaining all benign control prompts."
        ),
        "evidence_markers": ("180 synthetic attack cases", "blocked 91%"),
        "role_families": ("llm", "security", "evaluation"),
    },
    {
        "name": "ModelRouter",
        "claim": (
            "ModelRouter: implemented a typed Python service that routes model requests by task, "
            "validates structured outputs, and applies bounded retry and idempotent fallback rules."
        ),
        "resume_line": (
            "- ModelRouter: implemented a typed Python service that routes model requests by task, "
            "validates structured outputs, and applies bounded retry and idempotent fallback rules."
        ),
        "evidence_markers": ("structured outputs", "idempotent fallback rules"),
        "role_families": ("llm", "agents", "platform"),
    },
    {
        "name": "VisionEdge",
        "claim": (
            "VisionEdge: quantized a PyTorch defect-detection model for edge inference; median "
            "latency was 18 ms across 1,000 synthetic frames."
        ),
        "resume_line": (
            "- VisionEdge: quantized a PyTorch defect-detection model for edge inference; median "
            "latency was 18 ms across 1,000 synthetic frames."
        ),
        "evidence_markers": ("18 ms", "1,000 synthetic frames"),
        "role_families": ("computer_vision", "edge"),
    },
    {
        "name": "DataHarbor",
        "claim": (
            "DataHarbor: created a schema-validated Python ingestion pipeline for four synthetic "
            "datasets, with invalid-row quarantine and reproducible train/validation partitions."
        ),
        "resume_line": (
            "- DataHarbor: created a schema-validated Python ingestion pipeline for four synthetic "
            "datasets, with invalid-row quarantine and reproducible train/validation partitions."
        ),
        "evidence_markers": ("four synthetic datasets", "reproducible train/validation partitions"),
        "role_families": ("data", "platform", "evaluation"),
    },
    {
        "name": "DocExtract",
        "claim": (
            "DocExtract: built layout-aware PDF and table extraction for 320 synthetic documents; "
            "the exact-field F1 score was 0.93."
        ),
        "resume_line": (
            "- DocExtract: built layout-aware PDF and table extraction for 320 synthetic documents; "
            "the exact-field F1 score was 0.93."
        ),
        "evidence_markers": ("320 synthetic documents", "exact-field F1 score was 0.93"),
        "role_families": ("computer_vision", "data", "document_ai"),
    },
    {
        "name": "EvalHarness",
        "claim": (
            "EvalHarness: created a paired evaluation suite for 480 synthetic assistant questions, "
            "with deterministic fixtures and CI failures for regressions in retrieval and citation checks."
        ),
        "resume_line": (
            "- EvalHarness: created a paired evaluation suite for 480 synthetic assistant questions, "
            "with deterministic fixtures and CI failures for regressions in retrieval and citation checks."
        ),
        "evidence_markers": ("480 synthetic assistant questions", "CI failures for regressions"),
        "role_families": ("llm", "evaluation", "platform"),
    },
    {
        "name": "TraceKit",
        "claim": (
            "TraceKit: added event-level tracing for agent runs with redacted arguments, immutable "
            "version references, and request-to-artifact provenance."
        ),
        "resume_line": (
            "- TraceKit: added event-level tracing for agent runs with redacted arguments, immutable "
            "version references, and request-to-artifact provenance."
        ),
        "evidence_markers": ("redacted arguments", "request-to-artifact provenance"),
        "role_families": ("agents", "platform", "security"),
    },
)

SYNTHETIC_PREPARATION_ROLE_FAMILIES = {
    "SYN-01": "llm",
    "SYN-02": "computer_vision",
}
RECRUITER_ROLE_KEYWORD_GROUPS = {
    "llm": (("retrieval", "rag"), ("python",), ("pytorch", "torch"), ("citation", "evaluation", "evaluate")),
    "computer_vision": (("computer vision", "perception", "vision"), ("pytorch", "torch"), ("edge", "inference")),
}
HIRING_MANAGER_ROLE_KEYWORD_GROUPS = {
    "llm": (("retrieval", "rag"), ("python",), ("pytorch", "torch"), ("citation", "evaluation")),
    "computer_vision": (("computer vision", "perception", "vision"), ("pytorch", "torch"), ("edge", "inference")),
}


class BenchmarkRunError(RuntimeError):
    def __init__(
        self,
        phase: str,
        error_type: str,
        *,
        usage_recovery: dict[str, Any] | None = None,
    ) -> None:
        detail = (
            f"; local usage ledger retained at {usage_recovery['directory']}"
            if usage_recovery and usage_recovery.get("directory")
            else ""
        )
        super().__init__(f"{phase}: {error_type}{detail}")
        self.phase = phase
        self.error_type = error_type
        self.usage_recovery = usage_recovery


def _unresolved_provider_usage(database_path: Path) -> list[dict[str, Any]]:
    if not database_path.is_file():
        return []
    with connect(database_path) as db:
        openai_rows = db.execute(
            """SELECT id, request_id, stage, status, reserved_usd
               FROM openai_usage WHERE status IN ('reserved', 'unknown')
               ORDER BY created_at, id"""
        ).fetchall()
        jev_rows = db.execute(
            """SELECT id, run_id, status, reserved_usd
               FROM jev_usage WHERE status IN ('reserved', 'unknown')
               ORDER BY started_at, id"""
        ).fetchall()
    return [
        {
            "provider": "openai",
            "ledger_row_id": row["id"],
            "request_id": row["request_id"],
            "stage": row["stage"],
            "status": row["status"],
            "reserved_usd": float(row["reserved_usd"] or 0),
        }
        for row in openai_rows
    ] + [
        {
            "provider": "jev",
            "ledger_row_id": row["id"],
            "run_id": row["run_id"],
            "status": row["status"],
            "reserved_usd": float(row["reserved_usd"] or 0),
        }
        for row in jev_rows
    ]


class _EvaluationTemporaryDirectory:
    """Delete synthetic evaluation data only when both provider ledgers are settled."""

    def __init__(
        self,
        prefix: str,
        database_path_for: Callable[[Path], Path],
        *,
        parent_directory: Path | None = None,
    ) -> None:
        self.name = tempfile.mkdtemp(
            prefix=prefix,
            dir=str(parent_directory) if parent_directory else None,
        )
        self._database_path_for = database_path_for
        self._finalized = False
        self.usage_recovery: dict[str, Any] | None = None

    def __enter__(self) -> str:
        return self.name

    def __exit__(self, *_args: object) -> bool:
        self.cleanup()
        return False

    def cleanup(self) -> None:
        if self._finalized:
            return
        data_dir = Path(self.name).resolve()
        try:
            unresolved = _unresolved_provider_usage(self._database_path_for(data_dir))
        except Exception as exc:  # noqa: BLE001 - a failed ledger check must preserve the evidence.
            self.usage_recovery = {
                "directory": str(data_dir),
                "ledger_check_error": type(exc).__name__,
                "unresolved_usage": [],
            }
            self._finalized = True
            return
        if unresolved:
            self.usage_recovery = {
                "directory": str(data_dir),
                "unresolved_usage": unresolved,
            }
            self._finalized = True
            return
        temp_root = Path(tempfile.gettempdir()).resolve()
        if data_dir == temp_root or not data_dir.is_relative_to(temp_root):
            raise RuntimeError("Refusing to remove an evaluation directory outside the temp root.")
        shutil.rmtree(data_dir)
        self._finalized = True


def synthetic_benchmark() -> tuple[CandidateProfile, SearchCriteria, list[BenchmarkCase]]:
    """Return a synthetic junior AI-engineering benchmark aligned with Clue's search scope."""
    profile = CandidateProfile(
        summary=(
            "Synthetic junior AI engineer with project experience building Python services, "
            "machine-learning prototypes, and retrieval-augmented applications."
        ),
        target_roles="Junior AI Engineer, ML Engineer, LLM Engineer, Computer Vision Engineer",
        skills="Python, PyTorch, scikit-learn, FastAPI, RAG, Docker, Transformers",
        experience=(
            "Portfolio projects using Python and PyTorch; built a FastAPI service and a small "
            "retrieval-augmented assistant. No professional AI engineering tenure is claimed."
        ),
        education="Synthetic bachelor's degree in computer science",
        languages="English",
        profile_language="en",
    )
    criteria = SearchCriteria(
        work_from="Italy",
        roles="AI engineer, machine learning engineer, LLM engineer, computer vision engineer",
        workplace="remote",
        must_have="Python, PyTorch",
        nice_to_have="RAG, Docker",
        include_unknown_location=False,
        posted_within_days=365,
    )
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    stale = (datetime.now(timezone.utc) - timedelta(hours=25)).isoformat(timespec="seconds")
    examples = [
        (
            "SYN-01", "Junior LLM Engineer", "Synthetic Northstar AI", "Italy · Remote",
            """Paid, full-time junior role for candidates working remotely from Italy.
            Build and evaluate retrieval-augmented generation (RAG) features with Python and PyTorch.
            Applicants with 0–2 years of experience are welcome; the working language is English.
            The team provides mentorship and reviews model behavior with product engineers.""",
            "remote", 4, True, "match",
            "Clear target role, paid junior scope, Italy eligibility, Python/PyTorch, and English.", now,
        ),
        (
            "SYN-02", "Entry-Level Computer Vision Engineer", "Synthetic Alder Robotics", "Europe · Remote",
            """Paid entry-level position, remote across the EU including Italy. Help train and
            evaluate computer-vision models using Python and PyTorch. Zero to two years of
            experience is acceptable. English is the team's working language; Italian is optional.""",
            "remote", 4, True, "match",
            "Direct AI/vision duties; paid entry-level role, EU remote eligibility, and required tools.", now,
        ),
        (
            "SYN-03", "AI Integration Engineer", "Synthetic Meridian Systems", "Italy · Remote",
            """Paid, full-time AI engineering role, remote from Italy. Build Python and PyTorch
            integrations for model-serving systems and write tests for model-backed features.
            The description does not state an experience level or years requirement. English is
            used for technical documentation and team meetings.""",
            "remote", 3, True, "review",
            "Role and constraints fit, but the posting gives no evidence about seniority.", now,
        ),
        (
            "SYN-04", "Senior AI Platform Engineer", "Synthetic Lumen Cloud", "Italy · Remote",
            """Paid, full-time AI engineering role, remote from Italy. Own production ML platform
            architecture and lead cross-team delivery. Requires at least six years of experience
            and prior staff-level scope. The team builds with Python and PyTorch; English is required.""",
            "remote", 3, True, "conflict",
            "Local search should retain it; Jev should flag its explicit senior scope as conflict.", now,
        ),
        (
            "SYN-05", "Junior AI Engineer Intern", "Synthetic Cedar Labs", "Europe · Remote",
            """Unpaid volunteer internship, remote across Europe including Italy. Assist with
            Python and PyTorch experiments on language models. No prior experience is required;
            English is used by the research group.""",
            "remote", 2, True, "conflict",
            "Local search should retain it; Jev should conflict with the explicit unpaid condition.", stale,
        ),
        (
            "SYN-06", "AI Engineer — On-site Milan", "Synthetic Milan Lab", "Milan, Italy",
            """Paid junior AI engineer position based on-site in Milan. Build Python and PyTorch
            model prototypes and document evaluation results. English is used by the team.""",
            "onsite", 3, False, "unassessed",
            "Outside the remote-only workplace criterion in this deterministic search fixture.", now,
        ),
        (
            "SYN-07", "Junior ML Engineer — US Only", "Synthetic Redwood Models", "Remote · United States only",
            """Paid junior machine-learning role limited to US residents and work authorization.
            Build models with Python and PyTorch. English is the working language; the company
            cannot hire in Italy or elsewhere in Europe.""",
            "remote", 3, False, "unassessed",
            "Explicit country restriction is outside Italy; local geography gate must exclude it.", now,
        ),
        (
            "SYN-08", "Junior AI Product Engineer", "Synthetic Kestrel Apps", "Italy · Remote",
            """Paid junior AI product role, remote from Italy, with 0–2 years of experience.
            Build model-backed features using Java and React. English is required. Python and
            PyTorch are not used in this position.""",
            "remote", 3, False, "unassessed",
            "Explicit required Python/PyTorch evidence is absent; local must-have filter excludes it.", now,
        ),
    ]
    expected_checks_by_id = {
        "SYN-01": {key: "match" for key in ("language", "location", "pay", "requirements", "seniority", "workplace")},
        "SYN-02": {key: "match" for key in ("language", "location", "pay", "requirements", "seniority", "workplace")},
        "SYN-03": {
            **{key: "match" for key in ("language", "location", "pay", "requirements", "workplace")},
            "seniority": "review",
        },
        "SYN-04": {
            **{key: "match" for key in ("language", "location", "pay", "requirements", "workplace")},
            "seniority": "conflict",
        },
        "SYN-05": {
            **{key: "match" for key in ("language", "location", "requirements", "seniority", "workplace")},
            "pay": "conflict",
        },
        "SYN-06": {},
        "SYN-07": {},
        "SYN-08": {},
    }
    cases = []
    for (
        external_id, title, company, location, description, workplace, relevance, passes,
        expected_jev_filter_status, rationale, checked,
    ) in examples:
        url = f"https://benchmark.example/jobs/{external_id.casefold()}"
        cases.append(
            BenchmarkCase(
                job=NormalizedJob(
                    source_id="jobicy",
                    source_name="Synthetic benchmark",
                    external_id=external_id,
                    title=title,
                    company=company,
                    description=" ".join(description.split()),
                    source_url=url,
                    canonical_url=url,
                    location_raw=location,
                    workplace_type=workplace,
                    employment_type=("internship" if external_id == "SYN-05" else "full-time"),
                    posted_at=now,
                    source_credit="Synthetic benchmark (no external source)",
                    last_checked_at=checked,
                ),
                relevance=relevance,
                should_pass_hard_filters=passes,
                expected_jev_filter_status=expected_jev_filter_status,
                expected_jev_filter_checks=expected_checks_by_id[external_id],
                rationale=rationale,
            )
        )
    return profile, criteria, cases

def synthetic_grounding_benchmark() -> list[GroundingCase]:
    """Return synthetic assertion/evidence pairs for Jev's separate support-check role."""
    return [
        GroundingCase(
            "GND-01",
            "Built APIs using Python.",
            "Built REST APIs using Python.",
            "supported",
            "Equivalent wording with the same scope and technology.",
        ),
        GroundingCase(
            "GND-02",
            "Reduced query latency to 120 ms.",
            "Reduced query latency to 120 milliseconds in load tests.",
            "supported",
            "The result, number, unit, and context are explicit.",
        ),
        GroundingCase(
            "GND-03",
            "Partnered with product teams on reporting.",
            "Worked with product teams to build operational reports.",
            "supported",
            "The paraphrase preserves the documented collaboration and work.",
        ),
        GroundingCase(
            "GND-04",
            "Presented findings to business partners.",
            "Explained analysis results to business stakeholders.",
            "supported",
            "Equivalent audience and communication activity.",
        ),
        GroundingCase(
            "GND-05",
            "Led a 20-person engineering team.",
            "Worked as an individual contributor; no direct reports.",
            "contradicted",
            "Evidence explicitly conflicts with the management and team-size claim.",
        ),
        GroundingCase(
            "GND-06",
            "Launched the service in 2024.",
            "The service launched in 2022.",
            "contradicted",
            "The asserted year conflicts with the documented year.",
        ),
        GroundingCase(
            "GND-07",
            "Earned a master's degree in statistics.",
            "Bachelor's degree in statistics.",
            "contradicted",
            "The asserted qualification conflicts with the documented degree.",
        ),
        GroundingCase(
            "GND-08",
            "Increased annual revenue by 300%.",
            "Built a Python service; no revenue or business-impact metric is recorded.",
            "unresolved",
            "The evidence contains no revenue outcome or metric.",
        ),
        GroundingCase(
            "GND-09",
            "Managed a €2 million budget.",
            "Implemented analytics dashboards and data validation.",
            "unresolved",
            "No budget responsibility or financial amount is documented.",
        ),
        GroundingCase(
            "GND-10",
            "Available to start immediately.",
            "The owner has not provided availability information.",
            "unresolved",
            "Availability is not present in the evidence.",
        ),
        GroundingCase(
            "GND-11",
            "Fluent in Italian.",
            "English language skills are documented; Italian is not mentioned.",
            "unresolved",
            "Missing language evidence does not establish fluency or its absence.",
        ),
        GroundingCase(
            "GND-12",
            "Improved reliability through monitoring and tests.",
            "Added service monitoring and automated tests to improve reliability.",
            "supported",
            "The actions and outcome are stated in the evidence.",
        ),
    ]


def keyword_hits(job: dict[str, Any], profile: CandidateProfile) -> int:
    text = " ".join(str(job.get(field) or "") for field in ("title", "description")).casefold()
    terms = [
        "ai engineer",
        "machine learning",
        "llm",
        "computer vision",
        "python",
        "pytorch",
        "rag",
        "docker",
    ]
    candidate_text = " ".join(profile.fit_fields().values()).casefold()
    return sum(1 for term in terms if term in candidate_text and term in text)


def ndcg_at_k(ranked_ids: list[str], relevance: dict[str, int], k: int) -> float:
    def dcg(grades: list[int]) -> float:
        return sum((2**grade - 1) / math.log2(index + 2) for index, grade in enumerate(grades))

    actual = dcg([relevance.get(job_id, 0) for job_id in ranked_ids[:k]])
    ideal = dcg(sorted(relevance.values(), reverse=True)[:k])
    return actual / ideal if ideal else 0.0


def binary_metrics(expected: dict[str, bool], actual: dict[str, bool]) -> dict[str, Any]:
    """Calculate binary classification metrics without treating absent labels as negatives."""
    if set(expected) != set(actual):
        raise ValueError("Expected and actual labels must contain the same case IDs.")
    true_positive = sum(expected[key] and actual[key] for key in expected)
    false_positive = sum(not expected[key] and actual[key] for key in expected)
    false_negative = sum(expected[key] and not actual[key] for key in expected)
    true_negative = sum(not expected[key] and not actual[key] for key in expected)
    total = len(expected)
    precision_denominator = true_positive + false_positive
    recall_denominator = true_positive + false_negative
    precision = true_positive / precision_denominator if precision_denominator else None
    recall = true_positive / recall_denominator if recall_denominator else None
    f1 = (
        2 * precision * recall / (precision + recall)
        if precision is not None and recall is not None and precision + recall
        else None
    )
    return {
        "true_positive": true_positive,
        "false_positive": false_positive,
        "false_negative": false_negative,
        "true_negative": true_negative,
        "precision": precision,
        "recall": recall,
        "f1": f1,
        "accuracy": (true_positive + true_negative) / total if total else None,
        "case_count": total,
    }


def categorical_metrics(
    expected: dict[str, str],
    actual: dict[str, str],
    labels: tuple[str, ...] = GROUNDING_LABELS,
) -> dict[str, Any]:
    """Return a multi-class confusion matrix and per-class precision/recall/F1."""
    if set(expected) != set(actual):
        raise ValueError("Expected and actual labels must contain the same case IDs.")
    if any(value not in labels for value in (*expected.values(), *actual.values())):
        raise ValueError("Classification labels must belong to the declared label set.")
    matrix = {label: {predicted: 0 for predicted in labels} for label in labels}
    for case_id, expected_label in expected.items():
        matrix[expected_label][actual[case_id]] += 1
    per_class = {}
    for label in labels:
        true_positive = matrix[label][label]
        false_positive = sum(matrix[other][label] for other in labels if other != label)
        false_negative = sum(matrix[label][other] for other in labels if other != label)
        precision_denominator = true_positive + false_positive
        recall_denominator = true_positive + false_negative
        precision = true_positive / precision_denominator if precision_denominator else None
        recall = true_positive / recall_denominator if recall_denominator else None
        f1 = (
            2 * precision * recall / (precision + recall)
            if precision is not None and recall is not None and precision + recall
            else None
        )
        per_class[label] = {
            "precision": precision,
            "recall": recall,
            "f1": f1,
            "support": recall_denominator,
        }
    f1_values = [value["f1"] for value in per_class.values() if value["f1"] is not None]
    return {
        "case_count": len(expected),
        "accuracy": (
            sum(matrix[label][label] for label in labels) / len(expected) if expected else None
        ),
        "macro_f1": sum(f1_values) / len(f1_values) if f1_values else None,
        "labels": list(labels),
        "confusion_matrix": matrix,
        "per_class": per_class,
    }


def _bounded_evaluation_budget(base_budget_usd: float, requested_budget_usd: float) -> float:
    if (
        not isfinite(requested_budget_usd)
        or requested_budget_usd <= 0
        or requested_budget_usd > MAX_EVALUATION_BUDGET_USD
    ):
        raise ValueError(
            f"Evaluation budget must be greater than zero and no more than "
            f"{MAX_EVALUATION_BUDGET_USD:.2f} USD."
        )
    return min(max(0.0, base_budget_usd), 4.0, requested_budget_usd)


def _benchmark_job_id(job: dict[str, Any]) -> str:
    return str(job.get("title") or "")


def _link_is_structurally_valid(url: str) -> bool:
    parts = urlsplit(url)
    return (
        parts.scheme == "https"
        and bool(parts.hostname)
        and not parts.username
        and not parts.password
        and canonical_url(url) == url
    )


def run_live_benchmark(
    client_factory: Callable[..., Any] | None = None,
    *,
    max_budget_usd: float = MAX_EVALUATION_BUDGET_USD,
) -> dict[str, Any]:
    phase = "load local settings"
    temporary: _EvaluationTemporaryDirectory | None = None
    try:
        base_settings = Settings.from_environment()
        if not base_settings.api_key:
            raise RuntimeError("No TypeSafe API key is configured; no request was made.")
        evaluation_budget = _bounded_evaluation_budget(
            base_settings.monthly_jev_budget_usd,
            max_budget_usd,
        )
        profile, criteria, cases = synthetic_benchmark()
        run_id = "synthetic-jev-benchmark"
        phase = "create disposable local database"
        temporary = _EvaluationTemporaryDirectory(
            "clue-jev-benchmark-",
            lambda data_dir: replace(base_settings, data_dir=data_dir).database_path,
        )
        with temporary as temp_dir:
            settings = replace(
                base_settings,
                data_dir=Path(temp_dir),
                monthly_jev_budget_usd=evaluation_budget,
            )
            initialize(settings.database_path)
            set_jev_consent(settings.database_path, True)
            save_profile(settings.database_path, profile)
            save_search_run(settings.database_path, run_id, criteria)
            duplicate_report = save_jobs_with_report(
                settings.database_path,
                [case.job for case in cases] + [cases[0].job],
            )
            indexed_jobs = all_active_jobs(settings.database_path)
            filtered_jobs = filter_jobs(indexed_jobs, criteria)
            save_run_results(settings.database_path, run_id, filtered_jobs)
            stored_results = get_run_results(settings.database_path, run_id)
            phase = "score synthetic listings with Jev"
            fit_result = score_run(
                settings.database_path,
                settings,
                run_id,
                stored_results,
                profile,
                criteria,
                client_factory=client_factory,
            )
            phase = "read temporary evaluation results"
            results = get_run_results(settings.database_path, run_id)
            by_title = {_benchmark_job_id(job): job for job in results}
            case_by_title = {case.job.title: case for case in cases}
            relevance = {
                key: case.relevance for key, case in case_by_title.items() if key in by_title
            }
            case_order = [case.job.title for case in cases if case.job.title in by_title]
            jev_order = sorted(
                case_order,
                key=lambda key: (
                    by_title[key].get("combined_score") is not None,
                    float(by_title[key].get("combined_score") or -1.0),
                ),
                reverse=True,
            )
            keyword_order = sorted(
                case_order,
                key=lambda key: keyword_hits(by_title[key], profile),
                reverse=True,
            )
            expected = {case.job.title for case in cases if case.should_pass_hard_filters}
            actual = set(case_order)
            false_positives = sorted(actual - expected)
            false_negatives = sorted(expected - actual)
            expected_filter_labels = {
                case.job.title: case.should_pass_hard_filters for case in cases
            }
            actual_filter_labels = {
                case.job.title: case.job.title in actual for case in cases
            }
            expected_jev_labels = {
                case.job.title: case.expected_jev_filter_status
                for case in cases if case.should_pass_hard_filters
            }
            actual_jev_labels = {
                title: str(by_title[title].get("filter_status") or "unassessed")
                for title in expected_jev_labels
            }
            jev_check_metrics = {}
            jev_check_agreement = {}
            for check in ("language", "location", "pay", "requirements", "seniority", "workplace"):
                expected_by_title = {
                    case.job.title: case.expected_jev_filter_checks[check]
                    for case in cases
                    if case.should_pass_hard_filters and check in case.expected_jev_filter_checks
                }
                actual_by_title = {
                    title: str(
                        ((by_title[title].get("dimensions") or {}).get(f"filter_{check}") or {}).get("status")
                        or "unassessed"
                    )
                    for title in expected_by_title
                }
                jev_check_metrics[check] = categorical_metrics(
                    expected_by_title,
                    actual_by_title,
                    JEV_FILTER_LABELS,
                )
                for title, expected_status in expected_by_title.items():
                    case_id = case_by_title[title].job.external_id
                    jev_check_agreement.setdefault(case_id, {})[check] = {
                        "expected": expected_status,
                        "actual": actual_by_title[title],
                    }
            stale_ids = sorted(
                key for key in case_order if by_title[key].get("freshness_status") == "stale"
            )
            usage = monthly_jev_usage(settings.database_path, settings.monthly_jev_budget_usd)
            rows = []
            for rank, title in enumerate(jev_order, start=1):
                job = by_title[title]
                source_url = (job.get("sources") or [{}])[0].get("source_url") or ""
                rows.append(
                    {
                        "id": case_by_title[title].job.external_id,
                        "title": job["title"],
                        "human_relevance": case_by_title[title].relevance,
                        "expected_hard_filter": case_by_title[title].should_pass_hard_filters,
                        "expected_jev_filter_status": case_by_title[title].expected_jev_filter_status,
                        "hard_filter_rationale": case_by_title[title].rationale,
                        "keyword_hits": keyword_hits(job, profile),
                        "jev_rank": rank,
                        "fit_score": job.get("combined_score"),
                        "confidence": job.get("confidence"),
                        "score_state": job.get("score_state"),
                        "filter_status": job.get("filter_status"),
                        "eligibility_status": job.get("eligibility_status"),
                        "filter_checks": {
                            key: value
                            for key, value in (job.get("dimensions") or {}).items()
                            if key.startswith("filter_")
                        },
                        "freshness": job.get("freshness_status"),
                        "canonical_link_valid": _link_is_structurally_valid(source_url),
                    }
                )
            report = {
                "candidate_and_jobs_are_synthetic": True,
                "labels_source": "assistant_authored_synthetic_not_owner_validation",
                "api_key_value_printed": False,
                "api_requests": usage["requests"],
                "app_ledger_cost_usd": usage["used_usd"],
                "app_ledger_open_reserve_usd": sum(_usage_rows(settings.database_path)),
                "app_budget_usd": usage["budget_usd"],
                "duplicate_records_detected": duplicate_report.deduplicated,
                "hard_filter_expected_passes": len(expected),
                "hard_filter_actual_passes": len(actual),
                "hard_filter_false_positives": false_positives,
                "hard_filter_false_negatives": false_negatives,
                "hard_filter_metrics": binary_metrics(
                    expected_filter_labels,
                    actual_filter_labels,
                ),
                "jev_filter_metrics": categorical_metrics(
                    expected_jev_labels,
                    actual_jev_labels,
                    JEV_FILTER_LABELS,
                ),
                "jev_filter_check_metrics": jev_check_metrics,
                "jev_filter_check_case_agreement": jev_check_agreement,
                "jev_filter_case_agreement": {
                    case.job.external_id: {
                        "expected": expected_jev_labels.get(case.job.title),
                        "actual": actual_jev_labels.get(case.job.title, "not_sent_to_jev"),
                        "rationale": case.rationale,
                    }
                    for case in cases
                },
                "stale_filtered_results": stale_ids,
                "canonical_links_structurally_valid": all(
                    row["canonical_link_valid"] for row in rows
                ),
                "jev_ndcg_at_5": ndcg_at_k(jev_order, relevance, 5),
                "keyword_ndcg_at_5": ndcg_at_k(keyword_order, relevance, 5),
                "scored_count": fit_result.scored_count,
                "unscored_count": fit_result.unscored_count,
                "jev_filter_decision_counts": {
                    status: sum(job.get("filter_status") == status for job in results)
                    for status in ("match", "review", "conflict", "unassessed")
                },
                "jev_confidence_semantics": (
                    "Model-reported assessment confidence; not a probability of hiring or job success."
                ),
                "jev_message": fit_result.message,
                "ranked_results": rows,
            }
        if temporary.usage_recovery:
            report["usage_recovery"] = temporary.usage_recovery
        return report
    except Exception as exc:  # noqa: BLE001 - keep failure output free of payloads and credentials.
        if temporary is not None:
            temporary.cleanup()
        raise BenchmarkRunError(
            phase,
            type(exc).__name__,
            usage_recovery=temporary.usage_recovery if temporary else None,
        ) from None


def run_live_grounding_benchmark(
    client_factory: Callable[..., Any] | None = None,
    *,
    max_budget_usd: float = 0.025,
) -> dict[str, Any]:
    """Run the synthetic assertion/evidence set through Jev in a disposable ledger."""
    from clue_ai.claim_support import check_generated_claim_support

    budget = _bounded_evaluation_budget(max_budget_usd, max_budget_usd)
    base_settings = Settings.from_environment()
    temporary = _EvaluationTemporaryDirectory(
        "clue-jev-grounding-",
        lambda data_dir: replace(base_settings, data_dir=data_dir).database_path,
    )
    settings = replace(
        base_settings,
        data_dir=Path(temporary.name),
        monthly_jev_budget_usd=budget,
    )
    phase = "initialize synthetic grounding benchmark"
    result: dict[str, Any] | None = None
    try:
        initialize(settings.database_path)
        set_jev_consent(settings.database_path, True)
        cases = synthetic_grounding_benchmark()
        assertions = [
            {
                "id": case.case_id,
                "output_type": "synthetic_grounding_case",
                "text": case.assertion,
                "evidence": [
                    {
                        "source_type": "synthetic_approved_claim",
                        "source_id": case.case_id,
                        "text": case.evidence,
                    }
                ],
            }
            for case in cases
        ]
        phase = "send synthetic grounding cases to Jev"
        verdict = check_generated_claim_support(
            settings.database_path,
            settings,
            "synthetic-grounding-benchmark",
            assertions,
            client_factory=client_factory,
        )
        expected = {case.case_id: case.expected for case in cases}
        actual = {item["id"]: item["status"] for item in verdict["results"]}
        usage = monthly_jev_usage(settings.database_path, budget)
        result = {
            "candidate_and_evidence_are_synthetic": True,
            "labels_source": "assistant_authored_synthetic_not_owner_validation",
            "rubric_version": JEV_GROUNDING_RUBRIC_VERSION,
            "api_requests": usage["requests"],
            "app_ledger_cost_usd": usage["used_usd"],
            "app_ledger_open_reserve_usd": sum(_usage_rows(settings.database_path)),
            "app_budget_usd": usage["budget_usd"],
            "case_count": len(cases),
            "metrics": categorical_metrics(expected, actual),
            "case_results": {
                item["id"]: {
                    "expected": expected[item["id"]],
                    "actual": item["status"],
                    "model_status": item.get("model_status"),
                    "confidence": item.get("confidence"),
                }
                for item in verdict["results"]
            },
            "model": verdict.get("model"),
            "actual_input_tokens": verdict.get("actual_input_tokens"),
            "confidence_policy": "Recorded for diagnosis only; Jev's categorical support verdict is authoritative unless a deterministic numeric/date guard finds unsupported values.",
        }
        return result
    except Exception as exc:  # noqa: BLE001 - never report synthetic prompts, keys, or provider response bodies.
        temporary.cleanup()
        raise BenchmarkRunError(
            phase,
            type(exc).__name__,
            usage_recovery=temporary.usage_recovery,
        ) from None
    finally:
        temporary.cleanup()
        if result is not None and temporary.usage_recovery:
            result["usage_recovery"] = temporary.usage_recovery


def _synthetic_preparation_case(case_id: str = "SYN-01") -> BenchmarkCase:
    if case_id not in SYNTHETIC_PREPARATION_ROLE_FAMILIES:
        raise ValueError(f"Unsupported synthetic preparation case: {case_id}")
    _profile, _criteria, cases = synthetic_benchmark()
    case = next((item for item in cases if item.job.external_id == case_id), None)
    if case is None or not case.should_pass_hard_filters or case.expected_jev_filter_status != "match":
        raise ValueError(f"Synthetic preparation case is not an eligible match: {case_id}")
    return case


def run_live_preparation_benchmark(case_id: str = "SYN-01") -> dict[str, Any]:
    """Run a synthetic Jev→GPT→Jev→DOCX flow and Hiring Manager practice case."""
    from docx import Document

    from clue_ai.application_prep import (
        add_source,
        get_preparation,
        list_claims,
        request_preparation,
        review_claim,
        set_source_options,
        suggest_claims,
    )
    from clue_ai.application_workflow import (
        _call_stage,
        _cv_lines,
        _filter_diagnoser_references,
        _jev_read_only_context,
        _job_context,
        _quote_supported,
        _tailored_resume_text,
        create_practice_session,
        get_packet,
        queue_practice_answers,
        run_interview_practice_assessment,
        run_interview_practice_questions,
        run_preparation,
    )
    from clue_ai.config import load_local_environment
    from clue_ai.database import connect, get_settings, save_openai_controls, save_search_run
    from clue_ai.domain import utc_now
    from clue_ai.openai_provider import ResponsesClient
    from clue_ai.repository import (
        all_active_jobs,
        get_run_results,
        save_jobs,
        save_run_results,
        update_run,
    )

    load_local_environment(override_keys={"OPENAI_API_KEY"})
    base_settings = Settings.from_environment()
    if not base_settings.openai_api_key:
        raise BenchmarkRunError("check OpenAI configuration", "OpenAIConfigurationError")
    if not base_settings.api_key:
        raise BenchmarkRunError("check Jev configuration", "JevConfigurationError")
    temporary = _EvaluationTemporaryDirectory(
        "clue-preparation-eval-",
        lambda data_dir: replace(base_settings, data_dir=data_dir).database_path,
    )
    settings = replace(
        base_settings,
        data_dir=Path(temporary.name),
        monthly_jev_budget_usd=0.05,
    )
    phase = "initialize synthetic preparation database"
    result: dict[str, Any] | None = None
    try:
        initialize(settings.database_path)
        set_jev_consent(settings.database_path, True)
        rate_date = datetime.now(timezone.utc).date().isoformat()
        save_openai_controls(
            settings.database_path,
            consent=True,
            monthly_cap_usd=0.10,
            opportunity_cap_usd=0.05,
            input_usd_per_million=0.10,
            output_usd_per_million=0.50,
            rate_card_revision=rate_date,
        )
        profile, criteria, _ = synthetic_benchmark()
        profile = CandidateProfile(
            summary="Synthetic applied AI engineer with a multi-project portfolio across retrieval, agents, edge inference, and product integration.",
            target_roles="Junior AI Systems Engineer, LLM Engineer, Computer Vision Engineer, Edge AI Engineer",
            skills="Python, PyTorch, RAG, FastAPI, agent evaluation, Vitis AI, React, TypeScript, retrieval, evaluation",
            experience="\n".join(item["claim"] for item in SYNTHETIC_PREPARATION_PROJECTS),
            education="Synthetic bachelor's degree in computer science",
            languages="English",
            profile_language="en",
        )
        selected_case = _synthetic_preparation_case(case_id)
        role_family = SYNTHETIC_PREPARATION_ROLE_FAMILIES[case_id]
        job = selected_case.job
        phase = "save synthetic job and Jev input"
        save_jobs(settings.database_path, [job])
        with connect(settings.database_path) as db:
            job_id = db.execute(
                "SELECT id FROM jobs WHERE canonical_url = ?", (job.canonical_url,)
            ).fetchone()["id"]
        run_id = "synthetic-preparation-run"
        save_search_run(settings.database_path, run_id, criteria)
        save_run_results(
            settings.database_path,
            run_id,
            [{"id": job_id, "filter_status": "unassessed", "eligibility_status": "eligible"}],
        )
        phase = "score synthetic job with Jev"
        fit_result = score_run(
            settings.database_path,
            settings,
            run_id,
            [item for item in all_active_jobs(settings.database_path) if item["id"] == job_id],
            profile,
            criteria,
        )
        update_run(
            settings.database_path,
            run_id,
            status="complete",
            completed=True,
        )
        jev_results = get_run_results(settings.database_path, run_id)
        selected_result = next((item for item in jev_results if item["id"] == job_id), None)
        if not selected_result or selected_result.get("filter_status") != "match":
            jev_usage = _jev_usage_summary(settings.database_path)
            result = {
                "status": "stopped_before_openai",
                "reason": "Jev did not return an eligible match for the synthetic opportunity.",
                "jev_filter_status": selected_result.get("filter_status") if selected_result else "missing",
                "jev_filter_checks": selected_result.get("filter_checks", {}) if selected_result else {},
                "jev_score_reason": selected_result.get("score_reason", "") if selected_result else "",
                "jev_scored_count": fit_result.scored_count,
                "jev_usage": jev_usage,
                "openai_requests": 0,
                "synthetic_only": True,
            }
            return result

        phase = "prepare synthetic owner sources"
        technical = add_source(
            settings.database_path,
            settings,
            "synthetic-technical.md",
            "technical_profile",
            ("\n".join(item["claim"] for item in SYNTHETIC_PREPARATION_PROJECTS) + "\n").encode("utf-8"),
        )
        set_source_options(
            settings.database_path,
            technical["id"],
            permitted=True,
            default_cv=False,
            structure_policy="preserve",
        )
        suggest_claims(settings.database_path, technical["id"])
        measured_markers = tuple(
            marker
            for item in SYNTHETIC_PREPARATION_PROJECTS
            for marker in item["evidence_markers"]
        )
        for claim in list_claims(settings.database_path, technical["id"]):
            review_claim(
                settings.database_path,
                claim["id"],
                status="approved",
                evidence_level="measured_result"
                if any(marker in claim["claim_text"] for marker in measured_markers)
                else "implemented",
                category="synthetic_portfolio_project",
                role_family="applied_ai_llm",
                owner_note="Fictional candidate evidence for this isolated synthetic benchmark only.",
            )
        resume_text = (
            "Summary\nSynthetic applied AI engineer focused on retrieval, agent workflows, and deployable ML.\n"
            "Projects\n"
            + "\n".join(item["resume_line"] for item in SYNTHETIC_PREPARATION_PROJECTS)
            + "\nSkills\nPython, PyTorch, FastAPI, RAG, agent evaluation, Vitis AI, React, TypeScript\n"
        )
        cv = add_source(
            settings.database_path,
            settings,
            "synthetic-resume.md",
            "resume",
            resume_text.encode("utf-8"),
        )
        set_source_options(
            settings.database_path,
            cv["id"],
            permitted=True,
            default_cv=True,
            structure_policy="preserve",
        )
        descriptive = add_source(
            settings.database_path,
            settings,
            "synthetic-writing-preferences.md",
            "descriptive_profile",
            b"I prefer direct, modest writing. I value careful evaluation, useful documentation, and collaborative technical work.\n",
            authorship_label="owner_written",
        )
        set_source_options(
            settings.database_path,
            descriptive["id"],
            permitted=True,
            default_cv=False,
            structure_policy="preserve",
        )
        request, created = request_preparation(
            settings.database_path,
            job_id,
            run_id,
            cv["id"],
        )
        if not created:
            raise RuntimeError("Synthetic preparation trigger was not created.")

        class SyntheticCrawler:
            def __init__(self, _settings, allowed_urls):
                self.allowed_urls = set(allowed_urls)
                self.pages = {}

            def crawl(self, url, _research_question):
                if url not in self.allowed_urls:
                    raise ValueError("Synthetic page is outside the supplied URL allowlist.")
                role_text = {
                    "SYN-01": (
                        "Synthetic Northstar AI builds retrieval and agent tools for software teams. "
                        "Its junior LLM role evaluates RAG citations and model behavior with Python and PyTorch."
                    ),
                    "SYN-02": (
                        "Synthetic Alder Robotics builds perception software for mobile robots. "
                        "Its entry-level computer-vision role evaluates PyTorch models and edge inference."
                    ),
                }[job.external_id]
                page = {
                    "url": url,
                    "title": f"{job.company} engineering team",
                    "text": role_text,
                    "links": [],
                    "observed_at": utc_now(),
                }
                self.pages[url] = page
                return page

        phase = "run synthetic GPT-6 Luna preparation"
        run_preparation(
            settings.database_path,
            settings,
            request["id"],
            crawler_factory=SyntheticCrawler,
        )
        preparation = get_preparation(settings.database_path, request["id"])
        with connect(settings.database_path) as db:
            packet_row = db.execute(
                "SELECT id FROM preparation_packets WHERE request_id = ? ORDER BY revision DESC LIMIT 1",
                (request["id"],),
            ).fetchone()
        packet = get_packet(settings.database_path, packet_row["id"]) if packet_row else None
        documents = []
        generated_text = {}
        if packet:
            for artifact in packet["artifacts"]:
                path = Path(artifact["file_path"])
                if path.suffix.casefold() == ".docx":
                    document = Document(path)
                    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                    generated_text[artifact["artifact_type"]] = text
                    documents.append(
                        {
                            "artifact_type": artifact["artifact_type"],
                            "paragraph_count": len(document.paragraphs),
                            "body_paragraph_count": (
                                max(0, len(document.paragraphs) - 1)
                                if artifact["artifact_type"] == "cover_letter"
                                else None
                            ),
                            "reopened_as_docx": True,
                            "file_bytes": path.stat().st_size,
                        }
                    )
        output = packet["output"] if packet else {}
        grounding = output.get("grounding", {})
        grounding_results = [
            {
                "id": item.get("id"),
                "output_type": item.get("output_type"),
                "text": item.get("text", ""),
                "status": item.get("status", "unresolved"),
                "model_status": item.get("model_status", "unknown"),
                "confidence": item.get("confidence"),
                "reason": item.get("reason", "No reason was retained."),
                "unsupported_numbers": item.get("unsupported_numbers", []),
                "evidence_sources": [
                    {
                        "source_type": evidence.get("source_type"),
                        "source_id": evidence.get("source_id"),
                        "source_url": evidence.get("source_url", ""),
                    }
                    for evidence in item.get("evidence", [])
                ],
                "included_in_artifact": item.get("included_in_artifact", False),
            }
            for item in grounding.get("results", [])
        ]
        cover_letter = next(
            (item for item in documents if item["artifact_type"] == "cover_letter"),
            None,
        )
        if cover_letter is None:
            quality_issues = output.get("quality_check", {}).get("issues", [])
            review_reason = output.get("jev_tailored_resume_review", {}).get("reason", "")
            documents.append(
                {
                    "artifact_type": "cover_letter",
                    "generated": False,
                    "body_paragraph_count": 0,
                    "reason": (
                        quality_issues[0]
                        if quality_issues
                        else review_reason or "The packet did not pass its final quality gate."
                    ),
                }
            )
        diagnoser_challenge = {
            "status": "not_run",
            "reason": "A complete reviewable packet is required before the live parseability challenge.",
        }
        if packet:
            phase = "run synthetic Diagnoser parseability challenge"
            challenge_text = (
                "E X P E R I E N C E\n"
                "Synthetic Systems Lab | 2025\n"
                "- Built and evaluated a synthetic Python API.\n"
            )
            challenge_lines = _cv_lines(challenge_text)
            challenge_source_id = "synthetic-diagnoser-challenge"
            challenge_snapshot = preparation["snapshot"]
            challenge_context = {
                "job": _job_context(challenge_snapshot),
                "jev_read_only": _jev_read_only_context(challenge_snapshot),
                "cv": {
                    "source_id": challenge_source_id,
                    "filename": "synthetic-parseability-challenge.txt",
                    "structure_policy": "preserve",
                    "lines": challenge_lines,
                },
            }
            challenge_output = _call_stage(
                settings.database_path,
                settings,
                get_settings(settings.database_path),
                request,
                "diagnoser",
                challenge_context,
                ResponsesClient(settings.openai_api_key),
                [],
            )
            challenge_filtered = _filter_diagnoser_references(
                challenge_output,
                source_id=challenge_source_id,
                valid_line_ids={line["id"] for line in challenge_lines},
            )
            diagnoser_challenge = _diagnoser_quality_check(
                challenge_filtered,
                source_id=challenge_source_id,
                valid_line_ids={line["id"] for line in challenge_lines},
                expected_line_id=challenge_lines[0]["id"],
            )
            diagnoser_challenge["rejected_invalid_references"] = challenge_filtered[
                "rejected_invalid_references"
            ]
            diagnoser_challenge["raw_diagnostic_count"] = len(
                challenge_output.get("diagnostics", [])
                if isinstance(challenge_output.get("diagnostics", []), list)
                else []
            )
        diagnoser_line_ids_valid = all(
            item.get("line_id", "") in {line["id"] for line in _cv_lines(resume_text)}
            for item in output.get("diagnoser", {}).get("diagnostics", [])
        )
        recruiter_selected_cv_correct = (
            output.get("recruiter", {}).get("selected_cv_id") == cv["id"]
        )
        recruiter_quality = _recruiter_quality_check(
            output.get("recruiter", {}),
            selected_cv_id=cv["id"],
            valid_cv_line_ids={line["id"] for line in _cv_lines(resume_text)},
            valid_claim_ids={
                claim["id"]
                for source_id in (cv["id"], technical["id"])
                for claim in list_claims(settings.database_path, source_id)
            },
            role_keyword_groups=RECRUITER_ROLE_KEYWORD_GROUPS[role_family],
        )
        cover_letter_source_bindings = [
            {
                "source_url_count": len(item.get("source_urls", [])),
                "selected_listing_fallback_used": (
                    item.get("clue_source_url_binding") == "selected_listing_fallback"
                ),
            }
            for item in output.get("rewriter", {}).get("cover_letter_paragraphs", [])
        ]
        resume_line_order_complete = (
            output.get("rewriter", {}).get("resume_line_order")
            == [line["id"] for line in _cv_lines(resume_text)]
        )
        rewriter_output = output.get("rewriter", {})
        portfolio_resume_text = generated_text.get("resume") or _tailored_resume_text(
            _cv_lines(resume_text), rewriter_output
        )
        portfolio_letter_text = generated_text.get("cover_letter") or "\n".join(
            [rewriter_output.get("cover_letter_title", "")]
            + [
                str(item.get("text") or "")
                for item in rewriter_output.get("cover_letter_paragraphs", [])
            ]
        )
        writing_rubric = _synthetic_writing_rubric(
            grounding_results=grounding_results,
            diagnoser_line_ids_valid=diagnoser_line_ids_valid,
            recruiter_selected_cv_correct=recruiter_selected_cv_correct,
            resume_line_order_complete=resume_line_order_complete,
            artifact_checks=documents,
            tailored_resume_review=output.get("jev_tailored_resume_review"),
            packet_quality_check=output.get("quality_check"),
            portfolio_specificity=_portfolio_specificity_check(
                portfolio_resume_text,
                portfolio_letter_text,
                target_family=role_family,
            ),
        )
        portfolio_coverage = _portfolio_specificity_check(
            portfolio_resume_text,
            portfolio_letter_text,
            target_family=role_family,
        )
        hiring_manager_check = {
            "status": "not_run",
            "synthetic_answers_only": True,
            "reason": "Interview practice requires a reviewable generated packet.",
        }
        hiring_manager_session_state = "not_run"
        hiring_manager_status_message = ""
        if packet and packet.get("status") in {"review", "approved"}:
            phase = "run synthetic Hiring Manager questions"
            practice_session = create_practice_session(settings.database_path, request["id"])
            run_interview_practice_questions(
                settings.database_path,
                settings,
                practice_session["id"],
            )
            with connect(settings.database_path) as db:
                question_row = db.execute(
                    "SELECT state, output_json FROM practice_sessions WHERE id = ?",
                    (practice_session["id"],),
                ).fetchone()
            question_state = question_row["state"] if question_row else "missing"
            question_output = json.loads(question_row["output_json"]) if question_row else {}
            practice_questions = question_output.get("questions", [])
            synthetic_answers = [
                {
                    "question": question,
                    "answer": (
                        "For this synthetic practice case, I would define the target behavior, "
                        "establish a representative baseline, measure errors and trade-offs, and "
                        "state clearly which results have not been measured."
                    ),
                }
                for question in practice_questions
                if isinstance(question, str) and question.strip()
            ]
            if question_state == "questions" and synthetic_answers:
                queue_practice_answers(
                    settings.database_path,
                    practice_session["id"],
                    json.dumps(synthetic_answers, ensure_ascii=False),
                )
                phase = "run synthetic Hiring Manager answer assessment"
                run_interview_practice_assessment(
                    settings.database_path,
                    settings,
                    practice_session["id"],
                )
            with connect(settings.database_path) as db:
                assessment_row = db.execute(
                    "SELECT state, output_json, status_message FROM practice_sessions WHERE id = ?",
                    (practice_session["id"],),
                ).fetchone()
            hiring_manager_session_state = assessment_row["state"] if assessment_row else "missing"
            hiring_manager_status_message = (
                assessment_row["status_message"] if assessment_row else "Practice session was not saved."
            )
            assessment_output = (
                json.loads(assessment_row["output_json"]) if assessment_row else {}
            )
            hiring_manager_check = _hiring_manager_quality_check(
                session_state=hiring_manager_session_state,
                questions=practice_questions,
                expected_answers=synthetic_answers,
                output=assessment_output,
                role_keyword_groups=HIRING_MANAGER_ROLE_KEYWORD_GROUPS[role_family],
            )
        with connect(settings.database_path) as db:
            openai_rows = db.execute(
                """SELECT stage, status, COUNT(*) AS calls,
                          SUM(actual_usd) AS actual_cost_usd,
                          SUM(reserved_usd) AS reserved_usd,
                          SUM(reserved_input_tokens) AS reserved_input_tokens,
                          SUM(max_output_tokens) AS reserved_output_tokens,
                          SUM(input_tokens) AS input_tokens,
                          SUM(output_tokens) AS output_tokens
                   FROM openai_usage GROUP BY stage, status ORDER BY stage, status"""
            ).fetchall()
            research_pages = db.execute(
                "SELECT page_url, page_text FROM research_pages WHERE request_id = ?",
                (request["id"],),
            ).fetchall()
        research = output.get("research", {})
        research_findings = research.get("findings", [])
        research_contacts = research.get("contacts", [])
        research_text_by_url = {row["page_url"]: row["page_text"] for row in research_pages}
        researcher_provenance_valid = (
            bool(research_findings)
            and bool(research_text_by_url)
            and all(
                item.get("source_url") in research_text_by_url
                and _quote_supported(
                    item.get("quote", ""),
                    research_text_by_url[item.get("source_url", "")],
                )
                for item in research_findings
            )
            and all(
                contact.get("source_url") in research_text_by_url
                and (
                    not contact.get("public_email")
                    or contact["public_email"].casefold()
                    in research_text_by_url[contact["source_url"]].casefold()
                )
                for contact in research_contacts
            )
            and (bool(research_contacts) or bool(research.get("no_contact_found_reason")))
        )
        diagnoser_actual_quality = _diagnoser_quality_check(
            output.get("diagnoser", {}),
            source_id=cv["id"],
            valid_line_ids={line["id"] for line in _cv_lines(resume_text)},
        )
        diagnoser_agent_passed = (
            diagnoser_actual_quality["status"] == "pass"
            and diagnoser_challenge.get("status") == "pass"
            and output.get("diagnoser", {}).get("rejected_invalid_references", 0) == 0
        )
        result = {
            "status": preparation["state"],
            "synthetic_only": True,
            "triggered_listing_count": 1,
            "synthetic_case_id": case_id,
            "target_role_family": role_family,
            "jev_filter_status": selected_result.get("filter_status"),
            "jev_fit_score": selected_result.get("combined_score"),
            "jev_filter_checks": selected_result.get("filter_checks", {}),
            "jev_score_reason": selected_result.get("score_reason", ""),
            "jev_usage": _jev_usage_summary(settings.database_path),
            "openai_usage_by_stage": [dict(row) for row in openai_rows],
            "openai_usage_total_usd": sum(
                float(row["actual_cost_usd"] or 0)
                for row in openai_rows
                if row["status"] == "settled"
            ),
            "openai_usage_unresolved": any(
                row["status"] in {"reserved", "unknown"} for row in openai_rows
            ),
            "openai_budget_usd": 0.10,
            "per_opportunity_budget_usd": 0.05,
            "request_status_message": preparation.get("status_message", ""),
            "packet_status": packet.get("status") if packet else "none",
            "jev_tailored_resume_review": output.get("jev_tailored_resume_review"),
            "packet_quality_check": output.get("quality_check"),
            "artifact_checks": documents,
            "grounding_counts": {
                "total": grounding.get("total_assertions", 0),
                "supported": grounding.get("supported_count", 0),
                "omitted": grounding.get("omitted_count", 0),
            },
            "grounding_results": grounding_results,
            "grounding_summary": [
                {
                    "id": item["id"],
                    "output_type": item["output_type"],
                    "status": item["status"],
                    "model_status": item["model_status"],
                    "reason": item["reason"],
                    "unsupported_numbers": item["unsupported_numbers"],
                    "evidence_source_types": sorted(
                        {
                            evidence.get("source_type", "unknown")
                            for evidence in item["evidence_sources"]
                        }
                    ),
                    "evidence_source_ids": sorted(
                        {
                            evidence.get("source_id", "")
                            for evidence in item["evidence_sources"]
                            if evidence.get("source_id")
                        }
                    ),
                    "included_in_artifact": item["included_in_artifact"],
                }
                for item in grounding_results
            ],
            "diagnoser_line_ids_valid": diagnoser_line_ids_valid,
            "diagnoser_actual_resume": diagnoser_actual_quality,
            "diagnoser_parseability_challenge": diagnoser_challenge,
            "recruiter_selected_cv_correct": recruiter_selected_cv_correct,
            "recruiter_quality": recruiter_quality,
            "cover_letter_source_bindings": cover_letter_source_bindings,
            "resume_line_order_complete": resume_line_order_complete,
            "writing_rubric": writing_rubric,
            "portfolio_coverage": portfolio_coverage,
            "hiring_manager_practice": hiring_manager_check,
            "hiring_manager_session_state": hiring_manager_session_state,
            "hiring_manager_status_message": hiring_manager_status_message,
            "agent_evaluation": {
                "diagnoser": {
                    "status": "pass" if diagnoser_agent_passed else "fail",
                    "cv_line_ids_valid": diagnoser_line_ids_valid,
                    "diagnostic_count": len(output.get("diagnoser", {}).get("diagnostics", [])),
                    "invalid_references_rejected": output.get("diagnoser", {}).get("rejected_invalid_references", 0),
                    "known_defect_challenge": diagnoser_challenge,
                },
                "recruiter": {
                    "status": recruiter_quality["status"],
                    "selected_cv_correct": recruiter_selected_cv_correct,
                    **{key: value for key, value in recruiter_quality.items() if key != "status"},
                },
                "researcher": {
                    "status": "pass" if researcher_provenance_valid else "fail",
                    "source_page_count": len(research_pages),
                    "finding_count": len(research_findings),
                    "contact_count": len(research_contacts),
                    "verified_provenance": researcher_provenance_valid,
                    "quote_support_verified": all(
                        item.get("source_url") in research_text_by_url
                        and _quote_supported(
                            item.get("quote", ""),
                            research_text_by_url[item.get("source_url", "")],
                        )
                        for item in research_findings
                    ),
                    "no_contact_reason_present": bool(research.get("no_contact_found_reason")),
                },
                "rewriter": {
                    "status": "pass"
                    if resume_line_order_complete
                    and output.get("jev_tailored_resume_review", {}).get("status") == "approved"
                    and output.get("quality_check", {}).get("status") == "passed"
                    else "fail",
                    "source_line_order_complete": resume_line_order_complete,
                    "tailored_resume_approved": output.get("jev_tailored_resume_review", {}).get("status") == "approved",
                    "packet_quality_passed": output.get("quality_check", {}).get("status") == "passed",
                    "cover_letter_paragraph_count": len(cover_letter_source_bindings),
                    "role_source_fallback_count": sum(
                        item["selected_listing_fallback_used"]
                        for item in cover_letter_source_bindings
                    ),
                    "all_cover_letter_paragraphs_have_role_sources": all(
                        item["source_url_count"] > 0 for item in cover_letter_source_bindings
                    ),
                },
                "hiring_manager": hiring_manager_check,
            },
            "generation_prompt_version": output.get("prompt_version"),
            "output_schema_version": output.get("output_schema_version"),
            "model": output.get("model"),
            "reasoning_effort": output.get("reasoning_effort"),
            "generated_text": generated_text,
            "visual_render": "not rendered by the live runner; page inspection is a separate acceptance check",
            "submission_or_outreach_sent": False,
        }
        return result
    except BenchmarkRunError as exc:
        temporary.cleanup()
        if temporary.usage_recovery and not exc.usage_recovery:
            raise BenchmarkRunError(
                exc.phase,
                exc.error_type,
                usage_recovery=temporary.usage_recovery,
            ) from None
        raise
    except Exception as exc:  # noqa: BLE001 - keep the result free of synthetic input and provider response bodies.
        temporary.cleanup()
        raise BenchmarkRunError(
            phase,
            type(exc).__name__,
            usage_recovery=temporary.usage_recovery,
        ) from None
    finally:
        temporary.cleanup()
        if result is not None and temporary.usage_recovery:
            result["usage_recovery"] = temporary.usage_recovery


def _portfolio_specificity_check(
    resume_text: str,
    cover_letter_text: str,
    *,
    target_family: str | None = None,
) -> dict[str, Any]:
    """Check broad source-project retention and job-family-specific letter evidence."""
    resume = " ".join(resume_text.casefold().split())
    letter = " ".join(cover_letter_text.casefold().split())
    resume_projects = [
        item for item in SYNTHETIC_PREPARATION_PROJECTS if item["name"].casefold() in resume
    ]
    letter_projects = [
        item for item in SYNTHETIC_PREPARATION_PROJECTS if item["name"].casefold() in letter
    ]
    evidence_projects = [
        item
        for item in letter_projects
        if any(marker.casefold() in letter for marker in item["evidence_markers"])
    ]
    relevant_projects = [
        item
        for item in SYNTHETIC_PREPARATION_PROJECTS
        if target_family is None or target_family in item["role_families"]
    ]
    relevant_names = {item["name"] for item in relevant_projects}
    resume_relevant_projects = [item for item in resume_projects if item["name"] in relevant_names]
    letter_relevant_projects = [item for item in evidence_projects if item["name"] in relevant_names]
    status = (
        "pass"
        if len(resume_projects) >= 8
        and resume_relevant_projects
        and (letter_relevant_projects if target_family else evidence_projects)
        else "fail"
    )
    return {
        "status": status,
        "target_family": target_family,
        "source_project_count": len(SYNTHETIC_PREPARATION_PROJECTS),
        "resume_project_count": len(resume_projects),
        "resume_projects": [item["name"] for item in resume_projects],
        "relevant_source_project_count": len(relevant_projects),
        "resume_relevant_project_count": len(resume_relevant_projects),
        "resume_relevant_projects": [item["name"] for item in resume_relevant_projects],
        "cover_letter_project_count": len(letter_projects),
        "cover_letter_projects": [item["name"] for item in letter_projects],
        "cover_letter_evidence_marker_count": sum(
            1
            for item in evidence_projects
            for marker in item["evidence_markers"]
            if marker.casefold() in letter
        ),
        "cover_letter_projects_with_supported_detail": [
            item["name"] for item in evidence_projects
        ],
        "cover_letter_relevant_projects_with_supported_detail": [
            item["name"] for item in letter_relevant_projects
        ],
    }


def _recruiter_quality_check(
    output: dict[str, Any],
    *,
    selected_cv_id: str,
    valid_cv_line_ids: set[str],
    valid_claim_ids: set[str],
    role_keyword_groups: tuple[tuple[str, ...], ...],
) -> dict[str, Any]:
    requirements = output.get("requirements", [])
    requirements = requirements if isinstance(requirements, list) else []
    references_valid = all(
        isinstance(item, dict)
        and set(item.get("cv_line_ids", [])) <= valid_cv_line_ids
        and set(item.get("claim_ids", [])) <= valid_claim_ids
        for item in requirements
    )
    coverage_cited = all(
        item.get("document_coverage") not in {"covered", "partial"}
        or bool(item.get("cv_line_ids"))
        for item in requirements
        if isinstance(item, dict)
    )
    requirement_text = " ".join(
        str(item.get("requirement") or "")
        for item in requirements
        if isinstance(item, dict)
    ).casefold()
    matched_groups = sum(
        any(term in requirement_text for term in group)
        for group in role_keyword_groups
    )
    role_coverage = matched_groups / len(role_keyword_groups) if role_keyword_groups else 1.0
    status = (
        "pass"
        if output.get("selected_cv_id") == selected_cv_id
        and len(requirements) >= 2
        and references_valid
        and coverage_cited
        and role_coverage >= 0.5
        and bool(str(output.get("coverage_note") or "").strip())
        else "fail"
    )
    return {
        "status": status,
        "selected_cv_correct": output.get("selected_cv_id") == selected_cv_id,
        "requirement_count": len(requirements),
        "references_valid": references_valid,
        "covered_or_partial_requirements_have_cv_lines": coverage_cited,
        "role_keyword_groups_matched": matched_groups,
        "role_keyword_group_count": len(role_keyword_groups),
        "role_keyword_coverage": role_coverage,
        "coverage_note_present": bool(str(output.get("coverage_note") or "").strip()),
    }


def _diagnoser_quality_check(
    output: dict[str, Any],
    *,
    source_id: str,
    valid_line_ids: set[str],
    expected_line_id: str | None = None,
) -> dict[str, Any]:
    diagnostics = output.get("diagnostics", [])
    diagnostics = diagnostics if isinstance(diagnostics, list) else []
    references_valid = all(
        isinstance(item, dict)
        and item.get("source_id") == source_id
        and item.get("line_id") in valid_line_ids
        for item in diagnostics
    )
    note_present = bool(str(output.get("overall_note") or "").strip())
    if expected_line_id is None:
        # The portfolio fixture has no seeded extraction defect. Empty findings are
        # allowed, but this is not a detection-recall measurement.
        status = "pass" if references_valid and note_present else "fail"
        return {
            "status": status,
            "diagnostic_count": len(diagnostics),
            "references_valid": references_valid,
            "overall_note_present": note_present,
            "detection_recall_measured": False,
        }

    detected = any(
        isinstance(item, dict)
        and item.get("source_id") == source_id
        and item.get("line_id") == expected_line_id
        and any(word in str(item.get("issue") or "").casefold() for word in ("heading", "header", "section"))
        and any(
            word in (str(item.get("issue") or "") + " " + str(item.get("suggested_fix") or "")).casefold()
            for word in ("spac", "split", "separat", "standard")
        )
        for item in diagnostics
    )
    status = "pass" if references_valid and note_present and detected else "fail"
    return {
        "status": status,
        "diagnostic_count": len(diagnostics),
        "references_valid": references_valid,
        "overall_note_present": note_present,
        "expected_line_detected": detected,
        "expected_line_id": expected_line_id,
        "detection_recall_measured": True,
    }


def _hiring_manager_quality_check(
    *,
    session_state: str,
    questions: list[str],
    expected_answers: list[dict[str, str]],
    output: dict[str, Any],
    role_keyword_groups: tuple[tuple[str, ...], ...] = (),
) -> dict[str, Any]:
    raw_assessments = output.get("answer_assessments", [])
    assessments = raw_assessments if isinstance(raw_assessments, list) else []
    fields = ("technical_evidence", "reasoning", "clarity", "suggested_practice")
    questions_preserved = output.get("questions") == questions
    answers_preserved = len(assessments) == len(expected_answers) and all(
        isinstance(item, dict)
        and item.get("question") == answer["question"]
        and item.get("answer") == answer["answer"]
        for item, answer in zip(assessments, expected_answers, strict=False)
    )
    assessment_fields_complete = len(assessments) == len(expected_answers) and all(
        isinstance(item, dict) and all(str(item.get(field, "")).strip() for field in fields)
        for item in assessments
    )
    role_specific_question_count = sum(
        any(term in question.casefold() for group in role_keyword_groups for term in group)
        for question in questions
    )
    required_role_specific_question_count = math.ceil(0.6 * len(questions)) if questions else 0
    questions_are_role_specific = (
        role_specific_question_count >= required_role_specific_question_count
        if role_keyword_groups
        else True
    )
    questions_are_distinct = len({question.casefold().strip() for question in questions}) == len(questions)
    status = (
        "pass"
        if session_state == "complete"
        and 3 <= len(questions) <= 5
        and len(expected_answers) == len(questions)
        and questions_preserved
        and questions_are_role_specific
        and questions_are_distinct
        and answers_preserved
        and assessment_fields_complete
        else "fail"
    )
    return {
        "status": status,
        "synthetic_answers_only": True,
        "question_count": len(questions),
        "answer_assessment_count": len(assessments),
        "questions_preserved": questions_preserved,
        "role_specific_question_count": role_specific_question_count,
        "questions_are_role_specific": questions_are_role_specific,
        "questions_are_distinct": questions_are_distinct,
        "synthetic_answers_preserved": answers_preserved,
        "assessment_fields_complete": assessment_fields_complete,
    }


def _synthetic_writing_rubric(
    *,
    grounding_results: list[dict[str, Any]],
    diagnoser_line_ids_valid: bool,
    recruiter_selected_cv_correct: bool,
    resume_line_order_complete: bool,
    artifact_checks: list[dict[str, Any]],
    tailored_resume_review: dict[str, Any] | None = None,
    packet_quality_check: dict[str, Any] | None = None,
    portfolio_specificity: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Score mechanical writing safeguards; leave subjective voice ratings to a human."""
    cover_letter = next(
        (item for item in artifact_checks if item["artifact_type"] == "cover_letter"),
        None,
    )
    body_paragraphs = int((cover_letter or {}).get("body_paragraph_count") or 0)
    tailored_resume_approved = (tailored_resume_review or {}).get("status") == "approved"
    packet_quality_passed = (packet_quality_check or {}).get("status") == "passed"
    portfolio_specificity = portfolio_specificity or {"status": "human_review_required"}
    evidence_gate_ok = all(
        bool(item.get("included_in_artifact")) == (item.get("status") == "supported")
        for item in grounding_results
    )
    provenance_ok = all(
        bool(item.get("evidence_sources"))
        for item in grounding_results
        if item.get("included_in_artifact")
    )
    mechanical_gates = {
        "only_jev_supported_blocks_included": "pass" if evidence_gate_ok else "fail",
        "diagnoser_line_references": "pass" if diagnoser_line_ids_valid else "fail",
        "recruiter_cv_selection": "pass" if recruiter_selected_cv_correct else "fail",
        "resume_source_line_order": "pass" if resume_line_order_complete else "fail",
        "cover_letter_body_present": "pass" if body_paragraphs >= 2 else "fail",
        "tailored_resume_jev_approval": "pass" if tailored_resume_approved else "fail",
        "packet_quality_gate": "pass" if packet_quality_passed else "fail",
        "portfolio_project_coverage": portfolio_specificity["status"],
    }
    criteria = {
        "evidence_traceability": "pass" if provenance_ok else "fail",
        "factual_accuracy": "pass" if evidence_gate_ok else "fail",
        "role_specificity": "fail" if body_paragraphs < 2 else "human_review_required",
        "tailored_resume_jev_approval": "pass" if tailored_resume_approved else "fail",
        "packet_quality": "pass" if packet_quality_passed else "fail",
        "voice": "human_review_required",
        "xyz_integrity": "human_review_required",
        "question_adherence": "not_applicable_no_questions",
        "research_provenance": "human_review_required",
        "concision": "human_review_required",
        "portfolio_specificity": portfolio_specificity["status"],
    }
    status = (
        "fail"
        if "fail" in (*mechanical_gates.values(), *criteria.values())
        else "incomplete_human_review"
    )
    return {
        "status": status,
        "criteria": criteria,
        "definitions": WRITING_RUBRIC,
        "mechanical_gates": mechanical_gates,
        "human_review_note": (
            "The harness checks provenance and document integrity; it does not treat a model's "
            "self-assessment as a reliable score for voice, relevance, concision, or XYZ quality."
        ),
    }


def _jev_usage_summary(database_path: Path) -> dict[str, Any]:
    from clue_ai.repository import connect

    with connect(database_path) as db:
        row = db.execute(
            """SELECT COUNT(*) AS requests,
                      SUM(CASE WHEN actual_usd IS NOT NULL THEN actual_usd ELSE reserved_usd END) AS cost_usd,
                      SUM(CASE WHEN status = 'reserved' THEN reserved_usd ELSE 0 END) AS open_reserve_usd
               FROM jev_usage"""
        ).fetchone()
    return {
        "requests": int(row["requests"] or 0),
        "cost_usd": float(row["cost_usd"] or 0),
        "open_reserve_usd": float(row["open_reserve_usd"] or 0),
    }


def _usage_rows(database_path: Path) -> list[float]:
    from clue_ai.repository import connect

    with connect(database_path) as db:
        return [
            float(row["reserved_usd"] or 0)
            for row in db.execute("SELECT reserved_usd FROM jev_usage WHERE status = 'reserved'")
        ]


def main() -> int:
    parser = argparse.ArgumentParser(description="Run synthetic-only Clue agent evaluation benchmarks.")
    parser.add_argument(
        "--live",
        action="store_true",
        help=(
            "Send this synthetic benchmark to Jev once; delete local data only after provider "
            "usage is settled, and retain unsettled ledgers for reconciliation."
        ),
    )
    parser.add_argument(
        "--max-budget-usd",
        type=float,
        default=MAX_EVALUATION_BUDGET_USD,
        help=f"Evaluation-only hard cap in USD (maximum {MAX_EVALUATION_BUDGET_USD:.2f}).",
    )
    parser.add_argument(
        "--grounding",
        action="store_true",
        help="Run the synthetic claim-support benchmark instead of the job-fit benchmark.",
    )
    parser.add_argument(
        "--live-preparation",
        action="store_true",
        help="Run one full synthetic application and Hiring Manager practice workflow.",
    )
    parser.add_argument(
        "--preparation-case",
        choices=tuple(SYNTHETIC_PREPARATION_ROLE_FAMILIES),
        default="SYN-01",
        help="Synthetic role case to use with --live-preparation.",
    )
    args = parser.parse_args()
    if not args.live:
        print("No API request made. Review the synthetic data, then pass --live to call the selected synthetic benchmark.")
        return 0
    if args.live_preparation and args.grounding:
        parser.error("--live-preparation and --grounding are separate benchmark modes.")
    try:
        if args.live_preparation:
            report = run_live_preparation_benchmark(args.preparation_case)
        elif args.grounding:
            report = run_live_grounding_benchmark(max_budget_usd=args.max_budget_usd)
        else:
            report = run_live_benchmark(max_budget_usd=args.max_budget_usd)
    except Exception as exc:  # noqa: BLE001 - keep CLI output free of request details and credentials.
        phase = exc.phase if isinstance(exc, BenchmarkRunError) else "initialize benchmark"
        error_type = exc.error_type if isinstance(exc, BenchmarkRunError) else type(exc).__name__
        print(json.dumps({"error": error_type, "phase": phase, "request_details_printed": False}))
        return 1
    print(json.dumps(report, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
